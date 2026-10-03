"""Marketplace import adapters.

Every adapter shares one header-driven pipeline:

    detect()    filename signature OR workbook sheet-name signature
    validate()  explicit, actionable errors for malformed / unsupported files
    parse()     per-row ImportRow records (a row is never silently dropped)
    normalize() canonical transaction contract consumed by the import pipeline

Platform-specific column knowledge is taken only from evidence that already
exists in this repository:

  * ``PLATFORM_ADAPTER_GUIDE.md``   - Amazon MTR, Flipkart GST report,
    Meesho orders report, Myntra (same layout as Flipkart), Custom Excel,
    GSTR-1 govt multi-sheet workbook.
  * ``README_GST_ONLINE_SELLER.md`` - GSTR-1 workbook sheet names.
  * ``app/templates/import.html``   - marketplaces offered in the UI
    (Amazon, Flipkart, Meesho, Myntra).
  * ``tests/fixtures/generate_test_data.py`` - the concrete headers written for
    each marketplace fixture.
  * ``app/services/gstr1_excel_writer.py``   - the headers of the GSTR-1
    workbook this project itself produces (used by the govt-sheet adapter).

Adapters that have no documented column evidence (JioMart, Snapdeal, Tata CLiQ,
Limeroad, Shopdeck, GlowRoad, Snapmint, Citymall, Roposo) stay filename-detected
adapters that fall back to the generic header matching. Their formats are NOT
invented here; they are reported by ``FORMAT_DOCUMENTED = False``.
"""
import csv
import io
import json
import os
import re
from collections import Counter
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterator, List, Optional, Sequence, Tuple

import openpyxl

from .base import ImportResult, ImportRow, ImportRowStatus, PlatformAdapter
from .canonical import CanonicalTransaction
from app.utils.csv_utils import read_csv_rows
from app.utils.date_utils import format_date_gst, parse_date
from app.utils.gstin_validator import validate_gstin
from app.utils.state_codes import STATE_CODES, resolve_pos_code


# ---------------------------------------------------------------------------
# Canonical contract
# ---------------------------------------------------------------------------

CANONICAL_FIELDS: Tuple[str, ...] = (
    'order_id', 'invoice_number', 'invoice_date', 'customer_name',
    'customer_gstin', 'place_of_supply', 'seller_gstin', 'hsn_sac',
    'description', 'quantity', 'taxable_value', 'cgst_rate', 'cgst_amount',
    'sgst_rate', 'sgst_amount', 'igst_rate', 'igst_amount', 'cess_rate',
    'cess_amount', 'tax_rate', 'invoice_value', 'reverse_charge',
    'ecommerce_gstin', 'note_type', 'cancellation_flag', 'return_flag',
)

FIELD_LABELS: Dict[str, str] = {
    'invoice_number': 'invoice number',
    'taxable_value': 'taxable value',
    'invoice_date': 'invoice date',
    'hsn_sac': 'HSN/SAC',
    'quantity': 'quantity',
    'invoice_value': 'invoice value',
    'customer_gstin': 'customer GSTIN',
    'place_of_supply': 'place of supply',
}

# Generic alias table: canonical field -> column names seen in generic
# marketplace / custom Excel exports. Used as the fallback when an adapter has
# no platform-specific mapping for a column, and by the Custom Excel adapter.
GENERIC_ALIASES: Dict[str, List[str]] = {
    'order_id': ['Order ID', 'Order Id', 'Order Number', 'OrderItemId',
                 'Order Item ID', 'Sub Order ID', 'Marketplace Order ID'],
    'order_item_id': ['Order Item ID', 'OrderItemId', 'Sub Order ID'],
    'sub_order_id': ['Sub Order ID'],
    'invoice_number': ['Invoice Number', 'Invoice No', 'Invoice No.',
                       'InvoiceNumber', 'Bill Number', 'Tax Invoice Number',
                       'Invoice ID'],
    'invoice_date': ['Invoice Date', 'Invoice date', 'Date', 'Order Date',
                     'Invoice Date (DD-MM-YYYY)'],
    'invoice_type': ['Invoice Type'],
    'document_type': ['Document Type', 'Doc Type'],
    'customer_name': ['Customer Name', 'Buyer Name', 'Bill To Name',
                      'Customer', 'Shipping Name'],
    'customer_gstin': ['Buyer GSTIN', 'Customer GSTIN', 'GSTIN/UIN of Recipient',
                       'Customer GST No', 'Buyer GST No', 'GSTIN', 'GSTIN/UIN'],
    'place_of_supply': ['Place Of Supply', 'Place of Supply', 'POS',
                        'Ship To State', 'Buyer State', 'Customer State',
                        'State', 'Ship To State Code'],
    'ship_to_state': ['Ship To State', 'Shipping State', 'Buyer State'],
    'ship_from_state': ['Ship From State', 'Shipping From State', 'Seller State'],
    'buyer_state': ['Buyer State', 'Customer State'],
    'seller_gstin': ['Seller GSTIN', 'Seller GST No', 'GSTIN of Supplier',
                     'Supplier GSTIN'],
    'hsn_sac': ['HSN/SAC', 'HSN', 'HSN Code', 'HSN/SAC Code', 'SAC', 'SAC Code'],
    'description': ['Product Description', 'Product Title', 'Product Name',
                    'Item Description', 'Description', 'Item Name'],
    'quantity': ['Quantity', 'Qty', 'Item Quantity', 'Total Quantity'],
    'uqc': ['UQC', 'Unit', 'Unit of Measure'],
    'item_code': ['SKU', 'SKU Code', 'Item Code', 'Seller SKU'],
    'unit_price': ['Selling Price', 'Unit Price', 'Price', 'Rate per Unit'],
    'discount': ['Discount', 'Discount Amount', 'Item Discount'],
    'taxable_value': ['Taxable Value', 'Principal Amount', 'Taxable Amount',
                      'Taxable Value (Rs)', 'Assessable Value'],
    'cgst_rate': ['CGST Rate', 'Central Tax Rate'],
    'cgst_amount': ['CGST Amount', 'CGST', 'Central Tax', 'Central Tax Amount'],
    'sgst_rate': ['SGST Rate', 'State/UT Tax Rate', 'State Tax Rate'],
    'sgst_amount': ['SGST Amount', 'SGST', 'State/UT Tax', 'State Tax',
                    'State/UT Tax Amount'],
    'igst_rate': ['IGST Rate', 'Integrated Tax Rate'],
    'igst_amount': ['IGST Amount', 'IGST', 'Integrated Tax',
                    'Integrated Tax Amount'],
    'cess_rate': ['Cess Rate', 'CESS Rate'],
    'cess_amount': ['Cess Amount', 'CESS', 'Cess'],
    'tax_rate': ['Tax Rate', 'GST Rate', 'Total Tax Rate', 'Rate'],
    'invoice_value': ['Invoice Amount', 'Invoice Value', 'Total Amount',
                      'Total', 'Gross Amount'],
    'reverse_charge': ['Reverse Charge', 'Reverse charge', 'Reverse Charge (Y/N)'],
    'ecommerce_gstin': ['E-Commerce GSTIN', 'Ecommerce GSTIN',
                        'E-Commerce Operator GSTIN', 'Marketplace GSTIN'],
    'marketplace_name': ['Marketplace', 'Marketplace Name', 'Source',
                         'Source Platform'],
    'supply_type': ['Supply Type', 'Supply Category'],
    'order_status': ['Order Status', 'Status', 'Return Status', 'Return',
                     'Order Item Status'],
    'note_type': ['Note Type', 'Document Type', 'Note/Refund Voucher Type'],
    'note_number': ['Note Number', 'Note/Refund Voucher Number',
                    'Credit Note No', 'Debit Note No', 'Refund Id',
                    'Credit Note Number', 'Debit Note Number'],
    'note_date': ['Note Date', 'Note/Refund Voucher date', 'Credit Note Date'],
    'original_invoice_number': ['Original Invoice Number',
                                'Invoice/Advance Receipt Number',
                                'Original Invoice No'],
    'original_invoice_date': ['Original Invoice Date',
                              'Invoice/Advance Receipt date',
                              'Original Invoice date'],
}

def _key(value: Any) -> str:
    """Normalise a header name for matching: lower case, alphanumerics only."""
    return re.sub(r'[^a-z0-9]+', ' ', str(value if value is not None else '').lower()).strip()


# Canonical aliases may themselves repeat a column name (e.g. 'Return' is both a
# status column and a flag). First mapping wins; keep the table order stable.
_ALIAS_INDEX: Dict[str, str] = {}
for _field, _aliases in GENERIC_ALIASES.items():
    for _alias in _aliases:
        _ALIAS_INDEX.setdefault(_key(_alias), _field)


_HEADER_MAP_CACHE: Dict[type, Dict[str, str]] = {}


def _platform_index(adapter_cls: type) -> Dict[str, str]:
    """Per-class index of platform header name -> canonical field."""
    index = _HEADER_MAP_CACHE.get(adapter_cls)
    if index is None:
        index = {}
        for header, field in (getattr(adapter_cls, 'HEADER_MAP', None) or {}).items():
            index.setdefault(_key(header), field)
        _HEADER_MAP_CACHE[adapter_cls] = index
    return index


class SheetSource:
    """Uniform, re-iterable view over one worksheet / CSV sheet / row list."""

    def __init__(self, name: str, row_source):
        self.name = name
        self._row_source = row_source

    def rows(self) -> Iterator[Sequence[Any]]:
        data = self._row_source() if callable(self._row_source) else self._row_source
        return iter(data)


class BaseGenericAdapter(PlatformAdapter):
    """Header-driven adapter shared by every platform.

    Subclasses declare platform knowledge with class attributes only:

    ``FILENAME_TOKENS``  filename substrings that identify the platform
    ``SHEET_NAMES``      sheet names that identify the platform's report
    ``HEADER_MAP``       platform column header -> canonical field
    ``REQUIRED_COLUMNS`` canonical columns the file must expose to be accepted
    """

    PLATFORM_NAME = 'Generic'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Upload the marketplace GST / sales report. Column headers are '
                    'matched automatically; unrecognised columns are reported as warnings.')
    TEMPLATE_URL = ''
    FORMAT_DOCUMENTED = False

    FILENAME_TOKENS: Tuple[str, ...] = ()
    SHEET_NAMES: Tuple[str, ...] = ()
    HEADER_MAP: Dict[str, str] = {}
    NON_TRANSACTION_SHEETS: Tuple[str, ...] = ()
    CONSOLIDATED_SHEETS: Tuple[str, ...] = ()
    NOTE_SHEETS: Tuple[str, ...] = ()
    REQUIRED_COLUMNS: Tuple[str, ...] = ('invoice_number',)
    REQUIRED_ROW_FIELDS: Tuple[str, ...] = ('invoice_number',)
    MULTI_SHEET = False
    REQUIRE_ANY_SHEET = False
    CATCH_ALL_DETECT = False
    MAX_HEADER_SCAN_ROWS = 10
    MIN_MAPPED_COLUMNS = 2

    # ------------------------------------------------------------------
    # source loading
    # ------------------------------------------------------------------
    def _load_sheets(self, workbook_or_data, file_name: str = ''):
        """Return (sheets, file_name, load_errors, diagnostics, wb)."""
        if workbook_or_data is None:
            return [], file_name, ['No workbook or data provided.'], {}, None

        if isinstance(workbook_or_data, openpyxl.Workbook):
            return list(self._sheets_from_workbook(workbook_or_data)), file_name, [], {}, workbook_or_data

        if isinstance(workbook_or_data, (bytes, io.BytesIO)):
            try:
                stream = io.BytesIO(workbook_or_data) if isinstance(workbook_or_data, bytes) else workbook_or_data
                workbook = openpyxl.load_workbook(stream, data_only=True, read_only=True)
                diagnostics = {'source_type': 'excel'}
                return list(self._sheets_from_workbook(workbook)), file_name, [], diagnostics, workbook
            except Exception as exc:
                return [], file_name, [f'Unsupported or unreadable file: {exc}'], {}, None

        if isinstance(workbook_or_data, (list, tuple)):
            rows = list(workbook_or_data)
            if not rows:
                return [], file_name, ['No data rows provided.'], {}, None
            if isinstance(rows[0], dict):
                headers = list(rows[0].keys())
                data = [headers] + [[row.get(header) for header in headers] for row in rows]
                return [SheetSource('Data', data)], file_name, [], {}, None
            return [SheetSource('Data', rows)], file_name, [], {}, None

        if isinstance(workbook_or_data, str):
            path = workbook_or_data
            file_name = file_name or os.path.basename(path)
            extension = os.path.splitext(path)[1].lower()
            if not os.path.exists(path):
                return [], file_name, [f'File not found: {path}'], {}, None
            if extension == '.csv':
                csv_result = read_csv_rows(path)
                if csv_result.errors:
                    return [], file_name, list(csv_result.errors), {}, None
                diagnostics = {
                    'source_type': 'csv',
                    'csv_encoding': csv_result.encoding,
                    'csv_delimiter': csv_result.delimiter_label,
                }
                return [SheetSource(os.path.basename(path), csv_result.rows)], file_name, [], diagnostics, None
            if extension not in ('.xlsx', '.xlsm', '.xltx'):
                return [], file_name, [
                    f'Unsupported file type {extension or "(none)"}: expected an Excel '
                    f'(.xlsx/.xlsm) or .csv export.'
                ], {}, None
            try:
                workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
            except Exception as exc:
                return [], file_name, [f'Unsupported or unreadable file {path}: {exc}'], {}, None
            diagnostics = {'source_type': 'excel'}
            return list(self._sheets_from_workbook(workbook)), file_name, [], diagnostics, workbook

        return [], file_name, [
            f'Unsupported data type {type(workbook_or_data).__name__}: expected an '
            f'openpyxl Workbook, a .xlsx/.csv path, or a list of rows.'
        ], {}, None

    @staticmethod
    def _sheets_from_workbook(workbook: openpyxl.Workbook) -> Iterator[SheetSource]:
        worksheets = list(workbook.worksheets)
        active = workbook.active
        if active is not None and active in worksheets:
            worksheets.remove(active)
            worksheets.insert(0, active)  # active sheet first: previous behaviour
        for worksheet in worksheets:
            yield SheetSource(
                worksheet.title,
                (lambda ws=worksheet: ws.iter_rows(values_only=True)),
            )

    # ------------------------------------------------------------------
    # header detection / column mapping
    # ------------------------------------------------------------------
    def resolve_headers(self, headers: Sequence[Any]) -> Tuple[Dict[str, int], List[str]]:
        """Map canonical field -> column position; collect unmatched headers."""
        mapping: Dict[str, int] = {}
        unknown: List[str] = []
        platform_index = _platform_index(type(self))
        for position, header in enumerate(headers):
            if header is None or not str(header).strip():
                continue
            header_key = _key(header)
            field = platform_index.get(header_key) or _ALIAS_INDEX.get(header_key)
            if not field:
                unknown.append(str(header).strip())
                continue
            mapping.setdefault(field, position)
        return mapping, unknown

    def detect_header_row(self, sheet: SheetSource) -> Optional[Dict[str, Any]]:
        """Locate the header row within the first ``MAX_HEADER_SCAN_ROWS`` rows."""
        best: Optional[Dict[str, Any]] = None
        for row_number, row in enumerate(sheet.rows(), start=1):
            if row_number > self.MAX_HEADER_SCAN_ROWS:
                break
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            headers = [str(value).strip() if value is not None else '' for value in row]
            mapping, unknown = self.resolve_headers(headers)
            info = {
                'row_number': row_number,
                'headers': headers,
                'mapping': mapping,
                'unknown_headers': unknown,
                'score': len(mapping),
            }
            if best is None or info['score'] > best['score']:
                best = info
            if row_number == 1 and info['score'] >= self.MIN_MAPPED_COLUMNS:
                break
        if best is None or best['score'] < self.MIN_MAPPED_COLUMNS:
            return None
        return best

    def sheet_plan(self, sheets: List[SheetSource]):
        """Return (usable, ignored, observed_unknown_headers)."""
        preferred = {_key(name) for name in self.SHEET_NAMES}
        ignored_names = {_key(name) for name in self.NON_TRANSACTION_SHEETS}
        ordered = sorted(sheets, key=lambda sheet: 0 if _key(sheet.name) in preferred else 1)
        usable, ignored, observed = [], [], []
        for sheet in ordered:
            if _key(sheet.name) in ignored_names:
                ignored.append(sheet.name)
                continue
            header = self.detect_header_row(sheet)
            if header is None:
                ignored.append(sheet.name)
                continue
            usable.append((sheet, header))
            observed.extend(header['unknown_headers'])
        if not self.MULTI_SHEET and usable:
            usable = usable[:1]
        return usable, ignored, observed

    def _iter_data_rows(self, sheet: SheetSource, header_info: Dict[str, Any]):
        for row_number, row in enumerate(sheet.rows(), start=1):
            if row_number <= header_info['row_number']:
                continue
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            yield row_number, row

    def sheet_required_columns(self, sheet_name: str) -> Tuple[str, ...]:
        if self.MULTI_SHEET:
            key = _key(sheet_name)
            if key in {_key(n) for n in self.CONSOLIDATED_SHEETS}:
                return ('taxable_value',)
            if key in {_key(n) for n in self.NOTE_SHEETS}:
                return ('note_number',)
        return self.REQUIRED_COLUMNS

    def _expected_columns_hint(self, limit: int = 12) -> str:
        names: List[str] = list(self.HEADER_MAP.keys())
        for aliases in GENERIC_ALIASES.values():
            names.extend(aliases[:2])
        unique = []
        for name in names:
            if name not in unique:
                unique.append(name)
        return ', '.join(unique[:limit]) + (' ...' if len(unique) > limit else '')

    # ------------------------------------------------------------------
    # detection
    # ------------------------------------------------------------------
    def detect(self, workbook_or_data, file_name: str = '') -> bool:
        name = (file_name or '').lower()
        if any(token in name for token in self.FILENAME_TOKENS):
            return True
        if not self.SHEET_NAMES:
            return False
        sheets, _name, errors, _diagnostics, _wb_ref = self._load_sheets(workbook_or_data, file_name)
        if errors:
            return False
        preferred = {_key(sheet_name) for sheet_name in self.SHEET_NAMES}
        return any(_key(sheet.name) in preferred for sheet in sheets)

    # ------------------------------------------------------------------
    # validation
    # ------------------------------------------------------------------
    def validate(self, workbook_or_data) -> Tuple[bool, List[str]]:
        sheets, _file_name, load_errors, _diagnostics, _wb_ref = self._load_sheets(workbook_or_data)
        if load_errors:
            return False, list(load_errors)
        if not sheets:
            return False, ['Workbook contains no worksheets.']

        usable, _ignored, _observed = self.sheet_plan(sheets)
        if not usable:
            return False, [
                f'Unrecognized {self.PLATFORM_NAME} file: no worksheet has a recognizable '
                f'header row (scanned the first {self.MAX_HEADER_SCAN_ROWS} rows of '
                f'{len(sheets)} sheet(s): {", ".join(sheet.name for sheet in sheets)}). '
                f'Recognized column names include: {self._expected_columns_hint()}.'
            ]

        errors: List[str] = []
        if self.REQUIRE_ANY_SHEET:
            accepted = [
                (sheet, header) for sheet, header in usable
                if all(field in header['mapping'] for field in self.sheet_required_columns(sheet.name))
            ]
            if not accepted:
                required = sorted({
                    field
                    for sheet, _header in usable
                    for field in self.sheet_required_columns(sheet.name)
                })
                errors.append(
                    f'No {self.PLATFORM_NAME} sheet exposes the required column(s): '
                    f'{", ".join(FIELD_LABELS.get(f, f) for f in required)}.'
                )
        else:
            sheet, header = usable[0]
            missing = [field for field in self.REQUIRED_COLUMNS if field not in header['mapping']]
            if missing:
                errors.append(
                    f"Sheet '{sheet.name}' is missing required column(s): "
                    f"{', '.join(FIELD_LABELS.get(f, f) for f in missing)}. "
                    f"Columns found: {', '.join(h for h in header['headers'] if h) or '(none)'}."
                )

        if not errors:
            has_rows = False
            for sheet, header in usable:
                for _row_number, _row in self._iter_data_rows(sheet, header):
                    has_rows = True
                    break
                if has_rows:
                    break
            if not has_rows:
                errors.append(
                    'No data rows found beneath the header row. The file has headers but no '
                    'transaction rows to import.'
                )

        return (not errors), errors

    # ------------------------------------------------------------------
    # parsing
    # ------------------------------------------------------------------
    def parse(self, workbook_or_data, file_name: str = '', stream: bool = False) -> ImportResult:
        sheets, resolved_name, load_errors, diagnostics, _wb_ref = self._load_sheets(workbook_or_data, file_name)
        result = ImportResult(platform=self.PLATFORM_NAME, file_name=resolved_name or '')
        result.sheet_count = len(sheets)
        result.metadata['platform'] = self.PLATFORM_NAME
        result.metadata['format_documented'] = self.FORMAT_DOCUMENTED
        result.metadata.update(diagnostics or {})
        result.metadata["_wb_ref"] = _wb_ref

        if load_errors:
            result.errors.extend(load_errors)
            return result

        usable, ignored, observed = self.sheet_plan(sheets)
        result.metadata['sheets_ignored'] = ignored
        if not usable:
            result.errors.append(
                f'Unrecognized {self.PLATFORM_NAME} file: no worksheet has a recognizable '
                f'header row (scanned the first {self.MAX_HEADER_SCAN_ROWS} rows of '
                f'{len(sheets)} sheet(s)). Recognized column names include: '
                f'{self._expected_columns_hint()}.'
            )
            return result

        header_rows: Dict[str, int] = {}
        sheet_columns: Dict[str, List[str]] = {}

        def _row_generator():
            for sheet, header in usable:
                header_rows[sheet.name] = header['row_number']
                sheet_columns[sheet.name] = sorted(header['mapping'])
                for row_number, row in self._iter_data_rows(sheet, header):
                    raw_data: Dict[str, Any] = {}
                    headers = header['headers']
                    for position in range(max(len(headers), len(row))):
                        if position < len(headers):
                            key = headers[position] or f'col_{position}'
                        else:
                            key = f'col_{position}'
                        raw_data[key] = row[position] if position < len(row) else None
                    if len(row) > len(headers):
                        raw_data['_extra_values'] = [row[i] for i in range(len(headers), len(row))]
                    row_object = self._build_import_row(raw_data, row_number, sheet.name)
                    yield row_object

        result.rows = _row_generator()

        if not stream:
            result.rows = list(result.rows)
            result.total_rows = len(result.rows)
            for row_object in result.rows:
                if row_object.status.name == 'ERROR':
                    result.error_rows += 1
                elif row_object.status.name == 'WARNING':
                    result.warning_rows += 1
                elif row_object.status.name == 'SKIPPED':
                    result.skipped_rows += 1
                else:
                    result.success_rows += 1

        result.metadata['header_rows'] = header_rows
        result.metadata['sheet_columns'] = sheet_columns
        result.metadata['sheets_parsed'] = list(header_rows)
        result.metadata['unrecognized_headers'] = sorted(set(observed))
        result.metadata['recognized_columns'] = sorted({
            field for columns in sheet_columns.values() for field in columns
        })
        if observed:
            result.warnings.append(
                f'{len(set(observed))} unrecognised column(s) recorded and ignored: '
                f'{", ".join(sorted(set(observed)))}'
            )

        self._finalize_result(result)
        return result

    def _finalize_result(self, result: ImportResult) -> None:
        if not isinstance(result.rows, list): return
        periods, gstins = Counter(), Counter()
        for row in result.rows:
            normalized = row.normalized_data or {}
            for field in ('invoice_date', 'note_date'):
                parsed = parse_date(normalized.get(field))
                if parsed:
                    periods[f'{parsed.month:02d}{parsed.year}'] += 1
            seller_gstin = str(normalized.get('seller_gstin') or '')
            if validate_gstin(seller_gstin)[0]:
                gstins[seller_gstin] += 1
        if periods:
            result.detected_period = periods.most_common(1)[0][0]
        if gstins:
            result.detected_gstin = gstins.most_common(1)[0][0]

    def _build_import_row(self, raw_data: Dict[str, Any], row_number: int, sheet_name: str) -> ImportRow:
        raw_data.setdefault('_sheet_name', sheet_name)
        raw_data.setdefault('_row_number', row_number)
        normalized = self.normalize(raw_data)
        warnings = list(normalized.pop('_warnings', []) or [])
        errors: List[str] = list(normalized.pop('_errors', []) or [])
        extra_values = raw_data.get('_extra_values')
        if extra_values:
            warnings.append(
                f'{len(extra_values)} value(s) beyond the header columns were kept as '
                f'col_N and are not mapped to any field'
            )
        for field in self.sheet_required_columns(sheet_name):
            if not normalized.get(field):
                errors.append(f'Missing {FIELD_LABELS.get(field, field)}')
        status = ImportRowStatus.ERROR if errors else (
            ImportRowStatus.WARNING if warnings else ImportRowStatus.SUCCESS
        )
        return ImportRow(
            row_number=row_number,
            status=status,
            raw_data=raw_data,
            normalized_data=normalized,
            errors=errors,
            warnings=warnings,
            sheet_name=sheet_name,
        )

    # ------------------------------------------------------------------
    # normalisation helpers
    # ------------------------------------------------------------------
    def _row_field_values(self, raw_row: Dict[str, Any]) -> Dict[str, Any]:
        """Resolve a raw row to canonical field -> value (platform map first)."""
        values: Dict[str, Any] = {}
        platform_index = _platform_index(type(self))
        for key, value in raw_row.items():
            if key is None or str(key).startswith('_'):
                continue
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            header_key = _key(key)
            field = platform_index.get(header_key) or _ALIAS_INDEX.get(header_key)
            if field:
                values.setdefault(field, value)
        return values

    def _value(self, values: Dict[str, Any], raw_row: Dict[str, Any], field: str) -> Any:
        if values:
            value = values.get(field)
            if value is not None and not (isinstance(value, str) and not value.strip()):
                return value
        for alias in GENERIC_ALIASES.get(field, []):
            if alias in raw_row:
                value = raw_row[alias]
                if value is not None and not (isinstance(value, str) and not value.strip()):
                    return value
        return None

    @staticmethod
    def _to_decimal(value: Any) -> Optional[Decimal]:
        if value is None:
            return None
        if isinstance(value, Decimal):
            return value
        if isinstance(value, bool):
            return None
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        text = str(value).strip().replace(',', '').replace('₹', '').replace('%', '')
        if text in ('', '-', 'NA', 'N/A', 'None'):
            return None
        text = text.strip('()') if text.startswith('(') and text.endswith(')') else text
        try:
            return Decimal(text)
        except (InvalidOperation, ValueError):
            return None

    def _decimal_field(self, values, raw_row, field, warnings=None, label=None) -> Decimal:
        value = self._value(values, raw_row, field)
        if value is None:
            return Decimal('0')
        amount = self._to_decimal(value)
        if amount is None:
            if warnings is not None:
                warnings.append(
                    f"{label or FIELD_LABELS.get(field, field)} value '{value}' is not numeric"
                )
            return Decimal('0')
        return amount

    def get_decimal(self, raw_row: Dict, keys: List[str], default: Decimal = Decimal('0'),
                    values: Optional[Dict[str, Any]] = None) -> Decimal:
        for key in keys:
            if values and values.get(key) is not None:
                parsed = self._to_decimal(values.get(key))
                if parsed is not None:
                    return parsed
            if key in raw_row and raw_row[key] is not None:
                parsed = self._to_decimal(raw_row[key])
                if parsed is not None:
                    return parsed
        return default

    def get_string(self, raw_row: Dict, keys: List[str], default: str = '',
                   values: Optional[Dict[str, Any]] = None) -> str:
        for key in keys:
            if values and values.get(key) is not None:
                return str(values.get(key)).strip()
            if key in raw_row and raw_row[key] is not None:
                return str(raw_row[key]).strip()
        return default

    # ------------------------------------------------------------------
    # field mapping
    # ------------------------------------------------------------------
    def map_invoice(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None,
                    warnings: Optional[List[str]] = None) -> Dict:
        reverse_charge_raw = self._value(values, raw_row, 'reverse_charge')
        reverse_charge = 'Y' if str(reverse_charge_raw or '').strip().upper() in (
            'Y', 'YES', 'TRUE', '1'
        ) else 'N'
        return {
            'order_id': str(self._value(values, raw_row, 'order_id') or ''),
            'order_item_id': str(
                self._value(values, raw_row, 'order_item_id')
                or self._value(values, raw_row, 'sub_order_id') or ''
            ),
            'sub_order_id': str(self._value(values, raw_row, 'sub_order_id') or ''),
            'invoice_number': str(self._value(values, raw_row, 'invoice_number') or ''),
            'invoice_date': self._value(values, raw_row, 'invoice_date'),
            'invoice_type': str(self._value(values, raw_row, 'invoice_type') or ''),
            'document_type': str(self._value(values, raw_row, 'document_type') or ''),
            'seller_gstin': str(self._value(values, raw_row, 'seller_gstin') or '').strip().upper(),
            'reverse_charge': reverse_charge,
            'ecommerce_gstin': str(self._value(values, raw_row, 'ecommerce_gstin') or '').strip().upper(),
        }

    def map_customer(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None,
                     warnings: Optional[List[str]] = None) -> Dict:
        return {
            'customer_name': str(self._value(values, raw_row, 'customer_name') or ''),
            'customer_gstin': str(self._value(values, raw_row, 'customer_gstin') or '').strip().upper(),
            'place_of_supply_raw': self._value(values, raw_row, 'place_of_supply'),
            'place_of_supply': self._value(values, raw_row, 'place_of_supply'),
            'ship_to_state': str(self._value(values, raw_row, 'ship_to_state') or ''),
            'buyer_state': str(self._value(values, raw_row, 'buyer_state') or ''),
        }

    def map_items(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None,
                  warnings: Optional[List[str]] = None) -> Dict:
        return {
            'hsn_sac': str(self._value(values, raw_row, 'hsn_sac') or '').strip(),
            'description': str(self._value(values, raw_row, 'description') or ''),
            'item_code': str(self._value(values, raw_row, 'item_code') or ''),
            'quantity': self._decimal_field(values, raw_row, 'quantity', None, 'Quantity'),
            'uqc': str(self._value(values, raw_row, 'uqc') or ''),
            'unit_price': self._decimal_field(values, raw_row, 'unit_price', None, 'Unit price'),
            'discount': self._decimal_field(values, raw_row, 'discount', None, 'Discount'),
            'supply_type': str(self._value(values, raw_row, 'supply_type') or '').strip(),
        }

    def map_taxes(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None,
                  warnings: Optional[List[str]] = None) -> Dict:
        return {
            'taxable_value': self._decimal_field(values, raw_row, 'taxable_value', warnings),
            'cgst_rate': self._decimal_field(values, raw_row, 'cgst_rate', warnings),
            'cgst_amount': self._decimal_field(values, raw_row, 'cgst_amount', warnings),
            'sgst_rate': self._decimal_field(values, raw_row, 'sgst_rate', warnings),
            'sgst_amount': self._decimal_field(values, raw_row, 'sgst_amount', warnings),
            'igst_rate': self._decimal_field(values, raw_row, 'igst_rate', warnings),
            'igst_amount': self._decimal_field(values, raw_row, 'igst_amount', warnings),
            'cess_rate': self._decimal_field(values, raw_row, 'cess_rate', warnings),
            'cess_amount': self._decimal_field(values, raw_row, 'cess_amount', warnings),
            'tax_rate': self._decimal_field(values, raw_row, 'tax_rate', warnings),
            'invoice_value': self._decimal_field(values, raw_row, 'invoice_value', warnings),
        }

    def map_returns(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None) -> Dict:
        status = str(self._value(values, raw_row, 'order_status') or '').strip().lower()
        supply_type = str(self._value(values, raw_row, 'supply_type') or '').strip().lower()
        document_type = str(self._value(values, raw_row, 'document_type') or '').strip().lower()
        text = ' '.join(part for part in (status, supply_type, document_type) if part)

        cancelled = 'cancel' in text
        returned = any(token in text for token in ('return', 'rto', 'refund', 'exchange'))
        return {
            'return_flag': bool(returned or cancelled),
            'cancellation_flag': bool(cancelled),
            'event_status_raw': ' '.join(
                part for part in (status, supply_type) if part
            ),
        }

    def map_notes(self, raw_row: Dict, values: Optional[Dict[str, Any]] = None) -> Dict:
        # A note may be expressed by a note/document column, or - as in the
        # marketplace fixtures - by the Supply Type text ('Credit Note').
        note_type_raw = str(
            self._value(values, raw_row, 'note_type')
            or self._value(values, raw_row, 'document_type')
            or ''
        ).strip()
        supply_type_text = str(self._value(values, raw_row, 'supply_type') or '').strip()
        lowered = ' '.join(part for part in (note_type_raw, supply_type_text) if part).lower()
        if 'credit' in lowered or lowered in ('cr', 'c'):
            note_type = 'CREDIT'
        elif 'debit' in lowered or lowered in ('dr', 'd'):
            note_type = 'DEBIT'
        else:
            note_type = ''
        if note_type and not note_type_raw:
            note_type_raw = supply_type_text
        return {
            'note_type': note_type,
            'note_type_raw': note_type_raw,
            'note_number': str(self._value(values, raw_row, 'note_number') or ''),
            'note_date': self._value(values, raw_row, 'note_date'),
            'original_invoice_number': str(self._value(values, raw_row, 'original_invoice_number') or ''),
            'original_invoice_date': self._value(values, raw_row, 'original_invoice_date'),
        }

    # ------------------------------------------------------------------
    # normalise
    # ------------------------------------------------------------------
    def normalize(self, raw_row: Dict) -> Dict:
        values = self._row_field_values(raw_row)
        warnings: List[str] = []

        normalized: Dict[str, Any] = {}
        normalized.update(self.map_invoice(raw_row, values, warnings))
        normalized.update(self.map_customer(raw_row, values, warnings))
        normalized.update(self.map_items(raw_row, values, warnings))
        normalized.update(self.map_taxes(raw_row, values, warnings))
        normalized.update(self.map_returns(raw_row, values))
        normalized.update(self.map_notes(raw_row, values))
        normalized['marketplace_name'] = (
            str(self._value(values, raw_row, 'marketplace_name') or '').strip()
            or self.PLATFORM_NAME
        )
        normalized['source_platform'] = self.PLATFORM_NAME
        normalized['_errors'] = []

        self.post_normalize(raw_row, normalized, values, warnings)

        existing_meta = normalized.get('source_metadata') or {}
        normalized['source_metadata'] = {
            'platform': self.PLATFORM_NAME,
            'marketplace': normalized.get('marketplace_name', ''),
            'columns': [str(key) for key in raw_row.keys() if not str(key).startswith('_')],
        }
        if isinstance(existing_meta, dict):
            normalized['source_metadata'].update(existing_meta)
        normalized['_warnings'] = warnings
        return normalized

    def post_normalize(self, raw_row: Dict, normalized: Dict[str, Any],
                       values: Dict[str, Any], warnings: List[str]) -> None:
        """Canonicalise dates, place of supply, GSTINs and tax components."""
        errors: List[str] = normalized.setdefault('_errors', [])

        # --- dates -> DD-MM-YYYY (canonical contract) ---
        for field in ('invoice_date', 'note_date', 'original_invoice_date'):
            raw_value = normalized.get(field)
            if raw_value in (None, ''):
                continue
            parsed = parse_date(raw_value)
            if parsed is None:
                # A date that is present but unreadable makes the row unusable:
                # it cannot be assigned to a return period.
                errors.append(
                    f"{FIELD_LABELS.get(field, field)} '{raw_value}' is not a recognized date"
                )
                continue
            normalized[field] = format_date_gst(parsed)
        if not normalized.get('invoice_date') and not normalized.get('note_date'):
            warnings.append('Source row carries no document date column')

        # --- place of supply -> 2-digit state code ---
        customer_gstin = str(normalized.get('customer_gstin') or '')
        pos_raw = ''
        for field in ('place_of_supply', 'ship_to_state', 'buyer_state'):
            value = normalized.get(field)
            if value not in (None, ''):
                pos_raw = str(value).strip()
                break
        if pos_raw:
            # The explicit POS column wins; the customer GSTIN is only used as a
            # fallback when its state prefix is itself a valid state code (a
            # malformed GSTIN must not override a good POS column).
            gstin_for_pos = customer_gstin if (
                len(customer_gstin) >= 2 and customer_gstin[:2] in STATE_CODES
            ) else None
            code = resolve_pos_code(pos_raw, gstin_for_pos, fallback_code='')
            if code:
                normalized['place_of_supply'] = code
                normalized['place_of_supply_raw'] = pos_raw
                if code not in STATE_CODES:
                    errors.append(
                        f"Place of supply '{pos_raw}' resolved to '{code}', which is not a "
                        f"valid GST state code"
                    )
            else:
                normalized['place_of_supply'] = pos_raw
                warnings.append(
                    f"Place of supply '{pos_raw}' could not be resolved to a state code"
                )
        elif len(customer_gstin) >= 2 and customer_gstin[:2].isdigit():
            normalized['place_of_supply'] = customer_gstin[:2]
            warnings.append('Place of supply derived from the customer GSTIN')
        else:
            normalized['place_of_supply'] = ''
            warnings.append('No place of supply column found for this row')

        # --- GSTINs: format-checked but never discarded ---
        for field, label in (('customer_gstin', 'Customer GSTIN'),
                             ('ecommerce_gstin', 'E-commerce operator GSTIN'),
                             ('seller_gstin', 'Seller GSTIN')):
            gstin = str(normalized.get(field) or '').strip().upper()
            normalized[field] = gstin
            if not gstin:
                continue
            if field == 'ecommerce_gstin':
                continue  # operator GSTIN is informational for the import contract
            valid, message = validate_gstin(gstin)
            if not valid:
                warnings.append(f"{label} '{gstin}' rejected by validation ({message}) but kept")

        # --- tax components ---
        component_rates = (
            normalized.get('cgst_rate') or Decimal('0')
        ) + (normalized.get('sgst_rate') or Decimal('0')) + (normalized.get('igst_rate') or Decimal('0'))
        if not normalized.get('tax_rate') and component_rates:
            normalized['tax_rate'] = component_rates
            normalized['tax_rate_derived'] = True
        if not normalized.get('taxable_value') and not normalized.get('invoice_value'):
            warnings.append('Row carries neither a taxable value nor an invoice value column')
        elif not normalized.get('taxable_value'):
            warnings.append('Row has an invoice value but no taxable value column')
        elif (component_rates or normalized.get('tax_rate')) and not self._any_present(
            values, raw_row,
            ('cgst_amount', 'sgst_amount', 'igst_amount', 'cess_amount'),
        ):
            warnings.append(
                'Tax amounts are not present in the source row; only the rate is available'
            )

        # --- impossible values ---
        for field, label in (
            ('taxable_value', 'Taxable value'), ('invoice_value', 'Invoice value'),
            ('cgst_amount', 'CGST amount'), ('sgst_amount', 'SGST amount'),
            ('igst_amount', 'IGST amount'), ('cess_amount', 'Cess amount'),
            ('quantity', 'Quantity'), ('tax_rate', 'Tax rate'),
            ('cgst_rate', 'CGST rate'), ('sgst_rate', 'SGST rate'),
            ('igst_rate', 'IGST rate'), ('cess_rate', 'Cess rate'),
        ):
            value = normalized.get(field)
            if isinstance(value, Decimal) and value < 0:
                errors.append(
                    f'{label} {value} is negative; negative values are not supported by this '
                    f'import contract'
                )

    def _any_present(self, values: Dict[str, Any], raw_row: Dict[str, Any],
                     fields: Tuple[str, ...]) -> bool:
        """True if any of the canonical fields has a column in this row.

        Column *presence* is checked, not truthiness: an explicit 0.00 in a tax
        amount column is data, and must not be reported as a missing column.
        """
        for field in fields:
            if values and values.get(field) is not None:
                return True
            for alias in GENERIC_ALIASES.get(field, []):
                if alias in raw_row and raw_row[alias] is not None and str(raw_row[alias]).strip() != '':
                    return True
        return False


# ---------------------------------------------------------------------------
# Documented marketplace adapters
# ---------------------------------------------------------------------------

class AmazonAdapter(BaseGenericAdapter):
    """Amazon MTR (Marketplace Tax Report) export."""

    PLATFORM_NAME = 'Amazon'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Amazon Seller Central > Reports > Tax Document Library > MTR '
                    '(Marketplace Tax Report) export. Sheet is named "MTR".')
    FORMAT_DOCUMENTED = True
    FILENAME_TOKENS = ('amazon', 'mtr')
    SHEET_NAMES = ('MTR',)
    HEADER_MAP = {
        'Order ID': 'order_id',
        'Invoice Number': 'invoice_number',
        'Invoice Date': 'invoice_date',
        'Ship From State': 'ship_from_state',
        'Ship To State': 'ship_to_state',
        'Buyer GSTIN': 'customer_gstin',
        'Seller GSTIN': 'seller_gstin',
        'HSN/SAC': 'hsn_sac',
        'Product Description': 'description',
        'Quantity': 'quantity',
        'Taxable Value': 'taxable_value',
        'CGST Rate': 'cgst_rate',
        'CGST Amount': 'cgst_amount',
        'SGST Rate': 'sgst_rate',
        'SGST Amount': 'sgst_amount',
        'IGST Rate': 'igst_rate',
        'IGST Amount': 'igst_amount',
        'Cess Rate': 'cess_rate',
        'Cess Amount': 'cess_amount',
        'Invoice Value': 'invoice_value',
        'Tax Rate': 'tax_rate',
        'Supply Type': 'supply_type',
        'Place of Supply': 'place_of_supply',
        'Reverse Charge': 'reverse_charge',
        'E-Commerce GSTIN': 'ecommerce_gstin',
        'Ecommerce GSTIN': 'ecommerce_gstin',
        'Marketplace': 'marketplace_name',
    }


class FlipkartAdapter(BaseGenericAdapter):
    """Flipkart GST report (sheet "GST Report")."""

    PLATFORM_NAME = 'Flipkart'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Flipkart Seller Hub > Reports > GST report. Sheet is named '
                    '"GST Report"; line identifier column is "Order Item ID".')
    FORMAT_DOCUMENTED = True
    FILENAME_TOKENS = ('flipkart',)
    SHEET_NAMES = ('GST Report',)
    HEADER_MAP = {
        'Order Item ID': 'order_item_id',
        'Order ID': 'order_id',
        'Invoice Number': 'invoice_number',
        'Invoice Date': 'invoice_date',
        'Ship From State': 'ship_from_state',
        'Ship To State': 'ship_to_state',
        'Buyer GSTIN': 'customer_gstin',
        'Seller GSTIN': 'seller_gstin',
        'HSN/SAC': 'hsn_sac',
        'Product Title': 'description',
        'Selling Price': 'unit_price',
        'Quantity': 'quantity',
        'Taxable Value': 'taxable_value',
        'CGST Rate': 'cgst_rate',
        'CGST Amount': 'cgst_amount',
        'SGST Rate': 'sgst_rate',
        'SGST Amount': 'sgst_amount',
        'IGST Rate': 'igst_rate',
        'IGST Amount': 'igst_amount',
        'Cess Rate': 'cess_rate',
        'Cess Amount': 'cess_amount',
        'Invoice Value': 'invoice_value',
        'Tax Rate': 'tax_rate',
        'Supply Type': 'supply_type',
        'Place of Supply': 'place_of_supply',
        'Reverse Charge': 'reverse_charge',
        'E-Commerce GSTIN': 'ecommerce_gstin',
        'Marketplace': 'marketplace_name',
    }


class MeeshoAdapter(BaseGenericAdapter):
    """Meesho orders report (sheet "Orders")."""

    PLATFORM_NAME = 'Meesho'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Meesho Supplier Panel > Orders > Download orders report. Sheet is '
                    'named "Orders"; line identifier column is "Sub Order ID".')
    FORMAT_DOCUMENTED = True
    FILENAME_TOKENS = ('meesho',)
    SHEET_NAMES = ('Orders',)
    HEADER_MAP = {
        'Order ID': 'order_id',
        'Sub Order ID': 'sub_order_id',
        'Invoice Number': 'invoice_number',
        'Invoice Date': 'invoice_date',
        'Buyer GSTIN': 'customer_gstin',
        'Seller GSTIN': 'seller_gstin',
        'Ship From State': 'ship_from_state',
        'HSN': 'hsn_sac',
        'HSN Code': 'hsn_sac',
        'Product Name': 'description',
        'Selling Price': 'unit_price',
        'Quantity': 'quantity',
        'Taxable Amount': 'taxable_value',
        'Total': 'invoice_value',
        'CGST': 'cgst_amount',
        'SGST': 'sgst_amount',
        'IGST': 'igst_amount',
        'Cess Amount': 'cess_amount',
        'CGST Rate': 'cgst_rate',
        'SGST Rate': 'sgst_rate',
        'IGST Rate': 'igst_rate',
        'Tax Rate': 'tax_rate',
        'Buyer State': 'place_of_supply',
        'Place of Supply': 'place_of_supply',
        'Supply Type': 'supply_type',
        'Reverse Charge': 'reverse_charge',
        'E-Commerce GSTIN': 'ecommerce_gstin',
        'Marketplace': 'marketplace_name',
    }


class MyntraAdapter(FlipkartAdapter):
    """Myntra partner portal export.

    ``PLATFORM_ADAPTER_GUIDE.md`` records Myntra as using the same report layout
    as Flipkart, so the Flipkart column mapping is reused verbatim. No separate
    Myntra-only column set is documented in this project.
    """

    PLATFORM_NAME = 'Myntra'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Myntra Partner Portal > Reports > GST report (same column layout '
                    'as the Flipkart GST report).')
    FORMAT_DOCUMENTED = True
    FILENAME_TOKENS = ('myntra',)
    SHEET_NAMES = ('Myntra GST Report',)


class CustomExcelAdapter(BaseGenericAdapter):
    """Fallback adapter: flexible header detection over the generic aliases."""

    PLATFORM_NAME = 'Custom_Excel'
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('Any Excel/CSV export whose headers match the canonical field '
                    'aliases. Unrecognised columns are reported as warnings.')
    FORMAT_DOCUMENTED = True
    REQUIRED_COLUMNS = ('invoice_number',)
    CATCH_ALL_DETECT = True

    def detect(self, workbook_or_data, file_name: str = '') -> bool:
        # Fallback adapter: used when no platform can be detected.
        return True


class GSTR1GovtAdapter(BaseGenericAdapter):
    """GSTR-1 workbook produced by the GST portal offline tool / this project.

    The sheet names and sheet headers are the ones defined in
    ``app/services/gstr1_excel_writer.py``. Sheets that hold consolidated rows
    (b2cs/nil/hsn/hsnb2c) legitimately have no invoice number and are imported
    as WARNING rows; the document-count sheet (docs) holds no transactions and
    is skipped and reported in the result metadata.
    """

    PLATFORM_NAME = 'GSTR1_Govt'
    SUPPORTED_FILE_TYPES = ['.xlsx']
    INSTRUCTIONS = ('GSTR-1 workbook with the GST portal sheet names: b2b, b2ba, b2cl, '
                    'b2cs, cdnr, cdnur, exp, nil, hsn, hsnb2c, docs.')
    FORMAT_DOCUMENTED = True
    FILENAME_TOKENS = ('gstr1', 'gstr-1', 'gstr_1')
    SHEET_NAMES = ('b2b', 'b2ba', 'b2cl', 'b2cla', 'b2cs', 'b2csa', 'cdnr', 'cdnra',
                   'cdnur', 'cdnura', 'exp', 'expa', 'nil', 'hsn', 'hsnb2c')
    NON_TRANSACTION_SHEETS = ('docs',)
    CONSOLIDATED_SHEETS = ('b2cs', 'b2csa', 'nil', 'hsn', 'hsnb2c')
    AMENDMENT_SHEETS = ('b2ba', 'b2cla', 'b2csa', 'cdnra', 'cdnura', 'expa')
    NOTE_SHEETS = ('cdnr', 'cdnra', 'cdnur', 'cdnura')
    MULTI_SHEET = True
    REQUIRE_ANY_SHEET = True
    HEADER_MAP = {
        'GSTIN/UIN of Recipient': 'customer_gstin',
        'Invoice Number': 'invoice_number',
        'Invoice date': 'invoice_date',
        'Invoice Date': 'invoice_date',
        'Invoice Value': 'invoice_value',
        'Place Of Supply': 'place_of_supply',
        'Reverse Charge': 'reverse_charge',
        'Applicable % of Tax Rate': 'applicable_tax_rate_pct',
        'Invoice Type': 'invoice_type',
        'E-Commerce GSTIN': 'ecommerce_gstin',
        'Rate': 'tax_rate',
        'Taxable Value': 'taxable_value',
        'Cess Amount': 'cess_amount',
        'Original Invoice Number': 'original_invoice_number',
        'Original Invoice date': 'original_invoice_date',
        'Original Invoice Date': 'original_invoice_date',
        'Revised Invoice Number': 'invoice_number',
        'Revised Invoice date': 'invoice_date',
        'Revised Invoice Date': 'invoice_date',
        'Note/Refund Voucher Number': 'note_number',
        'Note/Refund Voucher date': 'note_date',
        'Note/Refund Voucher Date': 'note_date',
        'Note/Refund Voucher Value': 'invoice_value',
        'Invoice/Advance Receipt Number': 'original_invoice_number',
        'Invoice/Advance Receipt date': 'original_invoice_date',
        'Invoice/Advance Receipt Date': 'original_invoice_date',
        'Note Supply Type': 'supply_type',
        'Note Type': 'note_type',
        'Original Note Number': 'original_invoice_number',
        'Original Note Date': 'original_invoice_date',
        'Revised Note Number': 'note_number',
        'Revised Note Date': 'note_date',
        'Note Value': 'invoice_value',
        'Export Type': 'export_type',
        'Port Code': 'port_code',
        'Shipping Bill Number': 'shipping_bill_number',
        'Shipping Bill Date': 'shipping_bill_date',
        'Type': 'b2cs_type',
        'Description': 'description',
        'Nil Rated Supplies': 'nil_rated_supplies',
        'Exempted (other than nil rated/non GST supply)': 'exempt_supplies',
        'Non-GST supplies': 'non_gst_supplies',
        'HSN': 'hsn_sac',
        'UQC': 'uqc',
        'Total Quantity': 'quantity',
        'Total Value': 'invoice_value',
        'Integrated Tax Amount': 'igst_amount',
        'Central Tax Amount': 'cgst_amount',
        'State/UT Tax Amount': 'sgst_amount',
    }

    def _bucket_value(self, values, raw_row, field, header):
        if values and values.get(field) is not None:
            return values[field]
        return raw_row.get(header)

    def post_normalize(self, raw_row, normalized, values, warnings):
        # CDNR/CDNUR rows are identified by their note/refund voucher number;
        # carry it as the document number so the row stays importable.
        if not normalized.get('invoice_number') and normalized.get('note_number'):
            normalized['invoice_number'] = str(normalized['note_number'])
            warnings.append('Document number taken from the note/refund voucher number column')
        elif not normalized.get('invoice_number'):
            # b2cs / nil / hsn sheets hold aggregate rows with no document
            # identity by design; say so instead of presenting them as invoices.
            warnings.append(
                'Consolidated GSTR-1 row: no document number in this sheet layout (aggregate row)'
            )

        # Consolidated nil/exempt/non-GST rows carry three separate buckets.
        if not normalized.get('taxable_value'):
            buckets = [
                ('NIL', self._bucket_value(values, raw_row, 'nil_rated_supplies', 'Nil Rated Supplies')),
                ('EXEMPT', self._bucket_value(values, raw_row, 'exempt_supplies',
                                              'Exempted (other than nil rated/non GST supply)')),
                ('NONGST', self._bucket_value(values, raw_row, 'non_gst_supplies', 'Non-GST supplies')),
            ]
            present = [
                (item_type, self._to_decimal(amount) or Decimal('0'))
                for item_type, amount in buckets
                if amount not in (None, '')
            ]
            non_zero = [(item_type, amount) for item_type, amount in present if amount != 0]
            if non_zero:
                item_type, amount = max(non_zero, key=lambda item: item[1])
                normalized['taxable_value'] = amount
                normalized['item_type'] = item_type
                warnings.append(
                    f'Consolidated GSTR-1 row: {item_type} bucket used as the taxable value'
                )
            elif present:
                normalized['taxable_value'] = Decimal('0')
        super().post_normalize(raw_row, normalized, values, warnings)


# ---------------------------------------------------------------------------
# Filename-detected adapters with no documented column evidence
# ---------------------------------------------------------------------------

class _GenericFilenameAdapter(BaseGenericAdapter):
    """Base for platforms whose export layout is not documented in this project.

    The platform is still selectable (and still files under its own name), but
    the columns are matched with the generic alias table because no
    platform-specific format is documented.
    """

    FORMAT_DOCUMENTED = False
    SUPPORTED_FILE_TYPES = ['.xlsx', '.csv']
    INSTRUCTIONS = ('No platform-specific column format is documented for this '
                    'marketplace; generic header matching is applied. Export the report '
                    'with invoice number, date, taxable value and tax columns.')


class JioMartAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'JioMart'
    FILENAME_TOKENS = ('jiomart', 'jio mart')


class SnapdealAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Snapdeal'
    FILENAME_TOKENS = ('snapdeal',)


class TataCliqAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'TataCLiQ'
    FILENAME_TOKENS = ('tatacliq', 'tata_cliq', 'tata cliq')


class LimeroadAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Limeroad'
    FILENAME_TOKENS = ('limeroad',)


class ShopdeckAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Shopdeck'
    FILENAME_TOKENS = ('shopdeck',)


class GlowRoadAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'GlowRoad'
    FILENAME_TOKENS = ('glowroad',)


class SnapmintAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Snapmint'
    FILENAME_TOKENS = ('snapmint',)


class CitymallAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Citymall'
    FILENAME_TOKENS = ('citymall',)


class RoposoAdapter(_GenericFilenameAdapter):
    PLATFORM_NAME = 'Roposo'
    FILENAME_TOKENS = ('roposo', 'clout')
