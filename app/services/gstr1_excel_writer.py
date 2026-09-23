import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from decimal import Decimal
from typing import Dict, Any, List

class GSTR1ExcelWriter:
    SHEET_CONFIGS = {
        "b2b": {
            "title": "b2b",
            "headers": ["GSTIN/UIN of Recipient", "Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value", "Cess Amount"]
        },
        "b2ba": {
            "title": "b2ba",
            "headers": ["GSTIN/UIN of Recipient", "Original Invoice Number", "Original Invoice date", "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value", "Cess Amount"]
        },
        "b2cl": {
            "title": "b2cl",
            "headers": ["Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"]
        },
        "b2cla": {
            "title": "b2cla",
            "headers": ["Original Invoice Number", "Original Invoice date", "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"]
        },
        "b2cs": {
            "title": "b2cs",
            "headers": ["Type", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"]
        },
        "b2csa": {
            "title": "b2csa",
            "headers": ["Type", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"]
        },
        "cdnr": {
            "title": "cdnr",
            "headers": ["GSTIN/UIN of Recipient", "Note/Refund Voucher Number", "Note/Refund Voucher date", "Invoice/Advance Receipt Number", "Invoice/Advance Receipt date", "Note/Refund Voucher Value", "Place Of Supply", "Reverse Charge", "Note Supply Type", "Note Type", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"]
        },
        "cdnra": {
            "title": "cdnra",
            "headers": ["GSTIN/UIN of Recipient", "Original Note Number", "Original Note Date", "Revised Note Number", "Revised Note Date", "Note Value", "Place Of Supply", "Note Supply Type", "Note Type", "Rate", "Taxable Value", "Cess Amount"]
        },
        "cdnur": {
            "title": "cdnur",
            "headers": ["Note/Refund Voucher Number", "Note/Refund Voucher date", "Invoice/Advance Receipt Number", "Invoice/Advance Receipt date", "Note/Refund Voucher Value", "Place Of Supply", "Note Supply Type", "Note Type", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"]
        },
        "cdnura": {
            "title": "cdnura",
            "headers": ["Original Note Number", "Original Note Date", "Revised Note Number", "Revised Note Date", "Note Value", "Place Of Supply", "Note Supply Type", "Note Type", "Rate", "Taxable Value", "Cess Amount"]
        },
        "exp": {
            "title": "exp",
            "headers": ["Export Type", "Invoice Number", "Invoice date", "Invoice Value", "Port Code", "Shipping Bill Number", "Shipping Bill Date", "Rate", "Taxable Value"]
        },
        "expa": {
            "title": "expa",
            "headers": ["Export Type", "Original Invoice Number", "Original Invoice Date", "Revised Invoice Number", "Revised Invoice Date", "Invoice Value", "Port Code", "Shipping Bill Number", "Shipping Bill Date", "Rate", "Taxable Value"]
        },
        "nil": {
            "title": "nil",
            "headers": ["Supply Type", "Nil Rated Supplies", "Exempted (other than nil rated/non GST supply)", "Non-GST supplies"]
        },
        "hsn": {
            "title": "hsn",
            "headers": ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Taxable Value", "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount", "Rate"]
        },
        "hsnb2c": {
            "title": "hsnb2c",
            "headers": ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Taxable Value", "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount", "Rate"]
        },
        "docs": {
            "title": "docs",
            "headers": ["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled"]
        }
    }

    @classmethod
    def generate_excel(cls, data: Dict[str, List[List[Any]]], output_path: str):
        wb = openpyxl.Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        font_bold = Font(bold=True)
        align_center = Alignment(horizontal="center", vertical="center")

        for sheet_key, config in cls.SHEET_CONFIGS.items():
            if sheet_key in data and data[sheet_key]:
                ws = wb.create_sheet(title=config["title"])
                headers = config["headers"]
                ws.append(headers)

                # Format headers
                for col_idx, cell in enumerate(ws[1], 1):
                    cell.font = font_bold
                    cell.alignment = align_center
                    ws.column_dimensions[get_column_letter(col_idx)].width = max(15, len(headers[col_idx-1]) + 2)

                ws.freeze_panes = "A2"

                for row_data in data[sheet_key]:
                    formatted_row = []
                    for item in row_data:
                        if isinstance(item, Decimal):
                            formatted_row.append(float(round(item, 2)))
                        else:
                            formatted_row.append(item)
                    ws.append(formatted_row)

                # Format amounts
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        if isinstance(cell.value, float):
                            cell.number_format = '0.00'

        if not wb.sheetnames:
            wb.create_sheet("Empty")

        wb.save(output_path)
        return output_path

def generate_gstr1_excel(data: Dict[str, List[List[Any]]], output_path: str) -> str:
    """Convenience function to generate GSTR-1 Excel."""
    return GSTR1ExcelWriter.generate_excel(data, output_path)
