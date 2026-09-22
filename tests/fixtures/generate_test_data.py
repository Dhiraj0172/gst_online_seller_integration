"""
Comprehensive test fixtures and synthetic test data generator for GST Online Seller.

Generates realistic test Excel files with valid synthetic data for all test scenarios.
Run this file directly to generate all fixture files:
    python tests/fixtures/generate_test_data.py
"""
import os
import sys
from decimal import Decimal
from datetime import date, timedelta
import random

# Add project root to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# --- Constants ---
FIXTURE_DIR = os.path.dirname(os.path.abspath(__file__))

SELLER_GSTIN = '27AABCU9603R1ZM'
SELLER_STATE = '27'

BUYER_GSTINS = [
    '29AALCS5765L1ZP',  # Karnataka
    '07AADCB2230M1ZT',  # Delhi
    '33AAACR5055K1ZK',  # Tamil Nadu
    '09AABCT1332L1ZL',  # UP
    '24AABCU9603R1ZP',  # Gujarat
]

ECOM_GSTINS = {
    'Amazon': '27AABCA1234B1ZM',
    'Flipkart': '29AABCF5678D1ZP',
}

HSN_CODES = ['6109', '6205', '8517', '3304', '8471', '6110', '9403', '8528']
TAX_RATES = [5, 12, 18, 28]
STATES = {
    '27': 'Maharashtra', '29': 'Karnataka', '07': 'Delhi',
    '33': 'Tamil Nadu', '09': 'Uttar Pradesh', '24': 'Gujarat',
    '36': 'Telangana', '32': 'Kerala',
}
UQC_CODES = ['NOS', 'PCS', 'KGS', 'MTR', 'SET', 'DOZ']

DESCRIPTIONS = [
    'T-Shirt Cotton Round Neck', 'Mobile Phone Case', 'Laptop Bag Premium',
    'Face Cream Organic', 'Wireless Earbuds', 'Kurta Silk Men',
    'Wooden Shelf Unit', 'LED Television 43 inch',
]

BASE_DATE = date(2025, 1, 1)
RETURN_PERIOD = '012025'


def _rand_date(month=1, year=2025):
    """Random date within the month."""
    day = random.randint(1, 28)
    return date(year, month, day)


def _fmt_date(d):
    return d.strftime('%d-%m-%Y')


def _inv_num(prefix, idx):
    return f'{prefix}-2025-{idx:04d}'


def _make_b2b_row(idx, hsn=None, rate=None, buyer_gstin=None):
    """Create a B2B invoice row dict for Amazon-style format."""
    buyer = buyer_gstin or random.choice(BUYER_GSTINS)
    buyer_state = buyer[:2]
    is_intra = (buyer_state == SELLER_STATE)
    r = rate or random.choice(TAX_RATES)
    h = hsn or random.choice(HSN_CODES)
    qty = random.randint(1, 20)
    price = round(random.uniform(200, 15000), 2)
    taxable = round(price * qty, 2)
    if is_intra:
        cgst_rate = r / 2
        sgst_rate = r / 2
        igst_rate = 0
    else:
        cgst_rate = 0
        sgst_rate = 0
        igst_rate = r
    cgst = round(taxable * cgst_rate / 100, 2)
    sgst = round(taxable * sgst_rate / 100, 2)
    igst = round(taxable * igst_rate / 100, 2)
    total_tax = cgst + sgst + igst
    inv_value = round(taxable + total_tax, 2)
    return {
        'Order ID': f'ORD-{idx:06d}',
        'Invoice Number': _inv_num('INV', idx),
        'Invoice Date': _fmt_date(_rand_date()),
        'Ship From State': STATES.get(SELLER_STATE, 'Maharashtra'),
        'Ship To State': STATES.get(buyer_state, 'Other'),
        'Buyer GSTIN': buyer,
        'Seller GSTIN': SELLER_GSTIN,
        'HSN/SAC': h,
        'Product Description': random.choice(DESCRIPTIONS),
        'Quantity': qty,
        'Taxable Value': taxable,
        'CGST Rate': cgst_rate,
        'CGST Amount': cgst,
        'SGST Rate': sgst_rate,
        'SGST Amount': sgst,
        'IGST Rate': igst_rate,
        'IGST Amount': igst,
        'Cess Rate': 0,
        'Cess Amount': 0,
        'Invoice Value': inv_value,
        'Tax Rate': r,
        'Supply Type': 'Regular',
        'Place of Supply': f'{buyer_state}-{STATES.get(buyer_state, "Other")}',
        'Reverse Charge': 'N',
    }


def _make_b2c_row(idx, inter=False, large=False):
    """Create a B2C invoice row."""
    if inter:
        buyer_state = random.choice([k for k in STATES if k != SELLER_STATE])
    else:
        buyer_state = SELLER_STATE
    is_intra = (buyer_state == SELLER_STATE)
    r = random.choice(TAX_RATES)
    h = random.choice(HSN_CODES)
    qty = random.randint(1, 10)
    if large:
        price = round(random.uniform(30000, 80000), 2)
    else:
        price = round(random.uniform(100, 5000), 2)
    taxable = round(price * qty, 2)
    if is_intra:
        cgst = round(taxable * r / 200, 2)
        sgst = cgst
        igst = 0
    else:
        cgst = 0
        sgst = 0
        igst = round(taxable * r / 100, 2)
    total_tax = cgst + sgst + igst
    inv_value = round(taxable + total_tax, 2)
    return {
        'Order ID': f'ORD-{idx:06d}',
        'Invoice Number': _inv_num('INV', idx),
        'Invoice Date': _fmt_date(_rand_date()),
        'Ship From State': STATES.get(SELLER_STATE, 'Maharashtra'),
        'Ship To State': STATES.get(buyer_state, 'Other'),
        'Buyer GSTIN': '',
        'Seller GSTIN': SELLER_GSTIN,
        'HSN/SAC': h,
        'Product Description': random.choice(DESCRIPTIONS),
        'Quantity': qty,
        'Taxable Value': taxable,
        'CGST Rate': r / 2 if is_intra else 0,
        'CGST Amount': cgst,
        'SGST Rate': r / 2 if is_intra else 0,
        'SGST Amount': sgst,
        'IGST Rate': r if not is_intra else 0,
        'IGST Amount': igst,
        'Cess Rate': 0, 'Cess Amount': 0,
        'Invoice Value': inv_value,
        'Tax Rate': r,
        'Supply Type': 'Regular',
        'Place of Supply': f'{buyer_state}-{STATES.get(buyer_state, "Other")}',
        'Reverse Charge': 'N',
    }


def _make_cn_row(idx, registered=True):
    """Create a credit note row."""
    buyer = random.choice(BUYER_GSTINS) if registered else ''
    buyer_state = buyer[:2] if buyer else random.choice(list(STATES.keys()))
    is_intra = (buyer_state == SELLER_STATE)
    r = random.choice(TAX_RATES)
    taxable = round(random.uniform(500, 5000), 2)
    if is_intra:
        cgst = round(taxable * r / 200, 2); sgst = cgst; igst = 0
    else:
        cgst = 0; sgst = 0; igst = round(taxable * r / 100, 2)
    return {
        'Order ID': f'ORD-{idx:06d}',
        'Invoice Number': _inv_num('CN', idx),
        'Invoice Date': _fmt_date(_rand_date()),
        'Ship From State': STATES.get(SELLER_STATE),
        'Ship To State': STATES.get(buyer_state, 'Other'),
        'Buyer GSTIN': buyer,
        'Seller GSTIN': SELLER_GSTIN,
        'HSN/SAC': random.choice(HSN_CODES),
        'Product Description': random.choice(DESCRIPTIONS),
        'Quantity': random.randint(1, 5),
        'Taxable Value': taxable,
        'CGST Rate': r / 2 if is_intra else 0,
        'CGST Amount': cgst,
        'SGST Rate': r / 2 if is_intra else 0,
        'SGST Amount': sgst,
        'IGST Rate': r if not is_intra else 0,
        'IGST Amount': igst,
        'Cess Rate': 0, 'Cess Amount': 0,
        'Invoice Value': round(taxable + cgst + sgst + igst, 2),
        'Tax Rate': r,
        'Supply Type': 'Credit Note',
        'Document Type': 'Credit Note',
        'Original Invoice Number': _inv_num('INV', max(1, idx - 50)),
        'Original Invoice Date': _fmt_date(_rand_date()),
        'Place of Supply': f'{buyer_state}-{STATES.get(buyer_state, "Other")}',
        'Reverse Charge': 'N',
    }


def _write_rows_to_sheet(ws, rows):
    """Write list of dicts to worksheet with header row."""
    if not rows:
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, '') for h in headers])


def generate_amazon_sample():
    """Generate Amazon MTR format sample."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'MTR'
    rows = []
    for i in range(1, 51):
        rows.append(_make_b2b_row(i))
    for i in range(51, 151):
        rows.append(_make_b2c_row(i, inter=(i % 3 == 0)))
    for i in range(151, 161):
        rows.append(_make_cn_row(i, registered=True))
    for i in range(161, 166):
        rows.append(_make_cn_row(i, registered=False))
    # 3 cancelled orders
    for i in range(166, 169):
        r = _make_b2c_row(i)
        r['Supply Type'] = 'Cancelled'
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')
    return path


def generate_flipkart_sample():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'GST Report'
    rows = []
    for i in range(1, 31):
        r = _make_b2b_row(i + 200)
        r['Order Item ID'] = f'OI-{i:06d}'
        r['Product Title'] = r.pop('Product Description')
        r['Selling Price'] = r['Taxable Value'] / max(r['Quantity'], 1)
        rows.append(r)
    for i in range(31, 111):
        r = _make_b2c_row(i + 200, inter=(i % 4 == 0))
        r['Order Item ID'] = f'OI-{i:06d}'
        r['Product Title'] = r.pop('Product Description')
        rows.append(r)
    for i in range(111, 116):
        r = _make_cn_row(i + 200, registered=True)
        r['Product Title'] = r.pop('Product Description')
        rows.append(r)
    for i in range(116, 119):
        r = _make_b2c_row(i + 200)
        r['Supply Type'] = 'Return'
        r['Product Title'] = r.pop('Product Description')
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_meesho_sample():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Orders'
    rows = []
    for i in range(1, 41):
        if i <= 10:
            r = _make_b2b_row(i + 400)
        else:
            r = _make_b2c_row(i + 400, inter=(i % 3 == 0))
        r['Sub Order ID'] = f'SO-{i:06d}'
        r['Product Name'] = r.pop('Product Description')
        r['HSN'] = r.pop('HSN/SAC')
        r['Selling Price'] = r['Taxable Value'] / max(r['Quantity'], 1)
        r['Taxable Amount'] = r.pop('Taxable Value')
        r['Total'] = r.pop('Invoice Value')
        r['CGST'] = r.pop('CGST Amount')
        r['SGST'] = r.pop('SGST Amount')
        r['IGST'] = r.pop('IGST Amount')
        r['Buyer State'] = r.pop('Ship To State')
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'meesho_sample.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_b2b_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'B2B'
    rows = []
    for i in range(1, 26):
        r = _make_b2b_row(i + 500, buyer_gstin=BUYER_GSTINS[i % len(BUYER_GSTINS)])
        rows.append(r)
    # Add reverse charge
    r = _make_b2b_row(526)
    r['Reverse Charge'] = 'Y'
    rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'b2b_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_b2c_test(name, inter=False, large=False, count=20):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'B2C'
    rows = [_make_b2c_row(i + 600, inter=inter, large=large) for i in range(count)]
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, f'{name}.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_credit_debit_notes(name, note_type='credit', count=10):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Notes'
    rows = []
    for i in range(count):
        r = _make_cn_row(i + 700, registered=True)
        if note_type == 'debit':
            r['Invoice Number'] = _inv_num('DN', i + 700)
            r['Document Type'] = 'Debit Note'
            r['Supply Type'] = 'Debit Note'
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, f'{name}.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_nil_exempt_nongst(name, item_type='nil'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    for i in range(10):
        r = _make_b2c_row(i + 800, inter=(i % 2 == 0))
        r['Tax Rate'] = 0
        r['CGST Rate'] = 0; r['CGST Amount'] = 0
        r['SGST Rate'] = 0; r['SGST Amount'] = 0
        r['IGST Rate'] = 0; r['IGST Amount'] = 0
        r['Invoice Value'] = r['Taxable Value']
        if item_type == 'exempt':
            r['Supply Type'] = 'Exempt'
        elif item_type == 'non_gst':
            r['Supply Type'] = 'Non-GST'
        else:
            r['Supply Type'] = 'Nil Rated'
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, f'{name}.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_hsn_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'HSN'
    rows = []
    for i, hsn in enumerate(HSN_CODES):
        for rate in [5, 18]:
            r = _make_b2b_row(i * 10 + rate + 900, hsn=hsn, rate=rate)
            rows.append(r)
            r2 = _make_b2c_row(i * 10 + rate + 950, inter=False)
            r2['HSN/SAC'] = hsn
            r2['Tax Rate'] = rate
            rows.append(r2)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'hsn_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_ecom_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ecom'
    rows = []
    for i in range(20):
        r = _make_b2c_row(i + 1000, inter=(i % 2 == 0))
        r['E-Commerce GSTIN'] = ECOM_GSTINS.get('Amazon' if i % 2 == 0 else 'Flipkart', '')
        r['Marketplace'] = 'Amazon' if i % 2 == 0 else 'Flipkart'
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'ecom_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_duplicate_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Duplicates'
    r1 = _make_b2b_row(1100)
    r2 = dict(r1)  # exact duplicate
    r3 = _make_b2b_row(1101)
    _write_rows_to_sheet(ws, [r1, r2, r3])
    path = os.path.join(FIXTURE_DIR, 'duplicate_test.xlsx')
    wb.save(path)
    print(f'  Created {path} (3 rows, 1 duplicate)')


def generate_malformed_gstin_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    bad_gstins = ['INVALID', '123', '27AABCU960', 'XXAABCU9603R1ZM', '00AABCU9603R1ZM']
    for i, gstin in enumerate(bad_gstins):
        r = _make_b2b_row(i + 1200)
        r['Buyer GSTIN'] = gstin
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'malformed_gstin_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_missing_hsn_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    for i in range(5):
        r = _make_b2b_row(i + 1300)
        r['HSN/SAC'] = ''
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'missing_hsn_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_invalid_dates_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    bad_dates = ['99-99-9999', 'not-a-date', '2025/13/01', '', '00-00-0000']
    for i, d in enumerate(bad_dates):
        r = _make_b2b_row(i + 1400)
        r['Invoice Date'] = d
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'invalid_dates_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_invalid_tax_rate_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    bad_rates = [15, 25, 99, -5, 100]
    for i, rate in enumerate(bad_rates):
        r = _make_b2b_row(i + 1500, rate=rate)
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'invalid_tax_rate_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_rounding_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    # Specific rounding edge cases
    test_cases = [
        (333.33, 18), (100.01, 5), (0.01, 28),
        (99999999.99, 18), (1.00, 18), (999.995, 12),
    ]
    rows = []
    for i, (taxable, rate) in enumerate(test_cases):
        r = _make_b2c_row(i + 1600)
        r['Quantity'] = 1
        r['Taxable Value'] = taxable
        r['Tax Rate'] = rate
        cgst = round(taxable * rate / 200, 2)
        sgst = cgst
        r['CGST Rate'] = rate / 2
        r['CGST Amount'] = cgst
        r['SGST Rate'] = rate / 2
        r['SGST Amount'] = sgst
        r['IGST Rate'] = 0; r['IGST Amount'] = 0
        r['Invoice Value'] = round(taxable + cgst + sgst, 2)
        rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'rounding_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_zero_tax_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = [_make_b2c_row(i + 1700) for i in range(5)]
    for r in rows:
        r['Tax Rate'] = 0
        r['CGST Rate'] = 0; r['CGST Amount'] = 0
        r['SGST Rate'] = 0; r['SGST Amount'] = 0
        r['IGST Rate'] = 0; r['IGST Amount'] = 0
        r['Invoice Value'] = r['Taxable Value']
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'zero_tax_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_mixed_test():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Mixed'
    rows = []
    for i in range(1, 11): rows.append(_make_b2b_row(i + 1800))
    for i in range(11, 21): rows.append(_make_b2c_row(i + 1800, inter=(i % 2 == 0)))
    for i in range(21, 26): rows.append(_make_cn_row(i + 1800, registered=True))
    for i in range(26, 29): rows.append(_make_cn_row(i + 1800, registered=False))
    r = _make_b2c_row(1830)
    r['Tax Rate'] = 0; r['CGST Rate'] = 0; r['CGST Amount'] = 0
    r['SGST Rate'] = 0; r['SGST Amount'] = 0; r['IGST Rate'] = 0; r['IGST Amount'] = 0
    r['Invoice Value'] = r['Taxable Value']
    rows.append(r)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'mixed_test.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_large_test(count=10000):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = []
    for i in range(1, count + 1):
        if i % 5 == 0:
            rows.append(_make_b2b_row(i + 2000))
        elif i % 7 == 0:
            rows.append(_make_cn_row(i + 2000))
        else:
            rows.append(_make_b2c_row(i + 2000, inter=(i % 3 == 0)))
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'large_test_10k.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_empty_file():
    wb = openpyxl.Workbook()
    path = os.path.join(FIXTURE_DIR, 'empty_file.xlsx')
    wb.save(path)
    print(f'  Created {path} (empty)')


def generate_wrong_headers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(['Col A', 'Col B', 'Col C', 'Col D'])
    ws.append([1, 2, 3, 4])
    path = os.path.join(FIXTURE_DIR, 'wrong_headers.xlsx')
    wb.save(path)
    print(f'  Created {path}')


def generate_extra_columns():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    rows = [_make_b2b_row(i + 3000) for i in range(5)]
    for r in rows:
        r['Extra Column 1'] = 'extra_data'
        r['Internal ID'] = random.randint(10000, 99999)
    _write_rows_to_sheet(ws, rows)
    path = os.path.join(FIXTURE_DIR, 'extra_columns.xlsx')
    wb.save(path)
    print(f'  Created {path} ({len(rows)} rows)')


def generate_tcs_portal_report():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'TCS Report'
    headers = ['State Code', 'State Name', 'Taxable Value', 'TCS Amount', 'CGST TCS', 'SGST TCS', 'IGST TCS']
    ws.append(headers)
    for code, name in [('27', 'Maharashtra'), ('29', 'Karnataka'), ('07', 'Delhi'), ('33', 'Tamil Nadu')]:
        taxable = round(random.uniform(50000, 500000), 2)
        tcs = round(taxable * 0.01, 2)
        if code == SELLER_STATE:
            cgst_tcs = round(tcs / 2, 2)
            sgst_tcs = cgst_tcs
            igst_tcs = 0
        else:
            cgst_tcs = 0
            sgst_tcs = 0
            igst_tcs = tcs
        ws.append([code, name, taxable, tcs, cgst_tcs, sgst_tcs, igst_tcs])
    path = os.path.join(FIXTURE_DIR, 'tcs_portal_report.xlsx')
    wb.save(path)
    print(f'  Created {path}')


def generate_all():
    """Generate all test fixture files."""
    random.seed(42)  # Deterministic
    print('Generating test fixtures...')
    os.makedirs(FIXTURE_DIR, exist_ok=True)

    generate_amazon_sample()
    generate_flipkart_sample()
    generate_meesho_sample()
    generate_b2b_test()
    generate_b2c_test('b2c_intra', inter=False)
    generate_b2c_test('b2c_inter', inter=True)
    generate_b2c_test('b2cl_test', inter=True, large=True, count=10)
    generate_credit_debit_notes('credit_note_test', 'credit')
    generate_credit_debit_notes('debit_note_test', 'debit')
    generate_nil_exempt_nongst('nil_rated_test', 'nil')
    generate_nil_exempt_nongst('exempt_test', 'exempt')
    generate_nil_exempt_nongst('non_gst_test', 'non_gst')
    generate_hsn_test()
    generate_ecom_test()
    generate_duplicate_test()
    generate_malformed_gstin_test()
    generate_missing_hsn_test()
    generate_invalid_dates_test()
    generate_invalid_tax_rate_test()
    generate_rounding_test()
    generate_zero_tax_test()
    generate_mixed_test()
    generate_large_test(10000)
    generate_empty_file()
    generate_wrong_headers()
    generate_extra_columns()
    generate_tcs_portal_report()
    print('Done! All fixtures generated.')


if __name__ == '__main__':
    generate_all()
