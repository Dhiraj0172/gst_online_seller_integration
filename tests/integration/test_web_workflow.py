"""
Integration test for complete Web UI & API workflow:
Authentication -> Profile Setup -> File Upload -> DB Normalization -> 
Statement Querying -> Transaction Edit & Delete -> GSTR-1 Excel & JSON Generation -> File Downloads.
"""
import os
import sys
import io
import json
import uuid
import pytest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from app.models import User, GSTProfile, Transaction, ImportHistory, GSTR1Generation

def test_full_web_application_workflow(app, db, client):
    """Test full application flow through HTTP client with clean tenant isolation."""
    suffix = uuid.uuid4().hex[:6]
    email = f'seller_ent_{suffix}@example.com'
    password = 'SecurePassword123!'
    gstin = f'27AAB{suffix[:4].upper()}1ZM'

    try:
        # 1. Register
        reg_resp = client.post('/register', data={
            'name': 'Seller Enterprise',
            'email': email,
            'password': password
        }, follow_redirects=True)
        assert reg_resp.status_code == 200

        # 2. Login
        login_resp = client.post('/login', data={
            'email': email,
            'password': password
        }, follow_redirects=True)
        assert login_resp.status_code == 200

        # 3. Create GST Profile
        prof_resp = client.post('/profiles/create', data={
            'gstin': gstin,
            'legal_name': 'Seller Enterprise Private Limited',
            'trade_name': 'Seller Enterprise',
            'state': '27',
            'frequency': 'Monthly',
            'financial_year': '2024-25'
        }, follow_redirects=True)
        assert prof_resp.status_code == 200

        # 4. Upload file (Amazon sample fixture)
        fixture_path = os.path.join(os.path.dirname(__file__), '..', 'fixtures', 'amazon_sample.xlsx')
        assert os.path.exists(fixture_path), "Fixture amazon_sample.xlsx must exist"

        with open(fixture_path, 'rb') as f:
            file_bytes = io.BytesIO(f.read())

        upload_resp = client.post('/import/upload', data={
            'file': (file_bytes, 'amazon_sample.xlsx'),
            'platform': 'Amazon'
        }, content_type='multipart/form-data', follow_redirects=True)
        assert upload_resp.status_code == 200

        # Verify Database state (strictly profile-scoped)
        with app.app_context():
            profile = GSTProfile.query.filter_by(gstin=gstin).first()
            assert profile is not None
            profile_id = profile.id

            imp = ImportHistory.query.filter_by(profile_id=profile_id).first()
            assert imp is not None
            assert imp.processing_status == 'COMPLETED'
            assert imp.total_rows == 168

            tx_count = Transaction.query.filter_by(profile_id=profile_id, is_deleted=False).count()
            assert tx_count == 168

            b2b_count = Transaction.query.filter_by(profile_id=profile_id, supply_type='B2B', is_deleted=False).count()
            b2c_count = Transaction.query.filter(
                Transaction.profile_id == profile_id,
                Transaction.supply_type.in_(['B2CS', 'B2CL']),
                Transaction.is_deleted == False
            ).count()
            cdnr_count = Transaction.query.filter_by(profile_id=profile_id, supply_type='CDNR', is_deleted=False).count()
            cdnur_count = Transaction.query.filter_by(profile_id=profile_id, supply_type='CDNUR', is_deleted=False).count()

            assert b2b_count == 50
            assert b2c_count == 103
            assert cdnr_count == 10
            assert cdnur_count == 5

            first_tx = Transaction.query.filter_by(profile_id=profile_id).first()
            tx_id = first_tx.id
            old_val = float(first_tx.taxable_value)

        # 5. Test Statement Edit
        edit_resp = client.post(f'/statement/edit/{tx_id}', json={
            'taxable_value': old_val + 100.0,
            'reason': 'Adjusted value via test'
        })
        assert edit_resp.status_code == 200

        # 6. Test Statement Delete
        del_resp = client.post(f'/statement/delete/{tx_id}', json={
            'reason': 'Deleted via test'
        })
        assert del_resp.status_code == 200

        with app.app_context():
            deleted_tx = db.session.get(Transaction, tx_id)
            assert deleted_tx.is_deleted is True

        # 7. Generate GSTR-1
        gen_resp = client.post('/generate/run', data={
            'return_period': '012025'
        }, follow_redirects=True)
        assert gen_resp.status_code == 200

        with app.app_context():
            gen = GSTR1Generation.query.filter_by(profile_id=profile_id).first()
            assert gen is not None
            assert gen.generation_status == 'COMPLETED'
            gen_id = gen.id

        # 8. Download Excel and verify sheets
        excel_resp = client.get(f'/generate/download/excel/{gen_id}')
        assert excel_resp.status_code == 200
        assert len(excel_resp.data) > 0
        wb = openpyxl.load_workbook(io.BytesIO(excel_resp.data))
        for expected_sheet in ['b2b', 'b2cs', 'cdnr', 'cdnur', 'hsn']:
            assert expected_sheet in wb.sheetnames, f"Expected sheet {expected_sheet} not in {wb.sheetnames}"
        wb.close()

        # 9. Download JSON and verify structure
        json_resp = client.get(f'/generate/download/json/{gen_id}')
        assert json_resp.status_code == 200
        gstr1_json = json.loads(json_resp.data.decode('utf-8'))
        assert gstr1_json['gstin'] == gstin
        assert gstr1_json['fp'] == '012025'
        assert len(gstr1_json.get('b2b', [])) > 0
        assert len(gstr1_json.get('b2cs', [])) > 0

    finally:
        # Clean teardown: delete user to cascade-delete profile, imports, txs, generations
        with app.app_context():
            try:
                user = User.query.filter_by(email=email).first()
                if user is not None:
                    db.session.delete(user)
                    db.session.commit()
            except Exception:
                db.session.rollback()
