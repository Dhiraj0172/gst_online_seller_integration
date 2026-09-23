"""Focused tests for the marketplace import adapters.

Covers, for every adapter format this project actually supports/documents:

  * registration and the documented-vs-filename-only inventory
  * header / sheet-format detection
  * representative valid files (the project's own marketplace fixtures)
  * malformed and unsupported files failing explicitly
  * the normalized transaction contract

Evidence for the platform formats lives in PLATFORM_ADAPTER_GUIDE.md,
README_GST_ONLINE_SELLER.md, app/templates/import.html,
tests/fixtures/generate_test_data.py and app/services/gstr1_excel_writer.py.
No marketplace format is asserted here that is not evidence-backed there.
"""
import csv
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from app.adapters.base import ImportRowStatus
from app.adapters.all_adapters import (
    CANONICAL_FIELDS, AmazonAdapter, CustomExcelAdapter, FlipkartAdapter,
    GSTR1GovtAdapter, MeeshoAdapter, MyntraAdapter,
)
from app.adapters.registry import detect_platform, get_adapter, list_platforms
from app.services.gstr1_excel_writer import GSTR1ExcelWriter


FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')

DOCUMENTED_PLATFORMS = {'Amazon', 'Flipkart', 'Meesho', 'Myntra',
                        'Custom_Excel', 'GSTR1_Govt'}
FILENAME_ONLY_PLATFORMS = {'JioMart', 'Snapdeal', 'TataCLiQ', 'Limeroad',
                           'Shopdeck', 'GlowRoad', 'Snapmint', 'Citymall',
                           'Roposo'}


def _fixture(name):
    path = os.path.join(FIXTURE_DIR, name)
    if not os.path.exists(path):
        pytest.skip(f'fixture {name} not generated yet')
    return path


def _workbook(name):
    return openpyxl.load_workbook(_fixture(name), data_only=True)


def _sheet_workbook(headers, rows, title='Data'):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    return workbook


# ---------------------------------------------------------------------------
# Inventory
# ---------------------------------------------------------------------------

class TestAdapterInventory:
    def test_all_registered_adapters_instantiate(self):
        platforms = list_platforms()
        assert len(platforms) >= 16
        for entry in platforms:
            adapter = get_adapter(entry['name'])
            assert adapter is not None, f'adapter {entry["name"]} missing'

    def test_expected_platform_names_present(self):
        names = {entry['name'] for entry in list_platforms()}
        assert DOCUMENTED_PLATFORMS <= names
        assert FILENAME_ONLY_PLATFORMS <= names

    def test_documented_adapters_declare_a_header_map(self):
        for name in ('Amazon', 'Flipkart', 'Meesho', 'Myntra', 'GSTR1_Govt'):
            adapter = get_adapter(name)
            assert adapter.FORMAT_DOCUMENTED is True
            assert adapter.HEADER_MAP, f'{name} has no documented column map'
            assert adapter.INSTRUCTIONS

    def test_filename_only_adapters_do_not_invent_a_format(self):
        for name in FILENAME_ONLY_PLATFORMS:
            adapter = get_adapter(name)
            assert adapter.FORMAT_DOCUMENTED is False
            assert adapter.HEADER_MAP == {}

    def test_custom_excel_is_the_catch_all_fallback(self):
        adapter = get_adapter('Custom_Excel')
        assert isinstance(adapter, CustomExcelAdapter)
        assert adapter.detect(None, 'anything.xlsx') is True
        assert adapter.CATCH_ALL_DETECT is True

    def test_detection_prefers_specific_adapter_over_catch_all(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            assert isinstance(detect_platform(workbook, 'flipkart_gst_report.xlsx'), FlipkartAdapter)
            # no filename hint: the "GST Report" sheet identifies Flipkart
            assert isinstance(detect_platform(workbook, 'report.xlsx'), FlipkartAdapter)
        finally:
            workbook.close()


# ---------------------------------------------------------------------------
# Detection / header detection
# ---------------------------------------------------------------------------

class TestHeaderDetection:
    def test_amazon_detects_by_filename_and_sheet(self):
        adapter = AmazonAdapter()
        assert adapter.detect(None, 'Amazon_MTR_012025.xlsx') is True
        workbook = _workbook('amazon_sample.xlsx')
        try:
            assert adapter.detect(workbook, 'download.xlsx') is True
        finally:
            workbook.close()
        assert adapter.detect(None, 'flipkart_report.xlsx') is False

    def test_flipkart_detects_by_filename_and_sheet(self):
        adapter = FlipkartAdapter()
        assert adapter.detect(None, 'flipkart_orders.xlsx') is True
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            assert adapter.detect(workbook, 'download.xlsx') is True
        finally:
            workbook.close()

    def test_meesho_detects_by_filename_and_sheet(self):
        adapter = MeeshoAdapter()
        assert adapter.detect(None, 'meesho_orders.xlsx') is True
        workbook = _workbook('meesho_sample.xlsx')
        try:
            assert adapter.detect(workbook, 'download.xlsx') is True
        finally:
            workbook.close()

    def test_myntra_detects_by_filename(self):
        adapter = MyntraAdapter()
        assert adapter.detect(None, 'myntra_gst_report.xlsx') is True
        assert adapter.detect(None, 'flipkart_report.xlsx') is False

    def test_gstr1_govt_detects_by_sheet_names(self, tmp_path):
        path = str(tmp_path / 'gstr1.xlsx')
        GSTR1ExcelWriter.generate_excel(
            {'b2b': [['29AALCS5765L1ZP', 'INV-1', '15-01-2025', 11800, '29', 'N', '',
                       'Regular', '', 18, 10000, 0]]},
            path,
        )
        workbook = openpyxl.load_workbook(path)
        try:
            adapter = GSTR1GovtAdapter()
            assert adapter.detect(workbook, 'downloaded_from_portal.xlsx') is True
        finally:
            workbook.close()

    def test_header_row_is_found_below_a_preamble(self):
        workbook = _sheet_workbook(
            ['Seller GST report', '', ''],
            [
                ['Report generated for period 012025', '', ''],
                ['Invoice Number', 'Invoice Date', 'Taxable Value', 'Tax Rate'],
                ['INV-1', '15-01-2025', 1000, 18],
            ],
            title='Report',
        )
        try:
            adapter = CustomExcelAdapter()
            is_valid, errors = adapter.validate(workbook)
            assert is_valid, errors
            result = adapter.parse(workbook)
            assert result.total_rows == 1
            assert result.metadata['header_rows'] == {'Report': 3}
            assert result.rows[0].normalized_data['invoice_number'] == 'INV-1'
        finally:
            workbook.close()

    def test_unrecognized_headers_are_reported_not_dropped(self):
        workbook = _workbook('extra_columns.xlsx')
        try:
            result = CustomExcelAdapter().parse(workbook)
            assert result.total_rows == 5
            assert result.metadata['unrecognized_headers'] == ['Extra Column 1', 'Internal ID']
            assert any('unrecognised column' in warning for warning in result.warnings)
        finally:
            workbook.close()


# ---------------------------------------------------------------------------
# Amazon
# ---------------------------------------------------------------------------

class TestAmazonAdapter:
    def test_validate_accepts_the_mtr_fixture(self):
        workbook = _workbook('amazon_sample.xlsx')
        try:
            assert AmazonAdapter().validate(workbook) == (True, [])
        finally:
            workbook.close()

    def test_parse_mtr_fixture(self):
        workbook = _workbook('amazon_sample.xlsx')
        try:
            result = AmazonAdapter().parse(workbook, 'amazon_sample.xlsx')
        finally:
            workbook.close()
        assert result.platform == 'Amazon'
        assert result.total_rows == 168
        assert result.error_rows == 0
        assert result.metadata['header_rows'] == {'MTR': 1}
        assert result.metadata['sheets_parsed'] == ['MTR']
        assert result.detected_period == '012025'
        assert result.detected_gstin == '27AABCU9603R1ZM'
        assert result.rows[0].sheet_name == 'MTR'

    def test_cancelled_orders_are_flagged_not_discarded(self):
        workbook = _workbook('amazon_sample.xlsx')
        try:
            result = AmazonAdapter().parse(workbook)
        finally:
            workbook.close()
        cancelled = [row for row in result.rows if row.normalized_data['cancellation_flag']]
        assert len(cancelled) == 3
        assert all(row.normalized_data['supply_type'] == 'Cancelled' for row in cancelled)
        assert all(row.status is ImportRowStatus.SUCCESS for row in cancelled)

    def test_credit_notes_map_to_note_type_and_are_retained(self):
        workbook = _workbook('amazon_sample.xlsx')
        try:
            result = AmazonAdapter().parse(workbook)
        finally:
            workbook.close()
        notes = [row for row in result.rows if row.normalized_data['note_type'] == 'CREDIT']
        assert len(notes) == 15
        assert all(row.normalized_data['supply_type'] == 'Credit Note' for row in notes)

    def test_normalized_output_matches_the_canonical_contract(self):
        workbook = _workbook('amazon_sample.xlsx')
        try:
            normalized = AmazonAdapter().parse(workbook).rows[0].normalized_data
        finally:
            workbook.close()
        for field in CANONICAL_FIELDS:
            assert field in normalized, f'canonical field {field} missing'
        assert normalized['invoice_number'] == 'INV-2025-0001'
        # dates are canonicalised to DD-MM-YYYY
        assert normalized['invoice_date'].count('-') == 2
        assert len(normalized['invoice_date'].split('-')[2]) == 4
        # place of supply is a 2-digit state code, original text kept
        assert normalized['place_of_supply'].isdigit()
        assert len(normalized['place_of_supply']) == 2
        assert normalized['place_of_supply_raw']
        assert normalized['marketplace_name'] == 'Amazon'
        assert normalized['source_platform'] == 'Amazon'
        assert isinstance(normalized['source_metadata'], dict)
        assert normalized['reverse_charge'] in ('Y', 'N')

    def test_ecommerce_gstin_and_marketplace_columns_are_mapped(self):
        workbook = _workbook('ecom_test.xlsx')
        try:
            result = get_adapter('Custom_Excel').parse(workbook)
        finally:
            workbook.close()
        assert result.error_rows == 0
        gstins = {row.normalized_data['ecommerce_gstin'] for row in result.rows}
        assert gstins == {'27AABCA1234B1ZM', '29AABCF5678D1ZP'}
        marketplaces = {row.normalized_data['marketplace_name'] for row in result.rows}
        assert marketplaces == {'Amazon', 'Flipkart'}


# ---------------------------------------------------------------------------
# Flipkart / Myntra
# ---------------------------------------------------------------------------

class TestFlipkartAdapter:
    def test_validate_and_parse_gst_report(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            adapter = FlipkartAdapter()
            assert adapter.validate(workbook) == (True, [])
            result = adapter.parse(workbook, 'flipkart_sample.xlsx')
        finally:
            workbook.close()
        assert result.total_rows == 118
        assert result.error_rows == 0
        assert result.metadata['header_rows'] == {'GST Report': 1}

    def test_flipkart_specific_columns_are_mapped(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            result = FlipkartAdapter().parse(workbook)
        finally:
            workbook.close()
        first = result.rows[0].normalized_data
        assert first['order_item_id'].startswith('OI-')
        assert first['description']                      # 'Product Title'
        assert first['unit_price'] > 0                    # 'Selling Price'
        assert 'unit_price' in result.metadata['sheet_columns']['GST Report']

    def test_returns_are_flagged(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            result = FlipkartAdapter().parse(workbook)
        finally:
            workbook.close()
        returns = [row for row in result.rows
                   if row.normalized_data['return_flag']
                   and row.normalized_data['supply_type'] == 'Return']
        assert len(returns) == 3
        # a marketplace return is not a cancellation
        assert all(row.normalized_data['cancellation_flag'] is False for row in returns)

    def test_credit_notes_use_the_flipkart_supply_type_column(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            result = FlipkartAdapter().parse(workbook)
        finally:
            workbook.close()
        notes = [row for row in result.rows if row.normalized_data['note_type']]
        assert len(notes) == 5
        assert {row.normalized_data['note_type'] for row in notes} == {'CREDIT'}


class TestMyntraAdapter:
    def test_myntra_reuses_the_documented_flipkart_layout(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            adapter = MyntraAdapter()
            assert adapter.HEADER_MAP == FlipkartAdapter.HEADER_MAP
            result = adapter.parse(workbook, 'myntra_gst_report.xlsx')
        finally:
            workbook.close()
        assert result.platform == 'Myntra'
        assert result.total_rows == 118
        assert result.rows[0].normalized_data['marketplace_name'] == 'Myntra'
        assert result.rows[0].normalized_data['source_platform'] == 'Myntra'

    def test_myntra_does_not_claim_the_flipkart_sheet_name(self):
        workbook = _workbook('flipkart_sample.xlsx')
        try:
            assert MyntraAdapter().detect(workbook, 'download.xlsx') is False
        finally:
            workbook.close()


# ---------------------------------------------------------------------------
# Meesho
# ---------------------------------------------------------------------------

class TestMeeshoAdapter:
    def test_validate_and_parse_orders_report(self):
        workbook = _workbook('meesho_sample.xlsx')
        try:
            adapter = MeeshoAdapter()
            assert adapter.validate(workbook) == (True, [])
            result = adapter.parse(workbook, 'meesho_sample.xlsx')
        finally:
            workbook.close()
        assert result.total_rows == 40
        assert result.error_rows == 0
        assert result.metadata['header_rows'] == {'Orders': 1}
        assert result.metadata['unrecognized_headers'] == []

    def test_meesho_specific_columns_are_mapped(self):
        workbook = _workbook('meesho_sample.xlsx')
        try:
            result = MeeshoAdapter().parse(workbook)
        finally:
            workbook.close()
        first = result.rows[0].normalized_data
        assert first['sub_order_id'].startswith('SO-')     # 'Sub Order ID'
        assert first['order_item_id'] == first['sub_order_id']
        assert first['order_id'].startswith('ORD-')        # parent order id
        assert first['description']                        # 'Product Name'
        assert first['hsn_sac']                            # 'HSN'
        assert first['taxable_value'] > 0                  # 'Taxable Amount'
        assert first['invoice_value'] > 0                  # 'Total'
        assert first['place_of_supply'].isdigit() and len(first['place_of_supply']) == 2

    def test_meesho_tax_components_are_read_from_short_headers(self):
        headers = ['Order ID', 'Sub Order ID', 'Invoice Number', 'Invoice Date',
                   'HSN', 'Taxable Amount', 'Total', 'CGST', 'SGST', 'IGST',
                   'Buyer State', 'Tax Rate']
        workbook = _sheet_workbook(headers, [
            ['ORD-1', 'SO-1', 'INV-1', '15-01-2025', '6109', 1000, 1180, 90, 90, 0, 'Delhi', 18],
            ['ORD-2', 'SO-2', 'INV-2', '16-01-2025', '6109', 2000, 2360, 0, 0, 360, 'Karnataka', 18],
        ], title='Orders')
        try:
            result = MeeshoAdapter().parse(workbook)
        finally:
            workbook.close()
        assert result.total_rows == 2
        intra, inter = result.rows[0].normalized_data, result.rows[1].normalized_data
        assert (intra['cgst_amount'], intra['sgst_amount'], intra['igst_amount']) == (90, 90, 0)
        assert (inter['cgst_amount'], inter['sgst_amount'], inter['igst_amount']) == (0, 0, 360)
        assert intra['place_of_supply'] == '07'
        assert inter['place_of_supply'] == '29'


# ---------------------------------------------------------------------------
# GSTR-1 government workbook
# ---------------------------------------------------------------------------

class TestGSTR1GovtAdapter:
    def _workbook_path(self, tmp_path):
        path = str(tmp_path / 'gstr1_012025.xlsx')
        GSTR1ExcelWriter.generate_excel({
            'b2b': [['29AALCS5765L1ZP', 'INV-1', '15-01-2025', 11800, '29', 'N', '',
                     'Regular', '27AABCA1234B1ZM', 18, 10000, 0]],
            'b2cs': [['OE', '29', '', 18, 5000, 0, '']],
            'cdnr': [['29AALCS5765L1ZP', 'CN-1', '20-01-2025', 'INV-1', '15-01-2025',
                      1180, '29', 'N', 'Regular', 'Credit Note', '', 18, 1000, 0]],
            'nil': [['Nil rated supplies', 0, 2500, 0]],
            'hsn': [['6109', 'T-Shirt', 'NOS', 10, 11800, 10000, 0, 900, 900, 0, 18]],
            'docs': [['Invoices for outward supply', 1, 2, 2, 0]],
        }, path)
        return path

    def test_validate_and_multi_sheet_parse(self, tmp_path):
        workbook = openpyxl.load_workbook(self._workbook_path(tmp_path))
        try:
            adapter = GSTR1GovtAdapter()
            assert adapter.validate(workbook) == (True, [])
            result = adapter.parse(workbook, 'gstr1_012025.xlsx')
        finally:
            workbook.close()
        assert result.platform == 'GSTR1_Govt'
        assert result.error_rows == 0
        assert result.metadata['sheets_parsed'] == ['b2b', 'b2cs', 'cdnr', 'nil', 'hsn']
        # the document-count sheet holds no transactions and is reported, not dropped
        assert result.metadata['sheets_ignored'] == ['docs']
        # the govt sheet layout carries no seller-GSTIN column, so no seller GSTIN
        # can be detected from this file; the period is detected from row dates
        assert result.detected_gstin is None
        assert result.detected_period == '012025'

    def test_sheet_rows_are_normalized_per_sheet_kind(self, tmp_path):
        workbook = openpyxl.load_workbook(self._workbook_path(tmp_path))
        try:
            result = GSTR1GovtAdapter().parse(workbook)
        finally:
            workbook.close()
        by_sheet = {row.sheet_name: row for row in result.rows}

        b2b = by_sheet['b2b'].normalized_data
        assert b2b['invoice_number'] == 'INV-1'
        assert b2b['customer_gstin'] == '29AALCS5765L1ZP'
        assert b2b['place_of_supply'] == '29'
        assert b2b['taxable_value'] == 10000
        assert b2b['ecommerce_gstin'] == '27AABCA1234B1ZM'

        cdnr_row = by_sheet['cdnr']
        assert cdnr_row.normalized_data['note_type'] == 'CREDIT'
        assert cdnr_row.normalized_data['note_number'] == 'CN-1'
        assert cdnr_row.normalized_data['invoice_number'] == 'CN-1'
        assert cdnr_row.status is ImportRowStatus.WARNING

        # consolidated rows have no invoice number: kept as WARNING, never dropped
        assert by_sheet['b2cs'].status is ImportRowStatus.WARNING
        assert by_sheet['b2cs'].normalized_data['taxable_value'] == 5000
        assert by_sheet['nil'].normalized_data['item_type'] == 'EXEMPT'
        assert by_sheet['nil'].normalized_data['taxable_value'] == 2500
        assert by_sheet['hsn'].normalized_data['hsn_sac'] == '6109'

    def test_unrecognized_sheets_fail_explicitly(self):
        workbook = _sheet_workbook(
            ['Col A', 'Col B'],
            [[1, 2]],
            title='b2b',
        )
        try:
            is_valid, errors = GSTR1GovtAdapter().validate(workbook)
        finally:
            workbook.close()
        assert is_valid is False
        assert any('recognizable header row' in message for message in errors)


# ---------------------------------------------------------------------------
# Malformed / unsupported files
# ---------------------------------------------------------------------------

class TestMalformedAndUnsupportedFiles:
    @pytest.mark.parametrize('platform', ['Amazon', 'Flipkart', 'Meesho', 'Myntra',
                                          'Custom_Excel', 'Generic'])
    def test_unrecognized_headers_fail_explicitly(self, platform):
        workbook = _workbook('wrong_headers.xlsx')
        try:
            is_valid, errors = get_adapter(platform).validate(workbook)
        finally:
            workbook.close()
        assert is_valid is False
        assert errors and 'recognizable header row' in errors[0]
        assert 'Col A' not in errors[0]  # message lists expected columns, kept short

    def test_parse_reports_file_level_error_and_no_rows(self):
        workbook = _workbook('wrong_headers.xlsx')
        try:
            result = AmazonAdapter().parse(workbook, 'wrong_headers.xlsx')
        finally:
            workbook.close()
        assert result.total_rows == 0
        assert result.errors
        assert 'recognizable header row' in result.errors[0]

    def test_empty_workbook_fails_explicitly(self):
        workbook = _workbook('empty_file.xlsx')
        try:
            is_valid, errors = AmazonAdapter().validate(workbook)
        finally:
            workbook.close()
        assert is_valid is False
        assert errors

    def test_headers_without_data_rows_fail_explicitly(self):
        workbook = _sheet_workbook(['Invoice Number', 'Taxable Value'], [])
        try:
            is_valid, errors = AmazonAdapter().validate(workbook)
        finally:
            workbook.close()
        assert is_valid is False
        assert any('No data rows' in message or 'recognizable header row' in message
                   for message in errors)

    def test_missing_required_column_fails_explicitly(self):
        workbook = _sheet_workbook(['Order ID', 'Taxable Value', 'Tax Rate'],
                                   [['ORD-1', 100, 18]])
        try:
            is_valid, errors = AmazonAdapter().validate(workbook)
        finally:
            workbook.close()
        assert is_valid is False
        assert any('missing required column' in message.lower() for message in errors)
        assert any('invoice number' in message.lower() for message in errors)

    def test_no_input_fails_explicitly(self):
        assert AmazonAdapter().validate(None) == (False, ['No workbook or data provided.'])
        result = AmazonAdapter().parse(None)
        assert result.total_rows == 0
        assert result.errors == ['No workbook or data provided.']

    def test_missing_file_fails_explicitly(self):
        is_valid, errors = AmazonAdapter().validate('/no/such/file.xlsx')
        assert is_valid is False
        assert 'File not found' in errors[0]

    def test_unsupported_extension_fails_explicitly(self, tmp_path):
        path = tmp_path / 'orders.txt'
        path.write_text('Invoice Number,Taxable Value\nINV-1,100\n', encoding='utf-8')
        is_valid, errors = AmazonAdapter().validate(str(path))
        assert is_valid is False
        assert 'Unsupported file type' in errors[0]

    def test_invalid_gstin_and_dates_are_kept_with_warnings(self):
        workbook = _workbook('malformed_gstin_test.xlsx')
        try:
            result = CustomExcelAdapter().parse(workbook)
        finally:
            workbook.close()
        assert result.total_rows == 5
        assert result.error_rows == 0
        assert all(row.status is ImportRowStatus.WARNING for row in result.rows)
        assert all('rejected by validation' in warning
                   for row in result.rows for warning in row.warnings)
        # the supplied value is never silently replaced or discarded
        assert result.rows[0].normalized_data['customer_gstin'] == 'INVALID'

    def test_unparseable_dates_are_reported(self):
        workbook = _workbook('invalid_dates_test.xlsx')
        try:
            result = CustomExcelAdapter().parse(workbook)
        finally:
            workbook.close()
        assert result.total_rows == 5
        bad = [row for row in result.rows
               if any('not a recognized date' in warning for warning in row.warnings)]
        assert len(bad) >= 4

    def test_rows_missing_the_document_number_are_errors_not_dropped(self):
        workbook = _sheet_workbook(
            ['Invoice Number', 'Invoice Date', 'Taxable Value', 'Tax Rate',
             'Place of Supply', 'CGST Amount', 'SGST Amount'],
            [['', '15-01-2025', 1000, 18, '29', 0, 0],
             ['INV-2', '16-01-2025', 2000, 18, '29', 0, 0]],
        )
        try:
            result = CustomExcelAdapter().parse(workbook)
        finally:
            workbook.close()
        assert result.total_rows == 2
        assert result.error_rows == 1
        assert result.rows[0].status is ImportRowStatus.ERROR
        assert 'Missing invoice number' in result.rows[0].errors[0]
        # the row without a document number is kept as an ERROR row, never dropped
        assert result.rows[0].normalized_data['taxable_value'] == 1000
        # a row with every required column present is clean
        assert result.rows[1].status is ImportRowStatus.SUCCESS
        assert result.rows[1].warnings == []

    def test_non_numeric_money_values_are_reported(self):
        workbook = _sheet_workbook(
            ['Invoice Number', 'Invoice Date', 'Taxable Value', 'Tax Rate'],
            [['INV-1', '15-01-2025', 'not-a-number', 18]],
        )
        try:
            result = CustomExcelAdapter().parse(workbook)
        finally:
            workbook.close()
        assert result.rows[0].status is ImportRowStatus.WARNING
        assert any('not numeric' in warning for warning in result.rows[0].warnings)
        assert result.rows[0].normalized_data['taxable_value'] == 0


# ---------------------------------------------------------------------------
# CSV support (declared in SUPPORTED_FILE_TYPES)
# ---------------------------------------------------------------------------

class TestCsvSupport:
    def test_csv_file_is_parsed(self, tmp_path):
        path = tmp_path / 'custom_export.csv'
        with open(str(path), 'w', newline='', encoding='utf-8') as handle:
            writer = csv.writer(handle)
            writer.writerow(['Invoice Number', 'Invoice Date', 'Taxable Value',
                             'Tax Rate', 'Place of Supply', 'HSN'])
            writer.writerow(['INV-1', '15-01-2025', '1000', '18', '29-Karnataka', '6109'])
        adapter = CustomExcelAdapter()
        assert adapter.validate(str(path)) == (True, [])
        result = adapter.parse(str(path), 'custom_export.csv')
        assert result.total_rows == 1
        assert result.rows[0].normalized_data['invoice_number'] == 'INV-1'
        assert result.rows[0].normalized_data['place_of_supply'] == '29'


# ---------------------------------------------------------------------------
# Normalization contract
# ---------------------------------------------------------------------------

class TestNormalizationContract:
    def test_generic_normalize_covers_the_canonical_contract(self):
        row = {
            'Order ID': 'ORD-1',
            'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025',
            'Buyer GSTIN': '29AALCS5765L1ZP',
            'Place of Supply': '29-Karnataka',
            'Seller GSTIN': '27AABCU9603R1ZM',
            'HSN/SAC': '6109',
            'Product Description': 'T-Shirt',
            'Quantity': 2,
            'Taxable Value': '1,000.00',
            'IGST Rate': 18,
            'IGST Amount': 180,
            'Invoice Value': 1180,
            'Tax Rate': 18,
            'Reverse Charge': 'Y',
            'E-Commerce GSTIN': '27AABCA1234B1ZM',
            'Supply Type': 'Regular',
        }
        normalized = CustomExcelAdapter().normalize(row)
        for field in CANONICAL_FIELDS:
            assert field in normalized
        assert normalized['taxable_value'] == 1000          # thousands separator handled
        assert normalized['invoice_date'] == '15-01-2025'
        assert normalized['place_of_supply'] == '29'
        assert normalized['place_of_supply_raw'] == '29-Karnataka'
        assert normalized['reverse_charge'] == 'Y'
        assert normalized['ecommerce_gstin'] == '27AABCA1234B1ZM'
        assert normalized['return_flag'] is False
        assert normalized['cancellation_flag'] is False
        assert normalized['note_type'] == ''
        assert normalized['_warnings'] == []

    def test_derived_tax_rate_when_only_components_are_supplied(self):
        row = {
            'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025',
            'Taxable Value': 1000,
            'CGST Rate': 9,
            'SGST Rate': 9,
            'CGST Amount': 90,
            'SGST Amount': 90,
        }
        normalized = CustomExcelAdapter().normalize(row)
        assert normalized['tax_rate'] == 18
        assert normalized['tax_rate_derived'] is True

    def test_place_of_supply_falls_back_to_customer_gstin(self):
        row = {
            'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025',
            'Taxable Value': 1000,
            'Tax Rate': 18,
            'Buyer GSTIN': '29AALCS5765L1ZP',
        }
        normalized = CustomExcelAdapter().normalize(row)
        assert normalized['place_of_supply'] == '29'
        assert any('derived from the customer GSTIN' in warning
                   for warning in normalized['_warnings'])

    def test_rate_without_amounts_is_warned_about(self):
        row = {
            'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025',
            'Taxable Value': 1000,
            'Tax Rate': 18,
        }
        normalized = CustomExcelAdapter().normalize(row)
        assert any('Tax amounts are not present' in warning
                   for warning in normalized['_warnings'])

    def test_explicit_zero_tax_amounts_are_not_reported_as_missing(self):
        """Regression: 0.00 in a tax amount column is data, not a missing column."""
        row = {
            'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025',
            'Taxable Value': 1000,
            'Tax Rate': 18,
            'Place of Supply': '29',
            'CGST Amount': 0,
            'SGST Amount': 0,
            'IGST Amount': 0,
        }
        normalized = CustomExcelAdapter().normalize(row)
        assert normalized['_warnings'] == []
        assert normalized['cgst_amount'] == 0
        assert normalized['igst_amount'] == 0

    def test_debit_note_supply_type_is_recognized(self):
        row = {
            'Invoice Number': 'DN-1',
            'Invoice Date': '15-01-2025',
            'Taxable Value': 500,
            'Tax Rate': 18,
            'Supply Type': 'Debit Note',
        }
        normalized = CustomExcelAdapter().normalize(row)
        assert normalized['note_type'] == 'DEBIT'
        assert normalized['note_type_raw'] == 'Debit Note'

    def test_return_and_cancellation_status_text(self):
        base = {'Invoice Number': 'INV-1', 'Invoice Date': '15-01-2025',
                'Taxable Value': 1000, 'Tax Rate': 18}
        returned = CustomExcelAdapter().normalize(dict(base, **{'Order Status': 'Returned'}))
        cancelled = CustomExcelAdapter().normalize(dict(base, **{'Order Status': 'Cancelled'}))
        assert returned['return_flag'] is True and returned['cancellation_flag'] is False
        assert cancelled['cancellation_flag'] is True

    def test_normalized_survives_json_dumps(self):
        import json
        workbook = _workbook('meesho_sample.xlsx')
        try:
            result = MeeshoAdapter().parse(workbook)
        finally:
            workbook.close()
        encoded = json.dumps({key: str(value) for key, value in
                              result.rows[0].normalized_data.items()})
        assert 'INV-2025-0401' in encoded
