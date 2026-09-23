"""Focused tests for CSV import support and duplicate detection.

CSV is expected to behave exactly like the Excel path: same header-driven
detection, same normalization, same explicit errors. Duplicate detection is
expected to be conservative (source-aware) and to mark rows, never drop them.
"""
import csv
import json
import os
import sys
from decimal import Decimal

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from app.adapters.all_adapters import CustomExcelAdapter, FlipkartAdapter
from app.adapters.base import ImportRowStatus
from app.models import ImportHistory, Transaction
from app.services.duplicate_service import (
    DUPLICATE_EXISTING, DUPLICATE_FILE, DUPLICATE_IN_FILE, build_fingerprint,
    document_identity, file_duplicate_report, find_existing_transactions,
    find_prior_file_import,
)
from app.utils.csv_utils import read_csv_rows, sniff_delimiter

HEADERS = ['Invoice Number', 'Invoice Date', 'Place of Supply', 'Buyer GSTIN',
           'Seller GSTIN', 'HSN/SAC', 'Product Description', 'Quantity',
           'Taxable Value', 'Tax Rate', 'CGST Amount', 'SGST Amount',
           'IGST Amount', 'Invoice Value', 'Supply Type', 'Order Item ID']

ROW_INTRA = ['INV-1', '15-01-2025', '29-Karnataka', '29AALCS5765L1ZP',
             '27AABCU9603R1ZM', '6109', 'Cotton T-Shirt', '2', '1,000.00', '18',
             '90', '90', '0', '1,180.00', 'Regular', 'OI-1']
ROW_INTER = ['INV-2', '16-01-2025', '29-Karnataka', '', '27AABCU9603R1ZM',
             '6109', 'Cotton T-Shirt', '1', '500.00', '18', '0', '0', '90',
             '590.00', 'Regular', 'OI-2']


def write_csv(path, rows, delimiter=',', encoding='utf-8', bom=False):
    with open(str(path), 'w', newline='', encoding=encoding) as handle:
        if bom:
            handle.write('\ufeff')
        writer = csv.writer(handle, delimiter=delimiter)
        for row in rows:
            writer.writerow(row)
    return str(path)


def write_xlsx(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Data'
    for row in rows:
        sheet.append(row)
    workbook.save(str(path))
    workbook.close()
    return str(path)


# ---------------------------------------------------------------------------
# CSV reading
# ---------------------------------------------------------------------------

class TestCsvReading:
    def test_valid_comma_csv(self, tmp_path):
        path = write_csv(tmp_path / 'orders.csv', [HEADERS, ROW_INTRA])
        result = read_csv_rows(path)
        assert result.errors == []
        assert result.delimiter == ','
        assert result.encoding == 'utf-8-sig'
        assert len(result.rows) == 2
        assert result.rows[0][0] == 'Invoice Number'

    def test_utf8_bom_is_stripped(self, tmp_path):
        path = write_csv(tmp_path / 'bom.csv', [HEADERS, ROW_INTRA], bom=True)
        with open(path, 'rb') as handle:
            assert handle.read(3) == b'\xef\xbb\xbf'   # BOM really present
        result = read_csv_rows(path)
        assert result.errors == []
        assert result.rows[0][0] == 'Invoice Number'
        assert not result.rows[0][0].startswith('\ufeff')

    @pytest.mark.parametrize('delimiter', [';', '\t', '|'])
    def test_alternate_delimiters_are_detected(self, tmp_path, delimiter):
        path = write_csv(tmp_path / 'alt.csv', [HEADERS, ROW_INTRA], delimiter=delimiter)
        result = read_csv_rows(path)
        assert result.errors == []
        assert result.delimiter == delimiter
        assert len(result.rows[0]) == len(HEADERS)

    def test_delimiter_sniffing_defaults_to_comma(self):
        assert sniff_delimiter('a,b,c\n1,2,3') == ','
        assert sniff_delimiter('') == ','

    def test_cp1252_fallback_decoding(self, tmp_path):
        path = str(tmp_path / 'latin.csv')
        with open(path, 'wb') as handle:
            handle.write('Invoice Number,Product Description\nINV-1,Régulier café\n'.encode('cp1252'))
        result = read_csv_rows(path)
        assert result.errors == []
        assert result.encoding == 'cp1252'
        assert 'café' in result.rows[1][1]

    def test_empty_csv_is_an_explicit_error(self, tmp_path):
        path = write_csv(tmp_path / 'empty.csv', [])
        result = read_csv_rows(path)
        assert result.errors
        assert 'empty' in result.errors[0].lower()

    def test_binary_content_named_csv_is_rejected(self, tmp_path):
        path = str(tmp_path / 'really_xlsx.csv')
        openpyxl.Workbook().save(path)
        result = read_csv_rows(path)
        assert result.errors
        assert 'not text CSV' in result.errors[0]

    def test_missing_file_is_an_explicit_error(self):
        result = read_csv_rows('/no/such/file.csv')
        assert result.errors == ['File not found: /no/such/file.csv']


# ---------------------------------------------------------------------------
# CSV through the adapter pipeline
# ---------------------------------------------------------------------------

class TestCsvThroughAdapter:
    def test_csv_row_normalizes_like_the_excel_row(self, tmp_path):
        csv_path = write_csv(tmp_path / 'orders.csv', [HEADERS, ROW_INTRA])
        xlsx_path = write_xlsx(tmp_path / 'orders.xlsx', [HEADERS, ROW_INTRA])

        csv_result = CustomExcelAdapter().parse(csv_path, 'orders.csv')
        workbook = openpyxl.load_workbook(xlsx_path)
        try:
            xlsx_result = CustomExcelAdapter().parse(workbook, 'orders.xlsx')
        finally:
            workbook.close()

        assert csv_result.total_rows == xlsx_result.total_rows == 1
        csv_norm = csv_result.rows[0].normalized_data
        xlsx_norm = xlsx_result.rows[0].normalized_data
        for field in ('invoice_number', 'invoice_date', 'place_of_supply',
                      'customer_gstin', 'hsn_sac', 'description',
                      'taxable_value', 'cgst_amount', 'sgst_amount',
                      'igst_amount', 'invoice_value', 'tax_rate',
                      'marketplace_name', 'source_platform'):
            assert csv_norm[field] == xlsx_norm[field], field

    def test_csv_normalization_details(self, tmp_path):
        path = write_csv(tmp_path / 'orders.csv', [HEADERS, ROW_INTRA, ROW_INTER])
        result = CustomExcelAdapter().parse(path, 'orders.csv')
        assert result.error_rows == 0
        first, second = result.rows[0].normalized_data, result.rows[1].normalized_data

        assert first['invoice_number'] == 'INV-1'
        assert first['invoice_date'] == '15-01-2025'      # canonical DD-MM-YYYY
        assert first['place_of_supply'] == '29'           # 2-digit state code
        assert first['place_of_supply_raw'] == '29-Karnataka'
        assert first['customer_gstin'] == '29AALCS5765L1ZP'
        assert first['taxable_value'] == Decimal('1000.00')   # thousands separator
        assert first['cgst_amount'] == Decimal('90')
        assert first['sgst_amount'] == Decimal('90')
        assert first['igst_amount'] == Decimal('0')
        assert first['invoice_value'] == Decimal('1180.00')
        assert first['order_item_id'] == 'OI-1'
        assert first['source_metadata']['platform'] == 'Custom_Excel'
        assert second['igst_amount'] == Decimal('90')

    def test_csv_diagnostics_are_recorded(self, tmp_path):
        path = write_csv(tmp_path / 'orders.csv', [HEADERS, ROW_INTRA], delimiter=';')
        result = CustomExcelAdapter().parse(path, 'orders.csv')
        assert result.metadata['source_type'] == 'csv'
        assert result.metadata['csv_delimiter'] == 'semicolon'
        assert result.metadata['csv_encoding'] == 'utf-8-sig'

    def test_flipkart_shaped_csv_is_picked_up_by_the_flipkart_adapter(self, tmp_path):
        headers = ['Order Item ID', 'Order ID', 'Invoice Number', 'Invoice Date',
                   'Product Title', 'Selling Price', 'Quantity', 'Taxable Value',
                   'Tax Rate', 'Invoice Value', 'Place of Supply', 'Supply Type']
        row = ['OI-9', 'ORD-9', 'INV-9', '15-01-2025', 'Cotton T-Shirt', '100',
               '2', '200', '18', '236', '29-Karnataka', 'Regular']
        path = write_csv(tmp_path / 'flipkart_gst_report.csv', [headers, row])
        adapter = FlipkartAdapter()
        assert adapter.validate(path) == (True, [])
        result = adapter.parse(path, 'flipkart_gst_report.csv')
        assert result.total_rows == 1
        normalized = result.rows[0].normalized_data
        assert normalized['order_item_id'] == 'OI-9'
        assert normalized['description'] == 'Cotton T-Shirt'
        assert normalized['unit_price'] == Decimal('100')

    def test_unrecognized_csv_headers_are_rejected_explicitly(self, tmp_path):
        path = write_csv(tmp_path / 'junk.csv', [['Col A', 'Col B'], ['1', '2']])
        is_valid, errors = CustomExcelAdapter().validate(path)
        assert is_valid is False
        assert 'recognizable header row' in errors[0]

    def test_csv_without_document_number_keeps_the_row_as_error(self, tmp_path):
        row = list(ROW_INTRA)
        row[0] = ''
        path = write_csv(tmp_path / 'bad.csv', [HEADERS, row])
        result = CustomExcelAdapter().parse(path, 'bad.csv')
        assert result.total_rows == 1
        assert result.error_rows == 1
        assert result.rows[0].status is ImportRowStatus.ERROR
        assert 'Missing invoice number' in result.rows[0].errors[0]

    def test_csv_values_beyond_the_header_are_kept_not_dropped(self, tmp_path):
        path = write_csv(tmp_path / 'extra.csv', [HEADERS, ROW_INTRA + ['stray']])
        result = CustomExcelAdapter().parse(path, 'extra.csv')
        raw = result.rows[0].raw_data
        assert raw['col_16'] == 'stray'
        assert result.rows[0].status is ImportRowStatus.WARNING
        assert any('beyond the header columns' in warning
                   for warning in result.rows[0].warnings)

    def test_csv_short_row_is_kept_with_missing_values_as_none(self, tmp_path):
        short = ROW_INTRA[:3]                      # trailing values absent entirely
        path = write_csv(tmp_path / 'short.csv', [HEADERS, short])
        result = CustomExcelAdapter().parse(path, 'short.csv')
        assert result.total_rows == 1               # the row is not dropped
        row = result.rows[0]
        assert row.raw_data['Invoice Number'] == 'INV-1'
        assert row.raw_data['Invoice Value'] is None
        assert row.raw_data['Order Item ID'] is None
        assert row.status is not ImportRowStatus.SUCCESS

    def test_empty_csv_file_is_rejected_explicitly(self, tmp_path):
        path = write_csv(tmp_path / 'empty.csv', [])
        is_valid, errors = CustomExcelAdapter().validate(path)
        assert is_valid is False
        assert errors


# ---------------------------------------------------------------------------
# Fingerprint / business key
# ---------------------------------------------------------------------------

def normalized_row(**overrides):
    base = {
        'invoice_number': 'INV-1', 'invoice_date': '15-01-2025',
        'place_of_supply': '29', 'customer_gstin': '29AALCS5765L1ZP',
        'hsn_sac': '6109', 'description': 'Cotton T-Shirt',
        'quantity': Decimal('2'), 'taxable_value': Decimal('1000.00'),
        'cgst_amount': Decimal('90'), 'sgst_amount': Decimal('90'),
        'igst_amount': Decimal('0'), 'cess_amount': Decimal('0'),
        'invoice_value': Decimal('1180.00'), 'order_item_id': 'OI-1',
        'source_platform': 'Flipkart', 'marketplace_name': 'Flipkart',
        'note_number': '', 'note_type': '',
    }
    base.update(overrides)
    return base


class TestFingerprint:
    def test_same_row_same_fingerprint(self):
        assert build_fingerprint(normalized_row(), 1, 'Flipkart') == \
               build_fingerprint(normalized_row(), 1, 'Flipkart')

    def test_same_invoice_number_from_another_marketplace_is_not_a_duplicate(self):
        other = normalized_row(source_platform='Meesho', marketplace_name='Meesho')
        assert build_fingerprint(normalized_row(), 1, 'Flipkart') != \
               build_fingerprint(other, 1, 'Meesho')

    def test_different_profile_is_not_a_duplicate(self):
        assert build_fingerprint(normalized_row(), 1, 'Flipkart') != \
               build_fingerprint(normalized_row(), 2, 'Flipkart')

    @pytest.mark.parametrize('changed', [
        {'taxable_value': Decimal('1000.01')},
        {'invoice_value': Decimal('1180.01')},
        {'tax_rate': Decimal('12')},
        {'igst_amount': Decimal('90')},
        {'quantity': Decimal('3')},
        {'hsn_sac': '6205'},
        {'customer_gstin': ''},
        {'place_of_supply': '27'},
        {'invoice_date': '16-01-2025'},
        {'order_item_id': 'OI-2'},
        {'description': 'Cotton T-Shirt XL'},
    ])
    def test_any_material_difference_changes_the_key(self, changed):
        assert build_fingerprint(normalized_row(**changed), 1, 'Flipkart') != \
               build_fingerprint(normalized_row(), 1, 'Flipkart')

    def test_two_lines_of_one_invoice_stay_distinct(self):
        first = normalized_row(order_item_id='OI-1')
        second = normalized_row(order_item_id='OI-2')
        assert build_fingerprint(first, 1, 'Flipkart') != build_fingerprint(second, 1, 'Flipkart')

    def test_note_rows_use_the_note_number_as_identity(self):
        note = normalized_row(invoice_number='CN-1', note_number='CN-1', note_type='CREDIT')
        assert document_identity(note) == ('NOTE', 'CN-1')
        assert document_identity(normalized_row()) == ('INVOICE', 'INV-1')
        assert build_fingerprint(note, 1, 'Flipkart') != \
               build_fingerprint(normalized_row(), 1, 'Flipkart')

    def test_missing_optional_fields_are_stable(self):
        sparse = normalized_row(customer_gstin='', hsn_sac='', place_of_supply='', order_item_id='')
        assert build_fingerprint(sparse, 1, 'Flipkart') == \
               build_fingerprint(dict(sparse), 1, 'Flipkart')


# ---------------------------------------------------------------------------
# Duplicate lookups against the database
# ---------------------------------------------------------------------------

class TestDuplicateLookups:
    def test_prior_file_import_only_counts_completed_imports(self, app, db):
        with app.app_context():
            from app.models import GSTProfile, User
            import uuid
            suffix = uuid.uuid4().hex[:8]
            user = User(username=f'dup_{suffix}', email=f'dup_{suffix}@test.com')
            user.set_password('x')
            db.session.add(user)
            db.session.commit()
            profile = GSTProfile(
                user_id=user.id, gstin=f'29DUP{suffix[:4].upper()}1234A1Z5',
                legal_name='Dup', state_code='29', state_name='Karnataka',
                financial_year='2024-25', filing_frequency='monthly',
            )
            db.session.add(profile)
            db.session.commit()
            try:
                failed = ImportHistory(
                    user_id=user.id, profile_id=profile.id, file_name='bad.xlsx',
                    original_file_name='bad.xlsx', file_hash='hash-a',
                    platform_name='Flipkart', return_period='012025',
                    financial_year='2024-25', processing_status='FAILED', total_rows=0,
                )
                db.session.add(failed)
                db.session.commit()
                # a FAILED import of the same bytes must not block a retry
                assert find_prior_file_import(profile.id, 'hash-a') is None

                done = ImportHistory(
                    user_id=user.id, profile_id=profile.id, file_name='good.xlsx',
                    original_file_name='good.xlsx', file_hash='hash-b',
                    platform_name='Flipkart', return_period='012025',
                    financial_year='2024-25', processing_status='COMPLETED', total_rows=5,
                )
                db.session.add(done)
                db.session.commit()
                prior = find_prior_file_import(profile.id, 'hash-b')
                assert prior is not None and prior.id == done.id
                assert json.loads(json.dumps(file_duplicate_report(prior)))['code'] == DUPLICATE_FILE
            finally:
                db.session.delete(user)
                db.session.commit()

    def test_existing_transactions_are_found_by_fingerprint(self, app, db):
        with app.app_context():
            from app.models import GSTProfile, RawImport, User
            import uuid
            suffix = uuid.uuid4().hex[:8]
            user = User(username=f'fp_{suffix}', email=f'fp_{suffix}@test.com')
            user.set_password('x')
            db.session.add(user)
            db.session.commit()
            profile = GSTProfile(
                user_id=user.id, gstin=f'29FPX{suffix[:4].upper()}1234A1Z5',
                legal_name='Fp', state_code='29', state_name='Karnataka',
                financial_year='2024-25', filing_frequency='monthly',
            )
            db.session.add(profile)
            db.session.commit()
            try:
                history = ImportHistory(
                    user_id=user.id, profile_id=profile.id, file_name='f.xlsx',
                    original_file_name='f.xlsx', platform_name='Flipkart',
                    return_period='012025', financial_year='2024-25',
                    processing_status='COMPLETED', total_rows=1,
                )
                db.session.add(history)
                db.session.commit()
                raw = RawImport(import_history_id=history.id, sheet_name='S',
                                row_number=2, raw_data='{}', status='SUCCESS')
                db.session.add(raw)
                db.session.commit()

                fingerprint = build_fingerprint(normalized_row(), profile.id, 'Flipkart')
                transaction = Transaction(
                    import_history_id=history.id, profile_id=profile.id,
                    raw_import_id=raw.id, invoice_number='INV-1',
                    taxable_value=Decimal('1000.00'), row_fingerprint=fingerprint,
                    is_deleted=False,
                )
                db.session.add(transaction)
                db.session.commit()

                found = find_existing_transactions(profile.id, [fingerprint, 'other'])
                assert list(found) == [fingerprint]
                assert found[fingerprint].id == transaction.id

                # a soft-deleted row must not block a legitimate re-import
                transaction.is_deleted = True
                db.session.commit()
                assert find_existing_transactions(profile.id, [fingerprint]) == {}
            finally:
                db.session.delete(user)
                db.session.commit()
