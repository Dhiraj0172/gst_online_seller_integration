"""Phase 5B focused tests for Amazon MTR Production Adapter Hardening.

Covers all 15 required areas from Phase 5B:
1. Valid Amazon MTR detection (sheet + headers)
2. Invalid Amazon header rejection
3. Wrong sheet rejection (non-MTR sheet)
4. Malformed workbook rejection (corrupt, empty, unsupported extension)
5. Valid row parsing (Regular, Credit Note, Cancelled)
6. Malformed row handling (negative values, missing invoice number, unparseable date)
7. Decimal precision (exact Decimal parsing without float drift)
8. Canonical mapping (CanonicalTransaction contract and dictionary representation)
9. Duplicate rows in file (in-file duplicate marked SKIPPED)
10. Full-file re-import handling (duplicate file detection and DUPLICATE_EXISTING)
11. ImportHistory result recording (status, counts, metadata)
12. Persistence & reconciliation (Transactions, RawImports, exact taxable sum)
13. Rollback on injected failure (atomic rollback, zero orphaned records)
14. Tenant isolation (scoping across distinct profiles)
15. Unknown workbook not misclassified as Amazon (Flipkart, Meesho, generic)
"""
import io
import json
import os
import uuid
from decimal import Decimal
from unittest.mock import patch

import openpyxl
import pytest

from app.adapters.amazon import AmazonAdapter
from app.adapters.base import ImportResult, ImportRow, ImportRowStatus
from app.adapters.canonical import CanonicalTransaction
from app.adapters.registry import (
    detect_platform,
    get_adapter,
    register_adapter,
    auto_register_adapters,
    _ADAPTER_REGISTRY,
)
from app.extensions import db as _db
from app.models import GSTProfile, ImportHistory, RawImport, Transaction, User
from app.services.import_service import process_import


FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')


@pytest.fixture
def amazon_tenants(app, db):
    """Provide two distinct users and GST profiles for tenant isolation tests."""
    suffix = uuid.uuid4().hex[:8]
    u1_name = f'amz_user1_{suffix}'
    u2_name = f'amz_user2_{suffix}'
    u1_email = f'amz1_{suffix}@example.com'
    u2_email = f'amz2_{suffix}@example.com'
    gstin1 = f'27AABC{suffix[:4].upper()}1Z5'
    gstin2 = f'29BBCD{suffix[:4].upper()}1Z6'

    with app.app_context():
        u1 = User(username=u1_name, email=u1_email)
        u1.set_password('Pass123!')
        _db.session.add(u1)

        u2 = User(username=u2_name, email=u2_email)
        u2.set_password('Pass123!')
        _db.session.add(u2)
        _db.session.commit()

        p1 = GSTProfile(
            user_id=u1.id,
            gstin=gstin1,
            legal_name='Amazon Seller 1',
            trade_name='Store 1',
            state_code='27',
            state_name='Maharashtra',
            financial_year='2024-25',
            filing_frequency='Monthly'
        )
        _db.session.add(p1)

        p2 = GSTProfile(
            user_id=u2.id,
            gstin=gstin2,
            legal_name='Amazon Seller 2',
            trade_name='Store 2',
            state_code='29',
            state_name='Karnataka',
            financial_year='2024-25',
            filing_frequency='Monthly'
        )
        _db.session.add(p2)
        _db.session.commit()

        p1_id = p1.id
        p2_id = p2.id
        u1_id = u1.id
        u2_id = u2.id

    return {
        'user1_id': u1_id,
        'user2_id': u2_id,
        'profile1_id': p1_id,
        'profile2_id': p2_id,
        'gstin1': gstin1,
        'gstin2': gstin2,
    }


def _make_amazon_workbook(rows_data, sheet_name='MTR'):
    """Helper to build an in-memory workbook with Amazon headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        'Order ID', 'Invoice Number', 'Invoice Date', 'Ship From State',
        'Ship To State', 'Buyer GSTIN', 'Seller GSTIN', 'HSN/SAC',
        'Product Description', 'Quantity', 'Taxable Value', 'CGST Rate',
        'CGST Amount', 'SGST Rate', 'SGST Amount', 'IGST Rate',
        'IGST Amount', 'Cess Rate', 'Cess Amount', 'Invoice Value',
        'Tax Rate', 'Supply Type', 'Place of Supply', 'Reverse Charge'
    ]
    ws.append(headers)
    for row in rows_data:
        ws.append([row.get(h, '') for h in headers])
    return wb


# ==============================================================================
# 1. DETECTION TESTS (Requirements 1, 15)
# ==============================================================================

class TestAmazonDetection:
    """Deterministic Amazon MTR detection."""

    def test_01_valid_amazon_mtr_detection(self):
        """Valid Amazon MTR workbook is detected even with generic filename."""
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
        assert os.path.exists(fixture_path)

        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            adapter = detect_platform(wb, 'generic_download_9988.xlsx')
            assert adapter is not None
            assert adapter.PLATFORM_NAME == 'Amazon'
            assert isinstance(adapter, AmazonAdapter)
        finally:
            wb.close()

        # Pre-read filename hint check
        direct_adapter = AmazonAdapter()
        assert direct_adapter.detect(None, 'Amazon_MTR_January_2025.xlsx') is True
        assert direct_adapter.detect(None, 'unrelated_report.xlsx') is False

    def test_15_unknown_workbook_not_misclassified_as_amazon(self):
        """Unknown or non-Amazon workbooks must never be classified as Amazon."""
        # 1. Flipkart report
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        if os.path.exists(flipkart_path):
            wb_fk = openpyxl.load_workbook(flipkart_path, data_only=True)
            try:
                assert AmazonAdapter().detect(wb_fk, 'flipkart_sample.xlsx') is False
            finally:
                wb_fk.close()

        # 2. Meesho report
        meesho_path = os.path.join(FIXTURE_DIR, 'meesho_sample.xlsx')
        if os.path.exists(meesho_path):
            wb_me = openpyxl.load_workbook(meesho_path, data_only=True)
            try:
                assert AmazonAdapter().detect(wb_me, 'meesho_sample.xlsx') is False
            finally:
                wb_me.close()

        # 3. Random workbook with Amazon in filename but completely non-Amazon content
        wb_random = openpyxl.Workbook()
        ws = wb_random.active
        ws.title = 'Payroll'
        ws.append(['Employee ID', 'Employee Name', 'Department', 'Monthly Salary'])
        ws.append(['EMP-01', 'Alice', 'Engineering', 100000])
        try:
            # Filename says amazon, but content has no MTR sheet or Amazon headers
            assert AmazonAdapter().detect(wb_random, 'amazon_payroll_2025.xlsx') is False
        finally:
            wb_random.close()

        # 4. Sheet named MTR but with non-Amazon headers
        wb_fake_mtr = openpyxl.Workbook()
        ws_fake = wb_fake_mtr.active
        ws_fake.title = 'MTR'
        ws_fake.append(['Student ID', 'Grade', 'Course'])
        ws_fake.append(['S01', 'A', 'Math'])
        try:
            assert AmazonAdapter().detect(wb_fake_mtr, 'report.xlsx') is False
        finally:
            wb_fake_mtr.close()


# ==============================================================================
# 2. SCHEMA VALIDATION TESTS (Requirements 2, 3, 4)
# ==============================================================================

class TestAmazonSchemaValidation:
    """Explicit schema validation distinguishing headers, sheets, malformed files."""

    def test_02_invalid_amazon_header_rejection(self):
        """Amazon sheet with missing required headers is explicitly rejected."""
        adapter = AmazonAdapter()
        # Missing 'Invoice Number'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'MTR'
        ws.append(['Order ID', 'Taxable Value', 'Tax Rate'])
        ws.append(['ORD-001', 1500.0, 18.0])
        try:
            is_valid, errors = adapter.validate(wb)
            assert is_valid is False
            assert any('missing required column' in err.lower() for err in errors)
            assert any('invoice number' in err.lower() for err in errors)
        finally:
            wb.close()

    def test_03_wrong_sheet_rejection(self):
        """Excel workbook without an 'MTR' worksheet is explicitly rejected."""
        adapter = AmazonAdapter()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Orders'
        ws.append(['Order ID', 'Invoice Number', 'Taxable Value'])
        ws.append(['ORD-001', 'INV-001', 1000])
        try:
            is_valid, errors = adapter.validate(wb)
            assert is_valid is False
            assert any("expected worksheet named 'MTR'" in err or "recognizable header row" in err for err in errors)
        finally:
            wb.close()

    def test_04_malformed_workbook_rejection(self, tmp_path):
        """Malformed, corrupted, missing, or empty inputs return actionable errors without crashing."""
        adapter = AmazonAdapter()

        # 1. None input
        is_valid, errors = adapter.validate(None)
        assert is_valid is False
        assert errors == ['No workbook or data provided.']

        # 2. Missing file path
        is_valid, errors = adapter.validate('/non/existent/path/amazon.xlsx')
        assert is_valid is False
        assert any('file not found' in err.lower() for err in errors)

        # 3. Corrupted / unreadable file
        corrupt_file = tmp_path / 'corrupt.xlsx'
        corrupt_file.write_bytes(b'PK\x03\x04NOT_A_VALID_ZIP_ARCHIVE')
        is_valid, errors = adapter.validate(str(corrupt_file))
        assert is_valid is False
        assert any('unsupported or unreadable file' in err.lower() for err in errors)

        # 4. Unsupported file extension
        text_file = tmp_path / 'amazon.txt'
        text_file.write_text('Order ID,Invoice Number\nORD-1,INV-1\n', encoding='utf-8')
        is_valid, errors = adapter.validate(str(text_file))
        assert is_valid is False
        assert any('unsupported file type' in err.lower() for err in errors)

        # 5. Header row present but zero transaction data rows
        wb_empty = openpyxl.Workbook()
        ws = wb_empty.active
        ws.title = 'MTR'
        ws.append(['Order ID', 'Invoice Number', 'Taxable Value'])
        try:
            is_valid, errors = adapter.validate(wb_empty)
            assert is_valid is False
            assert any('no data rows found' in err.lower() for err in errors)
        finally:
            wb_empty.close()


# ==============================================================================
# 3. ROW PARSING, CANONICAL MAPPING & DECIMAL PRECISION (Requirements 5, 6, 7, 8)
# ==============================================================================

class TestAmazonRowParsingAndCanonical:
    """Row parsing into CanonicalTransaction, Decimal precision, error boundaries."""

    def test_05_valid_row_parsing_types(self):
        """Amazon adapter correctly parses Regular sales, Credit Notes, and Cancelled orders."""
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            adapter = AmazonAdapter()
            result = adapter.parse(wb, 'amazon_sample.xlsx')
        finally:
            wb.close()

        assert result.platform == 'Amazon'
        assert result.total_rows == 168
        assert result.success_rows == 168
        assert result.error_rows == 0

        # Verify Credit Notes
        credit_notes = [r for r in result.rows if r.canonical.note_type == 'CREDIT']
        assert len(credit_notes) == 15
        for cn_row in credit_notes:
            assert cn_row.canonical.supply_type == 'Credit Note'
            assert cn_row.canonical.invoice_number.startswith('CN-')
            assert cn_row.canonical.note_number == cn_row.canonical.invoice_number
            assert cn_row.canonical.taxable_value > Decimal('0')

        # Verify Cancelled orders
        cancelled = [r for r in result.rows if r.canonical.cancellation_flag]
        assert len(cancelled) == 3
        for c_row in cancelled:
            assert c_row.canonical.supply_type == 'Cancelled'
            assert c_row.canonical.return_flag is True
            assert c_row.status == ImportRowStatus.SUCCESS

        # Verify Regular sales
        regular = [r for r in result.rows if not r.canonical.note_type and not r.canonical.cancellation_flag]
        assert len(regular) == 150
        for reg in regular:
            assert reg.canonical.supply_type == 'Regular'
            assert reg.canonical.return_flag is False

    def test_06_malformed_row_handling(self):
        """Malformed rows (negative numbers, missing identity, bad dates) generate explicit row errors."""
        adapter = AmazonAdapter()

        rows = [
            # Row 1: Valid row
            {
                'Order ID': 'ORD-OK-01',
                'Invoice Number': 'INV-OK-01',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 1000.0,
                'CGST Rate': 9.0, 'CGST Amount': 90.0,
                'SGST Rate': 9.0, 'SGST Amount': 90.0,
                'Invoice Value': 1180.0,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            },
            # Row 2: Negative taxable value
            {
                'Order ID': 'ORD-BAD-02',
                'Invoice Number': 'INV-BAD-02',
                'Invoice Date': '15-01-2025',
                'Taxable Value': -500.0,
                'Invoice Value': -590.0,
                'Supply Type': 'Regular',
            },
            # Row 3: Missing invoice number
            {
                'Order ID': 'ORD-BAD-03',
                'Invoice Number': '',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 300.0,
                'Supply Type': 'Regular',
            },
            # Row 4: Unparseable date
            {
                'Order ID': 'ORD-BAD-04',
                'Invoice Number': 'INV-BAD-04',
                'Invoice Date': 'NOT_A_DATE_VALUE',
                'Taxable Value': 400.0,
                'Supply Type': 'Regular',
            },
            # Row 5: Malformed numeric string (warns and converts to 0, does not crash)
            {
                'Order ID': 'ORD-BAD-05',
                'Invoice Number': 'INV-BAD-05',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 'INVALID_AMOUNT',
                'Supply Type': 'Regular',
            },
        ]

        wb = _make_amazon_workbook(rows)
        try:
            result = adapter.parse(wb, 'test_malformed.xlsx')
        finally:
            wb.close()

        assert result.total_rows == 5

        # Row 1: Success
        assert result.rows[0].status == ImportRowStatus.SUCCESS

        # Row 2: Error due to negative taxable value
        assert result.rows[1].status == ImportRowStatus.ERROR
        assert any('negative' in err.lower() for err in result.rows[1].errors)

        # Row 3: Error due to missing invoice number
        assert result.rows[2].status == ImportRowStatus.ERROR
        assert any('missing invoice number' in err.lower() for err in result.rows[2].errors)

        # Row 4: Error due to unparseable date
        assert result.rows[3].status == ImportRowStatus.ERROR
        assert any('not a recognized date' in err.lower() for err in result.rows[3].errors)

        # Row 5: Error for non-numeric taxable value (defaults to 0, missing taxable value)
        assert result.rows[4].status == ImportRowStatus.ERROR
        assert any('not numeric' in w.lower() for w in result.rows[4].warnings)
        assert result.rows[4].canonical.taxable_value == Decimal('0')

    def test_07_decimal_precision(self):
        """Monetary and tax values parse to exact Decimal without float rounding corruption."""
        adapter = AmazonAdapter()
        rows = [
            {
                'Order ID': 'ORD-DEC-01',
                'Invoice Number': 'INV-DEC-01',
                'Invoice Date': '20-01-2025',
                'Taxable Value': 3424.26,
                'CGST Rate': 0, 'CGST Amount': 0,
                'SGST Rate': 0, 'SGST Amount': 0,
                'IGST Rate': 28, 'IGST Amount': 958.79,
                'Cess Rate': 0, 'Cess Amount': 0,
                'Invoice Value': 4383.05,
                'Tax Rate': 28,
                'Supply Type': 'Regular',
                'Place of Supply': '33-Tamil Nadu',
            }
        ]
        wb = _make_amazon_workbook(rows)
        try:
            result = adapter.parse(wb, 'dec_test.xlsx')
        finally:
            wb.close()

        row = result.rows[0]
        tx = row.canonical
        assert isinstance(tx.taxable_value, Decimal)
        assert isinstance(tx.igst_amount, Decimal)
        assert isinstance(tx.invoice_value, Decimal)
        assert tx.taxable_value == Decimal('3424.26')
        assert tx.igst_amount == Decimal('958.79')
        assert tx.invoice_value == Decimal('4383.05')
        assert tx.total_tax == Decimal('958.79')

    def test_08_canonical_mapping_contract(self):
        """Parsed rows provide both CanonicalTransaction and normalized dictionary representations."""
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            adapter = AmazonAdapter()
            result = adapter.parse(wb, 'amazon_sample.xlsx')
        finally:
            wb.close()

        first_row = result.rows[0]
        canonical = first_row.canonical

        # CanonicalTransaction instance check
        assert isinstance(canonical, CanonicalTransaction)
        assert canonical.invoice_number == 'INV-2025-0001'
        assert canonical.marketplace_name == 'Amazon'
        assert canonical.source_platform == 'Amazon'

        # Place of supply resolution
        assert canonical.place_of_supply.isdigit()
        assert len(canonical.place_of_supply) == 2
        assert canonical.place_of_supply_raw is not None

        # Mapping protocol compatibility
        assert canonical['invoice_number'] == 'INV-2025-0001'
        assert 'taxable_value' in canonical
        assert isinstance(canonical.to_dict(), dict)

        # Raw data separation
        assert 'Order ID' in first_row.raw_data
        assert 'Taxable Value' in first_row.raw_data


# ==============================================================================
# 4. DUPLICATES, PERSISTENCE & PIPELINE EXECUTION (Requirements 9, 10, 11, 12, 13, 14)
# ==============================================================================

class TestAmazonPipelinePersistence:
    """Full common pipeline execution with Amazon MTR, batched persistence and duplicate handling."""

    def test_09_in_file_duplicate_rows_marked_skipped(self, app, db, amazon_tenants, tmp_path):
        """Duplicate rows within the same Amazon MTR file are flagged and skipped."""
        p_id = amazon_tenants['profile1_id']
        u_id = amazon_tenants['user1_id']

        rows = [
            {
                'Order ID': 'ORD-DUP-01',
                'Invoice Number': 'INV-DUP-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 2000.0,
                'CGST Rate': 9.0, 'CGST Amount': 180.0,
                'SGST Rate': 9.0, 'SGST Amount': 180.0,
                'Invoice Value': 2360.0,
                'Supply Type': 'Regular',
                'Place of Supply': '27-Maharashtra',
            },
            # Exact duplicate of row 1
            {
                'Order ID': 'ORD-DUP-01',
                'Invoice Number': 'INV-DUP-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 2000.0,
                'CGST Rate': 9.0, 'CGST Amount': 180.0,
                'SGST Rate': 9.0, 'SGST Amount': 180.0,
                'Invoice Value': 2360.0,
                'Supply Type': 'Regular',
                'Place of Supply': '27-Maharashtra',
            }
        ]

        wb = _make_amazon_workbook(rows)
        test_file = str(tmp_path / 'amz_in_file_dup.xlsx')
        wb.save(test_file)
        wb.close()

        with app.app_context():
            res = process_import(
                file_path=test_file,
                profile_id=p_id,
                platform_name='Amazon',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
            )

            assert res.total_rows == 2
            assert res.success_rows == 1
            assert res.skipped_rows == 1
            assert res.error_rows == 0
            assert res.status == 'PARTIAL'

            # Exactly 1 transaction created in DB
            assert Transaction.query.filter_by(profile_id=p_id).count() == 1

    def test_10_full_file_reimport_handling(self, app, db, amazon_tenants):
        """Re-importing the same Amazon MTR file skips existing transactions and avoids double-counting."""
        p_id = amazon_tenants['profile1_id']
        u_id = amazon_tenants['user1_id']
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')

        with app.app_context():
            # Initial import
            res1 = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Amazon',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
            )
            assert res1.status == 'COMPLETED'
            assert res1.total_rows == 168
            assert res1.success_rows == 168
            assert res1.skipped_rows == 0
            assert Transaction.query.filter_by(profile_id=p_id).count() == 168

            # Re-import with allow_duplicate_file=True
            res2 = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Amazon',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
                allow_duplicate_file=True,
            )
            assert res2.status == 'PARTIAL'
            assert res2.total_rows == 168
            assert res2.success_rows == 0
            assert res2.skipped_rows == 168
            assert res2.error_rows == 0

            # Database row count strictly remains 168
            assert Transaction.query.filter_by(profile_id=p_id).count() == 168

    def test_11_import_history_result_recording(self, app, db, amazon_tenants):
        """ImportHistory correctly records processing metadata, status, and row metrics."""
        p_id = amazon_tenants['profile1_id']
        u_id = amazon_tenants['user1_id']
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')

        with app.app_context():
            res = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Amazon',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
            )

            ih = _db.session.get(ImportHistory, res.import_history_id)
            assert ih is not None
            assert ih.processing_status == 'COMPLETED'
            assert ih.platform_name == 'Amazon'
            assert ih.total_rows == 168
            assert ih.success_rows == 168
            assert ih.error_rows == 0
            assert ih.skipped_rows == 0
            assert ih.file_hash is not None
            assert ih.file_size > 0
            assert ih.processing_completed_at is not None

    def test_12_database_persistence_and_reconciliation(self, app, db, amazon_tenants):
        """Database transactions match fixture values down to the exact total taxable sum."""
        p_id = amazon_tenants['profile1_id']
        u_id = amazon_tenants['user1_id']
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')

        with app.app_context():
            res = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Amazon',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
            )
            assert res.status == 'COMPLETED'

            # Transactions created
            tx_list = Transaction.query.filter_by(profile_id=p_id).all()
            assert len(tx_list) == 168

            # Relationship-based RawImport linkage verified
            raw_imports = RawImport.query.filter_by(import_history_id=res.import_history_id).all()
            assert len(raw_imports) == 168
            for tx in tx_list:
                assert tx.raw_import is not None
                assert tx.raw_import_id is not None

            # Reconcile sum of taxable values against fixture
            persisted_sum = sum((tx.taxable_value for tx in tx_list), Decimal('0'))

            wb = openpyxl.load_workbook(fixture_path, data_only=True)
            ws = wb['MTR']
            headers = [c for c in next(ws.iter_rows(values_only=True))]
            tax_idx = headers.index('Taxable Value')
            fixture_sum = Decimal('0')
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[tax_idx] is not None:
                    fixture_sum += Decimal(str(row[tax_idx]))
            wb.close()

            assert persisted_sum == fixture_sum
            assert persisted_sum == Decimal('5548093.74')

    def test_13_rollback_on_injected_failure(self, app, db, amazon_tenants):
        """Injected failure during persistence triggers full rollback without orphaned rows."""
        p_id = amazon_tenants['profile1_id']
        u_id = amazon_tenants['user1_id']
        fixture_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')

        with app.app_context():
            with patch('app.extensions.db.session.flush', side_effect=RuntimeError('Injected DB Flush Error')):
                with pytest.raises(RuntimeError):
                    process_import(
                        file_path=fixture_path,
                        profile_id=p_id,
                        platform_name='Amazon',
                        user_id=u_id,
                        return_period='012025',
                        financial_year='2024-25',
                    )

            # Assert zero orphaned transactions or raw imports
            assert Transaction.query.filter_by(profile_id=p_id).count() == 0
            ih = ImportHistory.query.filter_by(profile_id=p_id).order_by(ImportHistory.id.desc()).first()
            assert ih is not None
            assert ih.processing_status == 'FAILED'
            assert RawImport.query.filter_by(import_history_id=ih.id).count() == 0

    def test_14_tenant_isolation(self, app, db, amazon_tenants, tmp_path):
        """Transactions imported for Profile A are strictly isolated from Profile B."""
        p1_id = amazon_tenants['profile1_id']
        u1_id = amazon_tenants['user1_id']
        p2_id = amazon_tenants['profile2_id']
        u2_id = amazon_tenants['user2_id']

        rows_p1 = [
            {
                'Order ID': 'ORD-P1-01',
                'Invoice Number': 'INV-P1-01',
                'Invoice Date': '12-01-2025',
                'Taxable Value': 1000.0,
                'Invoice Value': 1180.0,
                'Supply Type': 'Regular',
                'Place of Supply': '27-Maharashtra',
            }
        ]
        wb1 = _make_amazon_workbook(rows_p1)
        file_p1 = str(tmp_path / 'amz_p1.xlsx')
        wb1.save(file_p1)
        wb1.close()

        rows_p2 = [
            {
                'Order ID': 'ORD-P2-01',
                'Invoice Number': 'INV-P2-01',
                'Invoice Date': '14-01-2025',
                'Taxable Value': 2500.0,
                'Invoice Value': 2950.0,
                'Supply Type': 'Regular',
                'Place of Supply': '29-Karnataka',
            }
        ]
        wb2 = _make_amazon_workbook(rows_p2)
        file_p2 = str(tmp_path / 'amz_p2.xlsx')
        wb2.save(file_p2)
        wb2.close()

        with app.app_context():
            res1 = process_import(file_p1, p1_id, 'Amazon', u1_id, '012025', '2024-25')
            res2 = process_import(file_p2, p2_id, 'Amazon', u2_id, '012025', '2024-25')

            assert res1.status == 'COMPLETED'
            assert res2.status == 'COMPLETED'

            txs_p1 = Transaction.query.filter_by(profile_id=p1_id).all()
            txs_p2 = Transaction.query.filter_by(profile_id=p2_id).all()

            assert len(txs_p1) == 1
            assert len(txs_p2) == 1
            assert txs_p1[0].invoice_number == 'INV-P1-01'
            assert txs_p2[0].invoice_number == 'INV-P2-01'
            assert txs_p1[0].profile_id == p1_id
            assert txs_p2[0].profile_id == p2_id


# ==============================================================================
# 5. REGISTRY DETERMINISM TESTS (Phase 5B Follow-Up)
# ==============================================================================

class TestAmazonRegistryDeterminism:
    """Explicit determinism and precedence for Amazon adapter registration."""

    def test_get_adapter_amazon_resolves_to_specialized_module(self):
        """get_adapter('Amazon') must always resolve to app.adapters.amazon.AmazonAdapter."""
        adapter = get_adapter('Amazon')
        assert adapter is not None
        assert adapter.__class__.__module__ == 'app.adapters.amazon'
        assert adapter.__class__.__name__ == 'AmazonAdapter'

    def test_generic_adapter_registration_cannot_overwrite_specialized_adapter(self):
        """Generic all_adapters.AmazonAdapter must never overwrite the specialized adapter."""
        from app.adapters.all_adapters import AmazonAdapter as GenericAmazonAdapter
        from app.adapters.amazon import AmazonAdapter as SpecializedAmazonAdapter

        try:
            # Attempt to register generic legacy adapter
            register_adapter(GenericAmazonAdapter)
            resolved = get_adapter('Amazon')
            assert resolved.__class__.__module__ == 'app.adapters.amazon'
            assert resolved.__class__ is SpecializedAmazonAdapter
            assert _ADAPTER_REGISTRY['Amazon'] is SpecializedAmazonAdapter
        finally:
            auto_register_adapters()

    def test_registration_order_independence(self):
        """Precedence ensures specialized adapter wins regardless of registration order."""
        from app.adapters.all_adapters import AmazonAdapter as GenericAmazonAdapter
        from app.adapters.amazon import AmazonAdapter as SpecializedAmazonAdapter

        original = dict(_ADAPTER_REGISTRY)
        try:
            # Case A: Generic registered first, specialized registered second
            _ADAPTER_REGISTRY.clear()
            register_adapter(GenericAmazonAdapter)
            assert _ADAPTER_REGISTRY['Amazon'] is GenericAmazonAdapter
            register_adapter(SpecializedAmazonAdapter)
            assert _ADAPTER_REGISTRY['Amazon'] is SpecializedAmazonAdapter
            assert get_adapter('Amazon').__class__.__module__ == 'app.adapters.amazon'

            # Case B: Specialized registered first, generic registered second
            _ADAPTER_REGISTRY.clear()
            register_adapter(SpecializedAmazonAdapter)
            assert _ADAPTER_REGISTRY['Amazon'] is SpecializedAmazonAdapter
            register_adapter(GenericAmazonAdapter)
            assert _ADAPTER_REGISTRY['Amazon'] is SpecializedAmazonAdapter
            assert get_adapter('Amazon').__class__.__module__ == 'app.adapters.amazon'
        finally:
            _ADAPTER_REGISTRY.clear()
            _ADAPTER_REGISTRY.update(original)
            auto_register_adapters()

    def test_auto_register_idempotency_and_all_adapters_instantiate(self):
        """auto_register_adapters is idempotent and all platforms instantiate properly."""
        auto_register_adapters()
        assert get_adapter('Amazon').__class__.__module__ == 'app.adapters.amazon'
        assert get_adapter('Generic').__class__.__module__ == 'app.adapters.all_adapters'
        assert get_adapter('Generic').__class__.__name__ == '_GenericFilenameAdapter'

        auto_register_adapters()
        assert get_adapter('Amazon').__class__.__module__ == 'app.adapters.amazon'
        assert get_adapter('Generic').__class__.__module__ == 'app.adapters.all_adapters'
        assert get_adapter('Generic').__class__.__name__ == '_GenericFilenameAdapter'

        # Verify all registered adapters instantiate and match PLATFORM_NAME
        for platform_name in list(_ADAPTER_REGISTRY.keys()):
            adapter = get_adapter(platform_name)
            assert adapter is not None, f"Failed to instantiate {platform_name}"
            assert adapter.PLATFORM_NAME == platform_name

    def test_get_adapter_generic_resolves_to_generic_filename_adapter(self):
        """get_adapter('Generic') must resolve to app.adapters.all_adapters._GenericFilenameAdapter."""
        from app.adapters.all_adapters import _GenericFilenameAdapter
        adapter = get_adapter('Generic')
        assert adapter is not None
        assert adapter.__class__.__module__ == 'app.adapters.all_adapters'
        assert adapter.__class__.__name__ == '_GenericFilenameAdapter'
        assert adapter.__class__ is _GenericFilenameAdapter

    def test_base_generic_adapter_cannot_overwrite_generic_filename_adapter(self):
        """BaseGenericAdapter must never overwrite the more specialized _GenericFilenameAdapter."""
        from app.adapters.all_adapters import BaseGenericAdapter, _GenericFilenameAdapter

        try:
            register_adapter(BaseGenericAdapter)
            resolved = get_adapter('Generic')
            assert resolved.__class__ is _GenericFilenameAdapter
            assert _ADAPTER_REGISTRY['Generic'] is _GenericFilenameAdapter
        finally:
            auto_register_adapters()

    def test_generic_registration_order_independence(self):
        """Subclass precedence ensures _GenericFilenameAdapter wins regardless of registration order."""
        from app.adapters.all_adapters import BaseGenericAdapter, _GenericFilenameAdapter

        original = dict(_ADAPTER_REGISTRY)
        try:
            # Case A: BaseGenericAdapter registered first, _GenericFilenameAdapter registered second
            _ADAPTER_REGISTRY.clear()
            register_adapter(BaseGenericAdapter)
            assert _ADAPTER_REGISTRY['Generic'] is BaseGenericAdapter
            register_adapter(_GenericFilenameAdapter)
            assert _ADAPTER_REGISTRY['Generic'] is _GenericFilenameAdapter
            assert get_adapter('Generic').__class__ is _GenericFilenameAdapter

            # Case B: _GenericFilenameAdapter registered first, BaseGenericAdapter registered second
            _ADAPTER_REGISTRY.clear()
            register_adapter(_GenericFilenameAdapter)
            assert _ADAPTER_REGISTRY['Generic'] is _GenericFilenameAdapter
            register_adapter(BaseGenericAdapter)
            assert _ADAPTER_REGISTRY['Generic'] is _GenericFilenameAdapter
            assert get_adapter('Generic').__class__ is _GenericFilenameAdapter
        finally:
            _ADAPTER_REGISTRY.clear()
            _ADAPTER_REGISTRY.update(original)
            auto_register_adapters()
