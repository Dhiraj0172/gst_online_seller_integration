"""Phase 5C-3 focused tests for Myntra Marketplace Adapter Production Implementation.

Covers all required areas:
A. Deterministic detection
B. Explicit Myntra sheet ("Myntra GST Report")
C. Random filename with explicit Myntra sheet
D. Flipkart GST Report renamed myntra.xlsx (claimed by Flipkart, not Myntra)
E. Amazon MTR renamed myntra.xlsx (claimed by Amazon, not Myntra)
F. Meesho Orders renamed myntra.xlsx (claimed by Meesho, not Myntra)
G. Generic workbook renamed myntra.xlsx
H. Malformed Myntra headers
I. Empty workbook
J. Canonical field mapping
K. Invoice and date parsing
L. Return period derivation
M. GSTIN mapping
N. Place of Supply (POS) mapping
O. HSN/SAC mapping
P. Quantity mapping
Q. Taxable value mapping
R. CGST/SGST/IGST/Cess rates and amounts
S. Invoice value mapping
T. Supply classification
U. Credit notes and customer returns
V. E-commerce operator GSTIN
W. Duplicate detection and re-import behavior
X. Real database persistence
Y. Tenant isolation
Z. Malformed required numeric failures
AA. Mandatory >1000-row rollback crossing real batch flush boundary
AB. Registry determinism and precedence
AC. CSV limitation / safe behavior
"""
import io
import json
import os
import uuid
from collections import Counter
from decimal import Decimal
from unittest.mock import patch

import openpyxl
import pytest

from app.adapters.base import ImportRowStatus
from app.adapters.canonical import CanonicalTransaction
from app.adapters.flipkart import FlipkartAdapter
from app.adapters.myntra import MyntraAdapter
from app.adapters.registry import (
    _ADAPTER_REGISTRY,
    auto_register_adapters,
    detect_platform,
    get_adapter,
    register_adapter,
)
from app.extensions import db as _db
from app.models import GSTProfile, ImportHistory, RawImport, Transaction, User
from app.services.import_service import process_import

FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')


@pytest.fixture
def myntra_tenants(app, db):
    """Provide two distinct users and GST profiles for tenant isolation tests."""
    suffix = uuid.uuid4().hex[:8]
    u1_name = f'myntra_user1_{suffix}'
    u2_name = f'myntra_user2_{suffix}'
    u1_email = f'myntra1_{suffix}@example.com'
    u2_email = f'myntra2_{suffix}@example.com'
    gstin1 = f'27AABC{suffix[:4].upper()}1Z5'
    gstin2 = f'29BBED{suffix[:4].upper()}1Z6'

    with app.app_context():
        u1 = User(username=u1_name, email=u1_email)
        u1.set_password('TenantPass1!')
        u2 = User(username=u2_name, email=u2_email)
        u2.set_password('TenantPass2!')
        _db.session.add(u1)
        _db.session.add(u2)
        _db.session.commit()

        p1 = GSTProfile(
            user_id=u1.id,
            gstin=gstin1,
            legal_name='Tenant 1 Seller',
            state_code='27',
            state_name='Maharashtra',
            financial_year='2024-25',
            filing_frequency='Monthly',
        )
        p2 = GSTProfile(
            user_id=u2.id,
            gstin=gstin2,
            legal_name='Tenant 2 Seller',
            state_code='29',
            state_name='Karnataka',
            financial_year='2024-25',
            filing_frequency='Monthly',
        )
        _db.session.add(p1)
        _db.session.add(p2)
        _db.session.commit()

        p1_id = p1.id
        p2_id = p2.id
        u1_id = u1.id
        u2_id = u2.id

    return {
        'user1_id': u1_id,
        'user2_id': u2_id,
        'profile1_id': p1_id,
        'profile2_id': p2_id,
        'gstin1': gstin1,
        'gstin2': gstin2,
    }


def _make_myntra_workbook(rows_data, sheet_name='Myntra GST Report', extra_headers=None):
    """Helper to build an in-memory workbook with Myntra/Flipkart headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        'Order ID', 'Invoice Number', 'Invoice Date', 'Ship From State',
        'Ship To State', 'Buyer GSTIN', 'Seller GSTIN', 'HSN/SAC',
        'Quantity', 'Taxable Value', 'CGST Rate', 'CGST Amount',
        'SGST Rate', 'SGST Amount', 'IGST Rate', 'IGST Amount',
        'Cess Rate', 'Cess Amount', 'Invoice Value', 'Tax Rate',
        'Supply Type', 'Place of Supply', 'Reverse Charge',
        'Order Item ID', 'Product Title', 'Selling Price'
    ]
    if extra_headers:
        for eh in extra_headers:
            if eh not in headers:
                headers.append(eh)

    ws.append(headers)
    for r in rows_data:
        row_vals = [r.get(h, '') for h in headers]
        ws.append(row_vals)
    return wb


# ===========================================================================
# 1. Detection Contract Tests
# ===========================================================================
class TestMyntraDetection:
    def test_01_detect_explicit_myntra_sheet_and_signature_headers(self):
        """Explicit Myntra sheet with signature headers detects as Myntra."""
        rows = [{
            'Order ID': 'ORD-1', 'Order Item ID': 'OI-1', 'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025', 'Product Title': 'T-Shirt', 'Selling Price': 500,
            'Taxable Value': 500, 'Invoice Value': 590, 'Place of Supply': '27-Maharashtra',
            'Supply Type': 'Regular',
        }]
        wb = _make_myntra_workbook(rows, sheet_name='Myntra GST Report')
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'myntra_sales.xlsx') is True

    def test_02_detect_random_filename_with_explicit_myntra_sheet(self):
        """Explicit Myntra sheet with random filename detects as Myntra."""
        rows = [{
            'Order ID': 'ORD-1', 'Order Item ID': 'OI-1', 'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025', 'Product Title': 'T-Shirt', 'Selling Price': 500,
            'Taxable Value': 500, 'Invoice Value': 590, 'Place of Supply': '27-Maharashtra',
            'Supply Type': 'Regular',
        }]
        wb = _make_myntra_workbook(rows, sheet_name='Myntra GST Report')
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'random_export_9823.xlsx') is True

    def test_03_detect_flipkart_gst_report_renamed_myntra_rejected(self):
        """Flipkart GST Report renamed to myntra.xlsx is REJECTED by Myntra and claimed by Flipkart."""
        auto_register_adapters()
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            adapter = MyntraAdapter()
            assert adapter.detect(wb, 'myntra.xlsx') is False

            detected = detect_platform(wb, 'myntra.xlsx')
            assert detected is not None
            assert detected.PLATFORM_NAME == 'Flipkart'
        finally:
            wb.close()

    def test_04_detect_amazon_mtr_renamed_myntra_rejected(self):
        """Amazon MTR file renamed to myntra.xlsx is REJECTED by Myntra and claimed by Amazon."""
        auto_register_adapters()
        amazon_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
        wb = openpyxl.load_workbook(amazon_path, data_only=True)
        try:
            adapter = MyntraAdapter()
            assert adapter.detect(wb, 'myntra.xlsx') is False

            detected = detect_platform(wb, 'myntra.xlsx')
            assert detected is not None
            assert detected.PLATFORM_NAME == 'Amazon'
        finally:
            wb.close()

    def test_05_detect_meesho_orders_renamed_myntra_rejected(self):
        """Meesho Orders file renamed to myntra.xlsx is REJECTED by Myntra and claimed by Meesho."""
        auto_register_adapters()
        meesho_path = os.path.join(FIXTURE_DIR, 'meesho_sample.xlsx')
        wb = openpyxl.load_workbook(meesho_path, data_only=True)
        try:
            adapter = MyntraAdapter()
            assert adapter.detect(wb, 'myntra.xlsx') is False

            detected = detect_platform(wb, 'myntra.xlsx')
            assert detected is not None
            assert detected.PLATFORM_NAME == 'Meesho'
        finally:
            wb.close()

    def test_06_detect_generic_workbook_renamed_myntra_rejected(self):
        """Generic workbook with myntra.xlsx filename is rejected."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws.append(['Col A', 'Col B', 'Col C'])
        ws.append([1, 2, 3])
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'myntra.xlsx') is False

    def test_07_detect_empty_workbook_renamed_myntra_rejected(self):
        """Empty workbook renamed to myntra.xlsx is rejected."""
        wb = openpyxl.Workbook()
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'myntra.xlsx') is False

    def test_08_detect_myntra_sheet_with_malformed_headers_rejected(self):
        """Workbook with 'Myntra GST Report' sheet but unrelated/missing headers is rejected."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Myntra GST Report'
        ws.append(['Employee ID', 'Department', 'Salary'])
        ws.append([101, 'Engineering', 50000])
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'myntra.xlsx') is False

    def test_09_detect_wrong_sheet_with_myntra_headers_rejected(self):
        """Workbook with Myntra headers but wrong sheet name is rejected by Myntra."""
        rows = [{
            'Order ID': 'ORD-1', 'Order Item ID': 'OI-1', 'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025', 'Product Title': 'T-Shirt', 'Selling Price': 500,
            'Taxable Value': 500, 'Invoice Value': 590, 'Place of Supply': '27-Maharashtra',
            'Supply Type': 'Regular',
        }]
        wb = _make_myntra_workbook(rows, sheet_name='Other Data')
        adapter = MyntraAdapter()
        assert adapter.detect(wb, 'myntra.xlsx') is False

    def test_10_detect_pre_read_filename_token(self):
        """When workbook_or_data is None, filename token acts as pre-read hint."""
        adapter = MyntraAdapter()
        assert adapter.detect(None, 'myntra_report_jan.xlsx') is True
        assert adapter.detect(None, 'flipkart_report_jan.xlsx') is False

    def test_11_detect_csv_limitation_safe_rejection(self, tmp_path):
        """CSV files lack worksheet provenance and are safely NOT detected as Myntra."""
        csv_path = tmp_path / 'myntra_export.csv'
        csv_path.write_text(
            'Order ID,Invoice Number,Invoice Date,Taxable Value,Product Title,Selling Price\n'
            'ORD-1,INV-1,15-01-2025,1000,T-Shirt,1000\n',
            encoding='utf-8',
        )
        adapter = MyntraAdapter()
        assert adapter.detect(str(csv_path), 'myntra_export.csv') is False


# ===========================================================================
# 2. Schema Validation Tests
# ===========================================================================
class TestMyntraValidation:
    def test_12_validate_none_or_missing_input(self):
        """Validation fails explicitly on None input."""
        adapter = MyntraAdapter()
        valid, errors = adapter.validate(None)
        assert valid is False
        assert any('No workbook or data provided' in err for err in errors)

    def test_13_validate_empty_workbook(self):
        """Validation fails on empty workbook without worksheets."""
        adapter = MyntraAdapter()
        valid, errors = adapter.validate([])
        assert valid is False
        assert any('No data rows' in err or 'contains no worksheets' in err for err in errors)

    def test_14_validate_wrong_headers_fails_explicitly(self):
        """Validation fails when required column layout is not recognized."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Myntra GST Report'
        ws.append(['Col A', 'Col B', 'Col C'])
        ws.append([1, 2, 3])
        adapter = MyntraAdapter()
        valid, errors = adapter.validate(wb)
        assert valid is False
        assert any('Unrecognized Myntra file' in err for err in errors)

    def test_15_validate_missing_invoice_number_fails(self):
        """Validation fails when required invoice_number column is missing."""
        rows = [{
            'Order ID': 'ORD-1', 'Order Item ID': 'OI-1',
            'Invoice Date': '15-01-2025', 'Product Title': 'T-Shirt', 'Selling Price': 500,
            'Taxable Value': 500, 'Invoice Value': 590, 'Place of Supply': '27-Maharashtra',
            'Supply Type': 'Regular',
        }]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Myntra GST Report'
        # Headers without Invoice Number
        headers = [
            'Order ID', 'Order Item ID', 'Invoice Date', 'Product Title',
            'Selling Price', 'Taxable Value', 'Invoice Value', 'Place of Supply', 'Supply Type'
        ]
        ws.append(headers)
        ws.append([rows[0].get(h, '') for h in headers])

        adapter = MyntraAdapter()
        valid, errors = adapter.validate(wb)
        assert valid is False
        assert any('missing required column' in err and 'invoice number' in err.lower() for err in errors)

    def test_16_validate_missing_taxable_value_fails(self):
        """Validation fails when required taxable_value column is missing."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Myntra GST Report'
        headers = [
            'Order ID', 'Order Item ID', 'Invoice Number', 'Invoice Date',
            'Product Title', 'Selling Price', 'Invoice Value', 'Place of Supply', 'Supply Type'
        ]
        ws.append(headers)
        ws.append(['ORD-1', 'OI-1', 'INV-1', '15-01-2025', 'T-Shirt', 500, 590, '27', 'Regular'])

        adapter = MyntraAdapter()
        valid, errors = adapter.validate(wb)
        assert valid is False
        assert any('missing required column' in err and 'taxable value' in err.lower() for err in errors)

    def test_17_validate_header_only_zero_data_rows_fails(self):
        """Validation fails if headers are present but there are 0 data rows."""
        wb = _make_myntra_workbook([], sheet_name='Myntra GST Report')
        adapter = MyntraAdapter()
        valid, errors = adapter.validate(wb)
        assert valid is False
        assert any('No data rows found beneath the header row' in err for err in errors)

    def test_18_validate_valid_myntra_sheet_succeeds(self):
        """Valid Myntra sheet with rows passes validation cleanly."""
        rows = [{
            'Order ID': 'ORD-1', 'Order Item ID': 'OI-1', 'Invoice Number': 'INV-1',
            'Invoice Date': '15-01-2025', 'Product Title': 'T-Shirt', 'Selling Price': 500,
            'Taxable Value': 500, 'Invoice Value': 590, 'Place of Supply': '27-Maharashtra',
            'Supply Type': 'Regular',
        }]
        wb = _make_myntra_workbook(rows, sheet_name='Myntra GST Report')
        adapter = MyntraAdapter()
        valid, errors = adapter.validate(wb)
        assert valid is True
        assert errors == []

    def test_19_validate_shared_gst_report_sheet_succeeds(self):
        """Shared 'GST Report' sheet layout passes validation when evaluated by MyntraAdapter."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            adapter = MyntraAdapter()
            valid, errors = adapter.validate(wb)
            assert valid is True
            assert errors == []
        finally:
            wb.close()


# ===========================================================================
# 3. Parsing, Financial Reconciliation, and Canonical Mapping
# ===========================================================================
class TestMyntraParsingAndReconciliation:
    def test_20_parse_shared_fixture_totals_and_paise_reconciliation(self):
        """Parse shared fixture and verify exact financial totals down to the paise."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            adapter = MyntraAdapter()
            result = adapter.parse(wb, 'myntra_sample.xlsx')
        finally:
            wb.close()

        assert result.platform == 'Myntra'
        assert result.total_rows == 118
        assert result.success_rows == 118
        assert result.error_rows == 0
        assert result.warning_rows == 0

        taxable_sum = sum(r.normalized_data['taxable_value'] for r in result.rows)
        cgst_sum = sum(r.normalized_data['cgst_amount'] for r in result.rows)
        sgst_sum = sum(r.normalized_data['sgst_amount'] for r in result.rows)
        igst_sum = sum(r.normalized_data['igst_amount'] for r in result.rows)
        cess_sum = sum(r.normalized_data['cess_amount'] for r in result.rows)
        inv_sum = sum(r.normalized_data['invoice_value'] for r in result.rows)

        # Exact expected totals
        assert taxable_sum == Decimal('3363284.64')
        assert cgst_sum == Decimal('68093.44')
        assert sgst_sum == Decimal('68093.44')
        assert igst_sum == Decimal('391086.23')
        assert cess_sum == Decimal('0.00')
        assert inv_sum == Decimal('3890557.75')

        # Reconciliation: Taxable + Tax == Invoice
        total_tax = cgst_sum + sgst_sum + igst_sum + cess_sum
        assert taxable_sum + total_tax == inv_sum

    def test_21_parse_b2b_b2c_breakdown(self):
        """Verify 35 B2B and 83 B2C transactions parsed correctly."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            result = MyntraAdapter().parse(wb)
        finally:
            wb.close()

        b2b_rows = [r for r in result.rows if r.normalized_data.get('customer_gstin')]
        b2c_rows = [r for r in result.rows if not r.normalized_data.get('customer_gstin')]

        assert len(b2b_rows) == 35
        assert len(b2c_rows) == 83

    def test_22_parse_credit_notes_and_returns(self):
        """Verify 5 credit notes and 3 returns are identified with correct flags."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            result = MyntraAdapter().parse(wb)
        finally:
            wb.close()

        cns = [r for r in result.rows if r.normalized_data.get('note_type') == 'CREDIT']
        returns = [r for r in result.rows if r.normalized_data.get('return_flag') is True]

        assert len(cns) == 5
        assert len(returns) == 3
        for cn in cns:
            assert cn.normalized_data['note_number'] != ''
            assert cn.normalized_data['taxable_value'] > 0

    def test_23_parse_dates_and_return_period_derivation(self):
        """Verify invoice dates parsed to DD-MM-YYYY and return period '012025' detected."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            result = MyntraAdapter().parse(wb)
        finally:
            wb.close()

        assert result.detected_period == '012025'
        assert result.detected_gstin == '27AABCU9603R1ZM'
        for r in result.rows:
            inv_date = r.normalized_data['invoice_date']
            assert len(inv_date) == 10
            assert inv_date[2] == '-' and inv_date[5] == '-'

    def test_24_parse_pos_normalization(self):
        """Place of supply is normalized to 2-digit GST state code."""
        rows = [
            {'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1', 'Taxable Value': 100, 'Place of Supply': '27-Maharashtra'},
            {'Order ID': 'O2', 'Order Item ID': 'OI2', 'Invoice Number': 'I2', 'Taxable Value': 100, 'Place of Supply': 'Karnataka'},
            {'Order ID': 'O3', 'Order Item ID': 'OI3', 'Invoice Number': 'I3', 'Taxable Value': 100, 'Place of Supply': '07'},
        ]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.rows[0].normalized_data['place_of_supply'] == '27'
        assert result.rows[1].normalized_data['place_of_supply'] == '29'
        assert result.rows[2].normalized_data['place_of_supply'] == '07'

    def test_25_parse_hsn_sac_and_quantities(self):
        """HSN/SAC and quantity are mapped correctly as Decimal."""
        rows = [{
            'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1',
            'HSN/SAC': '6109', 'Quantity': '3', 'Selling Price': '250.00', 'Taxable Value': '750.00'
        }]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        r = result.rows[0].normalized_data
        assert r['hsn_sac'] == '6109'
        assert r['quantity'] == Decimal('3')
        assert r['unit_price'] == Decimal('250.00')
        assert r['taxable_value'] == Decimal('750.00')

    def test_26_parse_ecommerce_gstin(self):
        """E-commerce GSTIN column is extracted and normalized."""
        rows = [{
            'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1',
            'Taxable Value': 500, 'E-Commerce GSTIN': '29AABCF5678D1ZP'
        }]
        wb = _make_myntra_workbook(rows, extra_headers=['E-Commerce GSTIN'])
        result = MyntraAdapter().parse(wb)
        assert result.rows[0].normalized_data['ecommerce_gstin'] == '29AABCF5678D1ZP'

    def test_27_parse_canonical_transaction_contract(self):
        """Every parsed row instantiates a valid CanonicalTransaction."""
        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        wb = openpyxl.load_workbook(flipkart_path, data_only=True)
        try:
            result = MyntraAdapter().parse(wb)
        finally:
            wb.close()

        for r in result.rows:
            assert isinstance(r.canonical, CanonicalTransaction)
            assert r.canonical.marketplace_name == 'Myntra'
            assert r.canonical.source_platform == 'Myntra'
            assert r.canonical.order_item_id is not None
            assert r.canonical.invoice_number is not None

    def test_28_parse_tax_rate_slabs(self):
        """Tax rates and components parse accurately for standard slabs."""
        rows = [
            {'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1', 'Taxable Value': 1000, 'Tax Rate': 5, 'CGST Rate': 2.5, 'SGST Rate': 2.5, 'CGST Amount': 25, 'SGST Amount': 25},
            {'Order ID': 'O2', 'Order Item ID': 'OI2', 'Invoice Number': 'I2', 'Taxable Value': 1000, 'Tax Rate': 18, 'IGST Rate': 18, 'IGST Amount': 180},
        ]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.rows[0].normalized_data['cgst_amount'] == Decimal('25')
        assert result.rows[0].normalized_data['sgst_amount'] == Decimal('25')
        assert result.rows[1].normalized_data['igst_amount'] == Decimal('180')


# ===========================================================================
# 4. Error Handling and Malformed Data
# ===========================================================================
class TestMyntraErrorHandling:
    def test_29_malformed_taxable_numeric_produces_error(self):
        """Unparseable taxable value produces explicit row ERROR and is not coerced to zero."""
        rows = [
            {'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1', 'Taxable Value': 'VALID_100'},
        ]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.error_rows == 1
        r = result.rows[0]
        assert r.status == ImportRowStatus.ERROR
        assert any("Taxable value 'VALID_100' is not a valid numeric amount" in err for err in r.errors)

    def test_30_malformed_other_required_financial_numerics(self):
        """Malformed quantity or invoice value produces row ERRORs."""
        rows = [
            {'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1', 'Taxable Value': 100, 'Quantity': 'ten_pieces'},
            {'Order ID': 'O2', 'Order Item ID': 'OI2', 'Invoice Number': 'I2', 'Taxable Value': 100, 'Invoice Value': 'bad_val'},
        ]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.rows[0].status == ImportRowStatus.ERROR
        assert any("Quantity 'ten_pieces' is not a valid numeric amount" in err for err in result.rows[0].errors)
        assert result.rows[1].status == ImportRowStatus.ERROR
        assert any("Invoice value 'bad_val' is not a valid numeric amount" in err for err in result.rows[1].errors)

    def test_31_negative_values_produce_error(self):
        """Negative amounts are flagged as ERRORs per import contract direction conventions."""
        rows = [{
            'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1', 'Taxable Value': -500,
        }]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.error_rows == 1
        assert any('negative' in err.lower() for err in result.rows[0].errors)

    def test_32_unparseable_date_produces_error(self):
        """Unparseable date produces row ERROR."""
        rows = [{
            'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1',
            'Invoice Date': '99-99-9999', 'Taxable Value': 500
        }]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.error_rows == 1
        assert any('not a recognized date' in err for err in result.rows[0].errors)

    def test_33_malformed_gstin_kept_with_warning(self):
        """Malformed customer GSTIN records WARNING but is preserved without dropping."""
        rows = [{
            'Order ID': 'O1', 'Order Item ID': 'OI1', 'Invoice Number': 'I1',
            'Buyer GSTIN': 'INVALID_GSTIN_123', 'Taxable Value': 500
        }]
        wb = _make_myntra_workbook(rows)
        result = MyntraAdapter().parse(wb)
        assert result.rows[0].status == ImportRowStatus.WARNING
        assert result.rows[0].normalized_data['customer_gstin'] == 'INVALID_GSTIN_123'
        assert any('rejected by validation' in w for w in result.rows[0].warnings)


# ===========================================================================
# 5. Persistence, Duplicates, Tenant Isolation, and Rollback
# ===========================================================================
class TestMyntraDuplicatesAndPersistence:
    def test_34_in_file_duplicate_rows_marked_skipped(self, app, db, myntra_tenants):
        """Exact duplicate line item inside the same file is marked SKIPPED."""
        p1_id = myntra_tenants['profile1_id']
        u1_id = myntra_tenants['user1_id']

        rows = [
            {'Order ID': 'O1', 'Order Item ID': 'OI-DUP', 'Invoice Number': 'I1', 'Taxable Value': 100, 'Invoice Date': '15-01-2025', 'Product Title': 'Item A', 'Selling Price': 100, 'Invoice Value': 118, 'Place of Supply': '27-Maharashtra', 'Supply Type': 'Regular'},
            {'Order ID': 'O1', 'Order Item ID': 'OI-DUP', 'Invoice Number': 'I1', 'Taxable Value': 100, 'Invoice Date': '15-01-2025', 'Product Title': 'Item A', 'Selling Price': 100, 'Invoice Value': 118, 'Place of Supply': '27-Maharashtra', 'Supply Type': 'Regular'},
        ]
        wb = _make_myntra_workbook(rows)
        path = os.path.join(FIXTURE_DIR, 'test_myntra_in_file_dup.xlsx')
        wb.save(path)
        try:
            with app.app_context():
                res = process_import(path, p1_id, 'Myntra', u1_id, '012025', '2024-25', allow_duplicate_file=True)
                assert res.status == 'PARTIAL'
                assert res.total_rows == 2
                assert res.success_rows == 1
                assert res.skipped_rows == 1
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_35_full_file_duplicate_reimport(self, app, db, myntra_tenants):
        """Re-importing the same file without allow_duplicate_file is rejected as REJECTED_DUPLICATE_FILE."""
        p1_id = myntra_tenants['profile1_id']
        u1_id = myntra_tenants['user1_id']

        rows = [{'Order ID': 'O-RE', 'Order Item ID': 'OI-RE', 'Invoice Number': 'INV-RE', 'Taxable Value': 200, 'Invoice Date': '15-01-2025', 'Product Title': 'Item Re', 'Selling Price': 200, 'Invoice Value': 236, 'Place of Supply': '27-Maharashtra', 'Supply Type': 'Regular'}]
        wb = _make_myntra_workbook(rows)
        path = os.path.join(FIXTURE_DIR, 'test_myntra_reimport.xlsx')
        wb.save(path)
        try:
            with app.app_context():
                res1 = process_import(path, p1_id, 'Myntra', u1_id, '012025', '2024-25', allow_duplicate_file=False)
                assert res1.status == 'COMPLETED'

                res2 = process_import(path, p1_id, 'Myntra', u1_id, '012025', '2024-25', allow_duplicate_file=False)
                assert res2.status == 'REJECTED_DUPLICATE_FILE'
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_36_real_database_persistence_and_reconciliation(self, app, db, myntra_tenants):
        """Verify real database persistence of Myntra transactions and exact taxable sum."""
        p1_id = myntra_tenants['profile1_id']
        u1_id = myntra_tenants['user1_id']

        flipkart_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        with app.app_context():
            res = process_import(flipkart_path, p1_id, 'Myntra', u1_id, '012025', '2024-25', allow_duplicate_file=True)
            assert res.status == 'COMPLETED'
            assert res.total_rows == 118
            assert res.success_rows == 118
            assert res.skipped_rows == 0
            assert res.error_rows == 0

            txs = Transaction.query.filter_by(profile_id=p1_id).all()
            assert len(txs) == 118
            db_taxable = sum(tx.taxable_value for tx in txs)
            assert Decimal(str(db_taxable)) == Decimal('3363284.64')

            raw_imports = RawImport.query.filter_by(import_history_id=res.import_history_id).all()
            assert len(raw_imports) == 118

    def test_37_tenant_isolation(self, app, db, myntra_tenants):
        """Verify strict tenant isolation between profiles."""
        p1_id = myntra_tenants['profile1_id']
        p2_id = myntra_tenants['profile2_id']
        u1_id = myntra_tenants['user1_id']
        u2_id = myntra_tenants['user2_id']

        rows1 = [{'Order ID': 'O-T1', 'Order Item ID': 'OI-T1', 'Invoice Number': 'INV-T1', 'Taxable Value': 100, 'Invoice Date': '15-01-2025', 'Product Title': 'P1', 'Selling Price': 100, 'Invoice Value': 118, 'Place of Supply': '27-Maharashtra', 'Supply Type': 'Regular'}]
        rows2 = [{'Order ID': 'O-T2', 'Order Item ID': 'OI-T2', 'Invoice Number': 'INV-T2', 'Taxable Value': 200, 'Invoice Date': '15-01-2025', 'Product Title': 'P2', 'Selling Price': 200, 'Invoice Value': 236, 'Place of Supply': '29-Karnataka', 'Supply Type': 'Regular'}]

        wb1 = _make_myntra_workbook(rows1)
        wb2 = _make_myntra_workbook(rows2)
        path1 = os.path.join(FIXTURE_DIR, 'test_tenant1.xlsx')
        path2 = os.path.join(FIXTURE_DIR, 'test_tenant2.xlsx')
        wb1.save(path1)
        wb2.save(path2)
        try:
            with app.app_context():
                process_import(path1, p1_id, 'Myntra', u1_id, '012025', '2024-25', allow_duplicate_file=True)
                process_import(path2, p2_id, 'Myntra', u2_id, '012025', '2024-25', allow_duplicate_file=True)

                p1_txs = Transaction.query.filter_by(profile_id=p1_id).all()
                p2_txs = Transaction.query.filter_by(profile_id=p2_id).all()

                assert len(p1_txs) == 1
                assert p1_txs[0].invoice_number == 'INV-T1'
                assert len(p2_txs) == 1
                assert p2_txs[0].invoice_number == 'INV-T2'
        finally:
            for p in (path1, path2):
                if os.path.exists(p):
                    os.remove(p)

    def test_38_batch_boundary_rollback(self, app, db, myntra_tenants):
        """Mandatory >1000-row batch boundary rollback test.

        Verifies:
        1. First batch flushes cleanly at row 1000.
        2. Second batch raises failure on second flush.
        3. Entire import rolls back cleanly with ZERO orphan Transaction and RawImport rows.
        4. ImportHistory records FAILED status.
        5. Unrelated tenant data is completely preserved.
        """
        p1_id = myntra_tenants['profile1_id']
        p2_id = myntra_tenants['profile2_id']
        u1_id = myntra_tenants['user1_id']
        u2_id = myntra_tenants['user2_id']

        # Seed tenant 2 with 1 transaction
        p2_row = [{
            'Order ID': 'O-P2', 'Order Item ID': 'OI-P2', 'Invoice Number': 'INV-P2',
            'Invoice Date': '15-01-2025', 'Product Title': 'Item P2', 'Selling Price': 300,
            'Taxable Value': 300, 'Invoice Value': 354, 'Place of Supply': '29-Karnataka',
            'Supply Type': 'Regular',
        }]
        wb_p2 = _make_myntra_workbook(p2_row)
        path_p2 = os.path.join(FIXTURE_DIR, 'test_batch_seed_p2_myntra.xlsx')
        wb_p2.save(path_p2)

        # Generate 1050 rows crossing the 1000-row batch boundary
        rows_1050 = []
        for i in range(1, 1051):
            rows_1050.append({
                'Order ID': f'ORD-M-BATCH-{i:05d}',
                'Order Item ID': f'OI-M-BATCH-{i:05d}',
                'Invoice Number': f'INV-M-BATCH-{i:05d}',
                'Invoice Date': '15-01-2025',
                'Product Title': f'Myntra Product {i}',
                'Selling Price': 100,
                'Taxable Value': 100,
                'Invoice Value': 118,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            })
        wb_batch = _make_myntra_workbook(rows_1050)
        path_batch = os.path.join(FIXTURE_DIR, 'test_myntra_batch_1050.xlsx')
        wb_batch.save(path_batch)

        try:
            with app.app_context():
                res_p2 = process_import(path_p2, p2_id, 'Myntra', u2_id, '012025', '2024-25', allow_duplicate_file=True)
                assert res_p2.status == 'COMPLETED'
                p2_tx_count_before = Transaction.query.filter_by(profile_id=p2_id).count()
                assert p2_tx_count_before == 1

                p1_tx_count_before = Transaction.query.filter_by(profile_id=p1_id).count()

                flush_calls = [0]
                orig_flush = _db.session.flush

                def flush_mock(*args, **kwargs):
                    flush_calls[0] += 1
                    if flush_calls[0] > 1:
                        raise RuntimeError('Injected failure crossing 1000-row batch boundary')
                    return orig_flush(*args, **kwargs)

                with patch.object(_db.session, 'flush', side_effect=flush_mock):
                    with pytest.raises(RuntimeError) as exc_info:
                        process_import(
                            path_batch,
                            p1_id,
                            'Myntra',
                            u1_id,
                            '012025',
                            '2024-25',
                            allow_duplicate_file=True,
                        )
                    assert '1000-row batch boundary' in str(exc_info.value)

                # Confirm first flush reached and second flush triggered
                assert flush_calls[0] >= 2

                # Verify ZERO orphan transactions for profile 1
                p1_tx_count_after = Transaction.query.filter_by(profile_id=p1_id).count()
                assert p1_tx_count_after == p1_tx_count_before

                # Verify failed ImportHistory record exists with ZERO raw imports
                failed_ih = ImportHistory.query.filter_by(profile_id=p1_id, processing_status='FAILED').first()
                assert failed_ih is not None
                raw_count = RawImport.query.filter_by(import_history_id=failed_ih.id).count()
                assert raw_count == 0

                # Verify profile 2 data remains intact
                p2_tx_count_after = Transaction.query.filter_by(profile_id=p2_id).count()
                assert p2_tx_count_after == p2_tx_count_before
        finally:
            for p in (path_p2, path_batch):
                if os.path.exists(p):
                    os.remove(p)

    def test_39_registry_determinism(self):
        """Verify registry precedence and deterministic adapter resolution for Myntra."""
        auto_register_adapters()
        adapter = get_adapter('Myntra')
        assert adapter is not None
        assert adapter.__class__.__module__ == 'app.adapters.myntra'
        assert adapter.PLATFORM_NAME == 'Myntra'
        assert adapter.FORMAT_DOCUMENTED is True
