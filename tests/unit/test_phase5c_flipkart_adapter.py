"""Phase 5C-1 focused tests for Flipkart Marketplace Adapter Production Hardening.

Covers all 19 required areas:
1. Valid Flipkart detection (sheet "GST Report" + headers, generic filename)
2. Invalid workbook rejection (corrupt, empty, unsupported extension)
3. Wrong sheet rejection (non-Flipkart sheet)
4. Required-header rejection (missing invoice_number or taxable_value)
5. Malformed required numeric rejection (unparseable taxable_value generates row error)
6. Valid row parsing (118-row fixture, Order Item ID, Product Title, Selling Price)
7. Decimal precision (exact Decimal parsing without float drift)
8. Date parsing (standardized to DD-MM-YYYY, detected_period)
9. GSTIN mapping (buyer/seller GSTIN formatting, validation, detected_gstin)
10. POS mapping (state code + name resolution to 2-digit code)
11. Cancellation/return handling (Return rows, Credit Note rows, Cancelled rows)
12. Duplicate rows in file (in-file duplicate marked SKIPPED)
13. Full-file duplicate re-import (duplicate file rejection)
14. ImportHistory result recording (status, counts, metadata)
15. Real persistence & reconciliation (Transactions, RawImports, exact taxable sum)
16. Rollback on injected failure (atomic rollback, zero orphaned records)
17. Tenant isolation (scoping across distinct profiles)
18. Unknown workbook not misclassified as Flipkart (Amazon, Meesho, generic)
19. All existing adapter registrations remain valid (Flipkart, Amazon, Generic, etc.)
"""
import io
import os
import uuid
from collections import Counter
from decimal import Decimal
from unittest.mock import patch

import openpyxl
import pytest

from app.adapters.base import ImportRowStatus
from app.adapters.canonical import CanonicalTransaction
from app.adapters.flipkart import FlipkartAdapter
from app.adapters.registry import (
    _ADAPTER_REGISTRY,
    auto_register_adapters,
    detect_platform,
    get_adapter,
)
from app.extensions import db as _db
from app.models import GSTProfile, ImportHistory, RawImport, Transaction, User
from app.services.import_service import process_import

FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')


@pytest.fixture
def flipkart_tenants(app, db):
    """Provide two distinct users and GST profiles for tenant isolation tests."""
    suffix = uuid.uuid4().hex[:8]
    u1_name = f'fk_user1_{suffix}'
    u2_name = f'fk_user2_{suffix}'
    u1_email = f'fk1_{suffix}@example.com'
    u2_email = f'fk2_{suffix}@example.com'
    gstin1 = f'27AAEC{suffix[:4].upper()}1Z5'
    gstin2 = f'29BBED{suffix[:4].upper()}1Z6'

    with app.app_context():
        u1 = User(username=u1_name, email=u1_email)
        u1.set_password('Pass123!')
        _db.session.add(u1)

        u2 = User(username=u2_name, email=u2_email)
        u2.set_password('Pass123!')
        _db.session.add(u2)
        _db.session.commit()

        p1 = GSTProfile(
            user_id=u1.id,
            gstin=gstin1,
            legal_name='Flipkart Seller 1',
            trade_name='Store FK 1',
            state_code='27',
            state_name='Maharashtra',
            financial_year='2024-25',
            filing_frequency='Monthly',
        )
        _db.session.add(p1)

        p2 = GSTProfile(
            user_id=u2.id,
            gstin=gstin2,
            legal_name='Flipkart Seller 2',
            trade_name='Store FK 2',
            state_code='29',
            state_name='Karnataka',
            financial_year='2024-25',
            filing_frequency='Monthly',
        )
        _db.session.add(p2)
        _db.session.commit()

        p1_id = p1.id
        p2_id = p2.id
        u1_id = u1.id
        u2_id = u2.id

    return {
        'user1_id': u1_id,
        'user2_id': u2_id,
        'profile1_id': p1_id,
        'profile2_id': p2_id,
        'gstin1': gstin1,
        'gstin2': gstin2,
    }


def _make_flipkart_workbook(rows_data, sheet_name='GST Report'):
    """Helper to build an in-memory workbook with Flipkart headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        'Order ID', 'Invoice Number', 'Invoice Date', 'Ship From State',
        'Ship To State', 'Buyer GSTIN', 'Seller GSTIN', 'HSN/SAC',
        'Quantity', 'Taxable Value', 'CGST Rate', 'CGST Amount',
        'SGST Rate', 'SGST Amount', 'IGST Rate', 'IGST Amount',
        'Cess Rate', 'Cess Amount', 'Invoice Value', 'Tax Rate',
        'Supply Type', 'Place of Supply', 'Reverse Charge',
        'Order Item ID', 'Product Title', 'Selling Price'
    ]
    ws.append(headers)
    for row in rows_data:
        ws.append([row.get(h, '') for h in headers])
    return wb


# ==============================================================================
# 1. DETECTION TESTS (Requirements 1, 18)
# ==============================================================================

class TestFlipkartDetection:
    """Deterministic Flipkart GST report detection."""

    def test_01_valid_flipkart_detection(self):
        """Valid Flipkart workbook is detected even with generic filename."""
        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        assert os.path.exists(fixture_path)

        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            adapter = detect_platform(wb, 'generic_download_1234.xlsx')
            assert adapter is not None
            assert adapter.PLATFORM_NAME == 'Flipkart'
            assert isinstance(adapter, FlipkartAdapter)
        finally:
            wb.close()

        # Pre-read filename hint check
        direct_adapter = FlipkartAdapter()
        assert direct_adapter.detect(None, 'Flipkart_GST_Report_January_2025.xlsx') is True
        assert direct_adapter.detect(None, 'unrelated_report.xlsx') is False

    def test_18_unknown_workbook_not_misclassified_as_flipkart(self):
        """Unknown or non-Flipkart workbooks must never be classified as Flipkart."""
        # 1. Amazon MTR report
        amazon_path = os.path.join(FIXTURE_DIR, 'amazon_sample.xlsx')
        if os.path.exists(amazon_path):
            wb_amz = openpyxl.load_workbook(amazon_path, data_only=True)
            try:
                assert FlipkartAdapter().detect(wb_amz, 'amazon_sample.xlsx') is False
            finally:
                wb_amz.close()

        # 2. Meesho report
        meesho_path = os.path.join(FIXTURE_DIR, 'meesho_sample.xlsx')
        if os.path.exists(meesho_path):
            wb_meesho = openpyxl.load_workbook(meesho_path, data_only=True)
            try:
                assert FlipkartAdapter().detect(wb_meesho, 'meesho_sample.xlsx') is False
            finally:
                wb_meesho.close()

        # 3. Generic workbook with arbitrary sheet
        wb_gen = openpyxl.Workbook()
        ws = wb_gen.active
        ws.title = 'Sheet1'
        ws.append(['Col1', 'Col2', 'Col3'])
        ws.append(['Val1', 'Val2', 'Val3'])
        try:
            assert FlipkartAdapter().detect(wb_gen, 'arbitrary_data.xlsx') is False
        finally:
            wb_gen.close()


# ==============================================================================
# 2. SCHEMA VALIDATION TESTS (Requirements 2, 3, 4)
# ==============================================================================

class TestFlipkartSchemaValidation:
    """Flipkart schema and structure validation."""

    def test_02_invalid_workbook_rejection(self):
        """Malformed, empty, or unparseable workbooks are rejected."""
        adapter = FlipkartAdapter()

        # None input
        valid, errors = adapter.validate(None)
        assert valid is False
        assert any('No workbook or data provided' in err for err in errors)

        # Corrupted / unsupported file path
        valid, errors = adapter.validate('non_existent_file_path.xyz')
        assert valid is False
        assert len(errors) > 0

        # Empty workbook with no data rows beneath headers
        wb = _make_flipkart_workbook([])
        try:
            valid, errors = adapter.validate(wb)
            assert valid is False
            assert any('No data rows found beneath the header row' in err for err in errors)
        finally:
            wb.close()

    def test_03_wrong_sheet_rejection(self):
        """Workbooks without 'GST Report' sheet are rejected."""
        adapter = FlipkartAdapter()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Orders'
        ws.append(['Invoice Number', 'Taxable Value'])
        ws.append(['INV-01', 1000.0])

        try:
            valid, errors = adapter.validate(wb)
            assert valid is False
            assert any('Unrecognized Flipkart file' in err for err in errors)
        finally:
            wb.close()

    def test_04_required_header_rejection(self):
        """Missing required headers ('Invoice Number' or 'Taxable Value') fails validation."""
        adapter = FlipkartAdapter()

        # Missing 'Invoice Number'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'GST Report'
        ws.append(['Order ID', 'Order Item ID', 'Taxable Value', 'Selling Price'])
        ws.append(['ORD-01', 'OI-01', 1000.0, 1000.0])

        try:
            valid, errors = adapter.validate(wb)
            assert valid is False
            assert any('missing required column' in err.lower() for err in errors)
            assert any('invoice number' in err.lower() for err in errors)
        finally:
            wb.close()

        # Missing 'Taxable Value'
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = 'GST Report'
        ws2.append(['Order ID', 'Order Item ID', 'Invoice Number', 'Selling Price'])
        ws2.append(['ORD-01', 'OI-01', 'INV-01', 1000.0])

        try:
            valid2, errors2 = adapter.validate(wb2)
            assert valid2 is False
            assert any('missing required column' in err.lower() for err in errors2)
            assert any('taxable value' in err.lower() for err in errors2)
        finally:
            wb2.close()


# ==============================================================================
# 3. ROW PARSING & PRECISION TESTS (Requirements 5, 6, 7, 8, 9, 10, 11)
# ==============================================================================

class TestFlipkartRowParsingAndPrecision:
    """Row parsing, precision, canonicalization, and error handling."""

    def test_05_malformed_required_numeric_rejection(self):
        """Malformed required numeric amounts produce explicit row errors rather than silent zero."""
        adapter = FlipkartAdapter()

        rows = [
            # Row 1: Valid row
            {
                'Order ID': 'ORD-OK-01',
                'Invoice Number': 'INV-OK-01',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 1000.0,
                'CGST Rate': 9.0, 'CGST Amount': 90.0,
                'SGST Rate': 9.0, 'SGST Amount': 90.0,
                'Invoice Value': 1180.0,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            },
            # Row 2: Negative taxable value
            {
                'Order ID': 'ORD-BAD-02',
                'Invoice Number': 'INV-BAD-02',
                'Invoice Date': '15-01-2025',
                'Taxable Value': -500.0,
                'Invoice Value': -590.0,
                'Supply Type': 'Regular',
            },
            # Row 3: Missing invoice number
            {
                'Order ID': 'ORD-BAD-03',
                'Invoice Number': '',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 300.0,
                'Supply Type': 'Regular',
            },
            # Row 4: Unparseable date
            {
                'Order ID': 'ORD-BAD-04',
                'Invoice Number': 'INV-BAD-04',
                'Invoice Date': 'NOT_A_DATE_VALUE',
                'Taxable Value': 400.0,
                'Supply Type': 'Regular',
            },
            # Row 5: Malformed non-numeric taxable value string (explicit error)
            {
                'Order ID': 'ORD-BAD-05',
                'Invoice Number': 'INV-BAD-05',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 'INVALID_AMOUNT',
                'Supply Type': 'Regular',
            },
        ]

        wb = _make_flipkart_workbook(rows)
        try:
            result = adapter.parse(wb, 'test_malformed_fk.xlsx')
        finally:
            wb.close()

        assert result.total_rows == 5

        # Row 1: Success
        assert result.rows[0].status == ImportRowStatus.SUCCESS

        # Row 2: Error due to negative taxable value
        assert result.rows[1].status == ImportRowStatus.ERROR
        assert any('negative' in err.lower() for err in result.rows[1].errors)

        # Row 3: Error due to missing invoice number
        assert result.rows[2].status == ImportRowStatus.ERROR
        assert any('missing invoice number' in err.lower() for err in result.rows[2].errors)

        # Row 4: Error due to unparseable date
        assert result.rows[3].status == ImportRowStatus.ERROR
        assert any('not a recognized date' in err.lower() for err in result.rows[3].errors)

        # Row 5: Error for non-numeric taxable value
        assert result.rows[4].status == ImportRowStatus.ERROR
        assert any('not a valid numeric amount' in err.lower() or 'not numeric' in err.lower() for err in result.rows[4].errors)

    def test_06_valid_row_parsing(self):
        """Parse standard 118-row Flipkart fixture and verify column mappings."""
        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        adapter = FlipkartAdapter()

        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            result = adapter.parse(wb, 'flipkart_sample.xlsx')
        finally:
            wb.close()

        assert result.total_rows == 118
        assert result.error_rows == 0
        assert result.success_rows == 118

        first = result.rows[0]
        assert first.canonical.order_item_id.startswith('OI-')
        assert first.canonical.order_id.startswith('ORD-')
        assert first.canonical.invoice_number.startswith('INV-')
        assert first.canonical.description == 'LED Television 43 inch'
        assert first.canonical.unit_price == Decimal('14204.67')
        assert first.canonical.taxable_value == Decimal('28409.34')
        assert first.canonical.source_platform == 'Flipkart'
        assert first.canonical.marketplace_name == 'Flipkart'

    def test_07_decimal_precision(self):
        """Monetary values parse to exact Decimal without float rounding errors."""
        adapter = FlipkartAdapter()
        rows = [
            {
                'Order ID': 'ORD-DEC-01',
                'Invoice Number': 'INV-DEC-01',
                'Invoice Date': '20-01-2025',
                'Taxable Value': 3424.26,
                'CGST Rate': 0, 'CGST Amount': 0,
                'SGST Rate': 0, 'SGST Amount': 0,
                'IGST Rate': 18, 'IGST Amount': 616.37,
                'Cess Rate': 0, 'Cess Amount': 0,
                'Invoice Value': 4040.63,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            }
        ]
        wb = _make_flipkart_workbook(rows)
        try:
            result = adapter.parse(wb, 'decimal_precision.xlsx')
        finally:
            wb.close()

        row = result.rows[0]
        assert isinstance(row.canonical.taxable_value, Decimal)
        assert row.canonical.taxable_value == Decimal('3424.26')
        assert isinstance(row.canonical.igst_amount, Decimal)
        assert row.canonical.igst_amount == Decimal('616.37')
        assert isinstance(row.canonical.invoice_value, Decimal)
        assert row.canonical.invoice_value == Decimal('4040.63')
        assert row.canonical.taxable_value + row.canonical.igst_amount == row.canonical.invoice_value

    def test_08_date_parsing(self):
        """Dates are normalized to DD-MM-YYYY and return period is accurately detected."""
        adapter = FlipkartAdapter()
        rows = [
            {
                'Order ID': 'ORD-DT-01',
                'Invoice Number': 'INV-DT-01',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 100.0,
                'Place of Supply': '27-Maharashtra',
            },
            {
                'Order ID': 'ORD-DT-02',
                'Invoice Number': 'INV-DT-02',
                'Invoice Date': '2025-01-22',
                'Taxable Value': 200.0,
                'Place of Supply': '27-Maharashtra',
            },
        ]
        wb = _make_flipkart_workbook(rows)
        try:
            result = adapter.parse(wb, 'dates.xlsx')
        finally:
            wb.close()

        assert result.rows[0].canonical.invoice_date == '15-01-2025'
        assert result.rows[1].canonical.invoice_date == '22-01-2025'
        assert result.detected_period == '012025'

    def test_09_gstin_mapping(self):
        """GSTIN fields are formatted, validated, and detected on result."""
        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        adapter = FlipkartAdapter()

        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            result = adapter.parse(wb, 'flipkart_sample.xlsx')
        finally:
            wb.close()

        assert result.detected_gstin == '27AABCU9603R1ZM'
        # B2B rows have customer_gstin
        b2b_rows = [r for r in result.rows if r.canonical.customer_gstin]
        assert len(b2b_rows) >= 30
        for r in b2b_rows:
            assert len(r.canonical.customer_gstin) == 15
            assert r.canonical.customer_gstin == r.canonical.customer_gstin.upper()

    def test_10_pos_mapping(self):
        """Place of supply in 'Code-State' format resolves to 2-digit GST code."""
        adapter = FlipkartAdapter()
        rows = [
            {
                'Order ID': 'ORD-POS-01',
                'Invoice Number': 'INV-POS-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 500.0,
                'Place of Supply': '09-Uttar Pradesh',
            },
            {
                'Order ID': 'ORD-POS-02',
                'Invoice Number': 'INV-POS-02',
                'Invoice Date': '11-01-2025',
                'Taxable Value': 600.0,
                'Place of Supply': '27-Maharashtra',
            },
        ]
        wb = _make_flipkart_workbook(rows)
        try:
            result = adapter.parse(wb, 'pos.xlsx')
        finally:
            wb.close()

        assert result.rows[0].canonical.place_of_supply == '09'
        assert result.rows[0].canonical.place_of_supply_raw == '09-Uttar Pradesh'
        assert result.rows[1].canonical.place_of_supply == '27'
        assert result.rows[1].canonical.place_of_supply_raw == '27-Maharashtra'

    def test_11_cancellation_return_credit_note_handling(self):
        """Returns, Credit Notes, and Cancellations are mapped to their respective flags and types."""
        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')
        adapter = FlipkartAdapter()

        wb = openpyxl.load_workbook(fixture_path, data_only=True)
        try:
            result = adapter.parse(wb, 'flipkart_sample.xlsx')
        finally:
            wb.close()

        # Returns (3 rows)
        returns = [r for r in result.rows if r.canonical.return_flag and r.canonical.supply_type == 'Return']
        assert len(returns) == 3
        assert all(r.canonical.cancellation_flag is False for r in returns)

        # Credit Notes (5 rows)
        notes = [r for r in result.rows if r.canonical.note_type == 'CREDIT']
        assert len(notes) == 5
        assert all(r.canonical.note_type_raw == 'Credit Note' for r in notes)
        assert all(r.canonical.invoice_number.startswith('CN-') for r in notes)
        assert all(r.canonical.note_number == r.canonical.invoice_number for r in notes)

        # Test explicit Cancelled row
        cancelled_rows = [
            {
                'Order ID': 'ORD-CAN-01',
                'Invoice Number': 'INV-CAN-01',
                'Invoice Date': '15-01-2025',
                'Taxable Value': 500.0,
                'Supply Type': 'Cancelled',
                'Place of Supply': '27-Maharashtra',
            }
        ]
        wb_can = _make_flipkart_workbook(cancelled_rows)
        try:
            res_can = adapter.parse(wb_can, 'cancelled.xlsx')
        finally:
            wb_can.close()

        can_row = res_can.rows[0]
        assert can_row.canonical.cancellation_flag is True
        assert can_row.canonical.return_flag is True
        assert can_row.canonical.supply_type == 'Cancelled'


# ==============================================================================
# 4. PIPELINE PERSISTENCE & SYSTEM TESTS (Requirements 12-17, 19)
# ==============================================================================

class TestFlipkartPipelinePersistence:
    """End-to-end import pipeline persistence, batching, and tenant isolation."""

    def test_12_in_file_duplicate_rows_marked_skipped(self, app, db, flipkart_tenants, tmp_path):
        """Duplicate rows within the same file are detected and marked SKIPPED."""
        p_id = flipkart_tenants['profile1_id']
        u_id = flipkart_tenants['user1_id']

        rows = [
            {
                'Order ID': 'ORD-DUP-01',
                'Order Item ID': 'OI-DUP-01',
                'Invoice Number': 'INV-DUP-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 1000.0,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            },
            {
                'Order ID': 'ORD-DUP-01',
                'Order Item ID': 'OI-DUP-01',
                'Invoice Number': 'INV-DUP-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 1000.0,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            },
        ]
        wb = _make_flipkart_workbook(rows)
        path = str(tmp_path / 'test_flipkart_in_file_dup.xlsx')
        wb.save(path)
        wb.close()
        try:
            with app.app_context():
                res = process_import(
                    file_path=path,
                    profile_id=p_id,
                    platform_name='Flipkart',
                    user_id=u_id,
                    return_period='012025',
                    financial_year='2024-25',
                )
                assert res.status == 'PARTIAL'
                assert res.total_rows == 2
                assert res.success_rows == 1
                assert res.skipped_rows == 1
                assert res.error_rows == 0

                ih = _db.session.get(ImportHistory, res.import_history_id)
                assert ih.total_rows == 2
                assert ih.success_rows == 1
                assert ih.skipped_rows == 1
        finally:
            if os.path.exists(path):
                try:
                    os.remove(path)
                except OSError:
                    pass

    def test_13_full_file_duplicate_reimport(self, app, db, flipkart_tenants):
        """Re-uploading the exact same file is rejected as a duplicate file."""
        p_id = flipkart_tenants['profile1_id']
        u_id = flipkart_tenants['user1_id']

        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')

        with app.app_context():
            res1 = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Flipkart',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
            )
            assert res1.status == 'COMPLETED'
            assert res1.success_rows == 118

            # Second import without allow_duplicate_file is rejected as duplicate file
            res2 = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Flipkart',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
                allow_duplicate_file=False,
            )
            assert res2.status == 'REJECTED_DUPLICATE_FILE'
            assert 'already imported successfully' in res2.errors[0]

            # Re-import with allow_duplicate_file=True skips all rows and avoids double-counting
            res3 = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Flipkart',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
                allow_duplicate_file=True,
            )
            assert res3.status == 'PARTIAL'
            assert res3.total_rows == 118
            assert res3.success_rows == 0
            assert res3.skipped_rows == 118
            assert res3.error_rows == 0

            # Database row count strictly remains 118
            assert Transaction.query.filter_by(profile_id=p_id).count() == 118

    def test_14_import_history_and_15_database_persistence_and_reconciliation(self, app, db, flipkart_tenants):
        """Full 118-row fixture imports cleanly into database with exact financial reconciliation."""
        p_id = flipkart_tenants['profile1_id']
        u_id = flipkart_tenants['user1_id']

        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')

        with app.app_context():
            res = process_import(
                file_path=fixture_path,
                profile_id=p_id,
                platform_name='Flipkart',
                user_id=u_id,
                return_period='012025',
                financial_year='2024-25',
                allow_duplicate_file=True,
            )
            assert res.status == 'COMPLETED'
            assert res.total_rows == 118
            assert res.success_rows == 118
            assert res.error_rows == 0
            assert res.skipped_rows == 0

            ih = _db.session.get(ImportHistory, res.import_history_id)
            assert ih is not None
            assert ih.processing_status == 'COMPLETED'
            assert ih.total_rows == 118
            assert ih.success_rows == 118
            assert ih.processing_duration_ms > 0

            # Verify persisted Transactions
            txs = Transaction.query.filter_by(import_history_id=ih.id).all()
            assert len(txs) == 118

            # Verify persisted RawImports
            raws = RawImport.query.filter_by(import_history_id=ih.id).all()
            assert len(raws) == 118

            # Verify GSTR-1 classification breakdown
            table_counts = Counter(t.gstr1_table for t in txs)
            assert table_counts['b2b'] == 30
            assert table_counts['b2cs'] == 83
            assert table_counts['cdnr'] == 5

            # Mathematical reconciliation of taxable total
            total_taxable = sum(t.taxable_value for t in txs)
            assert total_taxable == Decimal('3363284.64')

    def test_16_rollback_on_injected_failure(self, app, db, flipkart_tenants):
        """Failure during persistence rolls back atomically leaving zero orphaned records."""
        p_id = flipkart_tenants['profile1_id']
        u_id = flipkart_tenants['user1_id']

        fixture_path = os.path.join(FIXTURE_DIR, 'flipkart_sample.xlsx')

        with app.app_context():
            initial_tx_count = Transaction.query.filter_by(profile_id=p_id).count()

            with patch('app.extensions.db.session.flush', side_effect=RuntimeError('Injected database failure')):
                with pytest.raises(RuntimeError):
                    process_import(
                        file_path=fixture_path,
                        profile_id=p_id,
                        platform_name='Flipkart',
                        user_id=u_id,
                        return_period='012025',
                        financial_year='2024-25',
                        allow_duplicate_file=True,
                    )

            # Confirm complete rollback
            assert Transaction.query.filter_by(profile_id=p_id).count() == initial_tx_count

            # Confirm ImportHistory was transitioned to FAILED and has zero raw imports
            failed_ih = ImportHistory.query.filter_by(profile_id=p_id, processing_status='FAILED').first()
            assert failed_ih is not None
            assert RawImport.query.filter_by(import_history_id=failed_ih.id).count() == 0

    def test_17_tenant_isolation(self, app, db, flipkart_tenants, tmp_path):
        """Transactions imported for Profile 1 are completely isolated from Profile 2."""
        p1_id = flipkart_tenants['profile1_id']
        u1_id = flipkart_tenants['user1_id']
        p2_id = flipkart_tenants['profile2_id']
        u2_id = flipkart_tenants['user2_id']

        rows_p1 = [
            {
                'Order ID': 'ORD-P1-01',
                'Invoice Number': 'INV-P1-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 1000.0,
                'Place of Supply': '27-Maharashtra',
                'Supply Type': 'Regular',
            }
        ]
        rows_p2 = [
            {
                'Order ID': 'ORD-P2-01',
                'Invoice Number': 'INV-P2-01',
                'Invoice Date': '10-01-2025',
                'Taxable Value': 2000.0,
                'Place of Supply': '29-Karnataka',
                'Supply Type': 'Regular',
            }
        ]

        wb1 = _make_flipkart_workbook(rows_p1)
        path1 = str(tmp_path / 'test_tenant_p1.xlsx')
        wb1.save(path1)
        wb1.close()

        wb2 = _make_flipkart_workbook(rows_p2)
        path2 = str(tmp_path / 'test_tenant_p2.xlsx')
        wb2.save(path2)
        wb2.close()

        try:
            with app.app_context():
                process_import(path1, p1_id, 'Flipkart', u1_id, '012025', '2024-25')
                process_import(path2, p2_id, 'Flipkart', u2_id, '012025', '2024-25')

                txs_p1 = Transaction.query.filter_by(profile_id=p1_id).all()
                txs_p2 = Transaction.query.filter_by(profile_id=p2_id).all()

                assert len(txs_p1) == 1
                assert len(txs_p2) == 1
                assert txs_p1[0].invoice_number == 'INV-P1-01'
                assert txs_p2[0].invoice_number == 'INV-P2-01'
                assert txs_p1[0].profile_id == p1_id
                assert txs_p2[0].profile_id == p2_id
        finally:
            for p in (path1, path2):
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except OSError:
                        pass

    def test_19_all_existing_adapter_registrations_remain_valid(self):
        """All existing adapters remain properly registered and resolve deterministically."""
        auto_register_adapters()

        fk = get_adapter('Flipkart')
        assert fk is not None
        assert fk.__class__.__module__ == 'app.adapters.flipkart'
        assert fk.__class__.__name__ == 'FlipkartAdapter'

        amz = get_adapter('Amazon')
        assert amz is not None
        assert amz.__class__.__module__ == 'app.adapters.amazon'
        assert amz.__class__.__name__ == 'AmazonAdapter'

        gen = get_adapter('Generic')
        assert gen is not None
        assert gen.__class__.__module__ == 'app.adapters.all_adapters'
        assert gen.__class__.__name__ == '_GenericFilenameAdapter'

        for platform_name in list(_ADAPTER_REGISTRY.keys()):
            adapter = get_adapter(platform_name)
            assert adapter is not None, f"Failed to instantiate {platform_name}"
            assert adapter.PLATFORM_NAME == platform_name
