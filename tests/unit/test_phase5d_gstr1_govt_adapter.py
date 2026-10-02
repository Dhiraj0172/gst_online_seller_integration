"""Phase 5D focused tests for GSTR-1 Government Excel Adapter & Output Compatibility.

Covers all 32 required areas:
1. Fixture loads
2. Every expected sheet recognized
3. Every expected header recognized
4. Amendment sheets recognized
5. Docs sheet behavior
6. Malformed sheet rejection
7. Unknown sheet rejection
8. Missing-header rejection
9. B2B normalization
10. B2BA normalization
11. B2CL normalization
12. B2CLA normalization
13. B2CS aggregate handling
14. B2CSA aggregate/amendment handling
15. CDNR normalization
16. CDNRA normalization
17. CDNUR normalization
18. CDNURA normalization
19. EXP normalization
20. EXPA normalization
21. NIL aggregate handling
22. HSN aggregate handling
23. HSN B2C handling
24. Document summary handling
25. Source-row preservation
26. No silent row loss
27. Warning/error accounting
28. Duplicate handling
29. Multi-tenant isolation
30. Return-period isolation
31. Re-import behavior
32. Malformed workbook handling
Plus:
33. Persistence without fake invoice numbers
34. Bidirectional round-trip (Import -> DB -> GSTR-1 Generator -> Excel & JSON)
35. Programmatic Excel and JSON regression & schema validation
"""
import io
import json
import os
import uuid
from collections import Counter
from decimal import Decimal
import openpyxl
import pytest

from app.adapters.base import ImportRowStatus
from app.adapters.canonical import CanonicalTransaction
from app.adapters.gstr1_govt import GSTR1GovtAdapter
from app.adapters.registry import (
    _ADAPTER_REGISTRY,
    auto_register_adapters,
    detect_platform,
    get_adapter,
)
from app.extensions import db as _db
from app.models import GSTProfile, ImportHistory, RawImport, Transaction, User
from app.services.gstr1_excel_writer import GSTR1ExcelWriter
from app.services.gstr1_generator import generate_gstr1
from app.services.import_service import process_import

FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')
SAMPLE_GSTR1_PATH = os.path.join(FIXTURE_DIR, 'gstr1_sample.xlsx')


@pytest.fixture
def gstr1_adapter():
    return GSTR1GovtAdapter()


@pytest.fixture
def seller_context(app, db):
    """Create a persistent user, GST profile, and return period context."""
    suffix = uuid.uuid4().hex[:8]
    email = f"gstr1_seller_{suffix}@example.com"
    digits = str(1000 + (int(suffix, 16) % 8999))
    gstin = f"29ABCDE{digits}F1Z5"

    with app.app_context():
        user = User(email=email, username=f"seller_{suffix}")
        user.set_password("SecurePass123!")
        db.session.add(user)
        db.session.commit()

        profile = GSTProfile(
            user_id=user.id,
            gstin=gstin,
            legal_name="GSTR1 Testing Merchant Ltd",
            trade_name="GSTR1 Testing",
            state_code="29",
            state_name="Karnataka",
            financial_year="2024-25",
            filing_frequency="Monthly",
        )
        db.session.add(profile)
        db.session.commit()

        context_data = {
            "user_id": user.id,
            "profile_id": profile.id,
            "gstin": gstin,
            "return_period": "072024",
            "financial_year": "2024-25",
        }

    yield context_data

    with app.app_context():
        try:
            u = User.query.filter_by(email=email).first()
            if u:
                db.session.delete(u)
                db.session.commit()
        except Exception:
            db.session.rollback()


class TestPhase5DFixtureAndDetection:
    """Requirement 1-8, 32: Fixture presence, detection, headers, sheets, validation."""

    def test_01_fixture_loads_successfully(self):
        assert os.path.exists(SAMPLE_GSTR1_PATH), f"Fixture not found at {SAMPLE_GSTR1_PATH}"
        wb = openpyxl.load_workbook(SAMPLE_GSTR1_PATH, data_only=True)
        expected_sheets = [
            'b2b', 'b2ba', 'b2cl', 'b2cla', 'b2cs', 'b2csa',
            'cdnr', 'cdnra', 'cdnur', 'cdnura', 'exp', 'expa',
            'nil', 'hsn', 'hsnb2c', 'docs'
        ]
        for name in expected_sheets:
            assert name in wb.sheetnames, f"Expected sheet '{name}' missing from fixture"
        wb.close()

    def test_02_every_expected_sheet_recognized(self, gstr1_adapter):
        wb = openpyxl.load_workbook(SAMPLE_GSTR1_PATH, data_only=True)
        assert gstr1_adapter.detect(wb, "some_generic_file.xlsx") is True
        for sheet_name in gstr1_adapter.SHEET_NAMES:
            single_sheet_wb = openpyxl.Workbook()
            ws = single_sheet_wb.active
            ws.title = sheet_name
            headers = GSTR1ExcelWriter.SHEET_CONFIGS[sheet_name]['headers']
            ws.append(headers)
            ws.append(['DUMMY'] * len(headers))
            assert gstr1_adapter.detect(single_sheet_wb, "report.xlsx") is True

    def test_03_every_expected_header_recognized(self, gstr1_adapter):
        for sheet_name in gstr1_adapter.ALL_KNOWN_SHEETS:
            headers = GSTR1ExcelWriter.SHEET_CONFIGS[sheet_name]['headers']
            mapping, unknown = gstr1_adapter.resolve_headers(headers)
            assert len(mapping) > 0, f"Sheet '{sheet_name}' mapped zero headers"
            assert 'taxable_value' in mapping or 'invoice_number' in mapping or sheet_name in ('nil', 'docs')

    def test_04_amendment_sheets_recognized(self, gstr1_adapter):
        expected_amendments = ('b2ba', 'b2cla', 'b2csa', 'cdnra', 'cdnura', 'expa')
        for sheet in expected_amendments:
            assert sheet in gstr1_adapter.AMENDMENT_SHEETS

    def test_05_docs_sheet_behavior(self, gstr1_adapter):
        assert 'docs' in gstr1_adapter.NON_TRANSACTION_SHEETS
        res = gstr1_adapter.parse(SAMPLE_GSTR1_PATH)
        assert res.errors == []
        assert 'docs_summary' in res.metadata
        assert len(res.metadata['docs_summary']) == 3
        # docs summary rows are skipped from transaction rows
        for row in res.rows:
            assert row.sheet_name != 'docs'

    def test_06_malformed_sheet_rejection(self, gstr1_adapter):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'b2b'
        ws.append(['Wrong Column 1', 'Wrong Column 2'])
        ws.append(['val1', 'val2'])
        is_valid, errors = gstr1_adapter.validate(wb)
        assert is_valid is False
        assert any('recognizable header row' in err for err in errors)

    def test_07_unknown_sheet_rejection(self, gstr1_adapter):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'UnknownSheet'
        ws.append(['col1', 'col2'])
        ws.append(['val1', 'val2'])
        is_valid, errors = gstr1_adapter.validate(wb)
        assert is_valid is False
        assert any('no recognized GSTR-1 sheet found' in err for err in errors)

    def test_08_missing_header_rejection(self, gstr1_adapter):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'b2b'
        # Missing required 'Taxable Value' column
        ws.append(['GSTIN/UIN of Recipient', 'Invoice Number', 'Invoice date', 'Invoice Value'])
        ws.append(['29AALCS5765L1ZP', 'INV-100', '15-01-2025', 1000])
        is_valid, errors = gstr1_adapter.validate(wb)
        assert is_valid is False
        assert any('missing required column' in err for err in errors)

    def test_32_malformed_workbook_handling(self, gstr1_adapter):
        res = gstr1_adapter.parse(b"NOT_A_VALID_EXCEL_BYTES", "corrupt.xlsx")
        assert len(res.errors) > 0
        assert res.total_rows == 0


class TestPhase5DSheetNormalizations:
    """Requirement 9-27: Per-sheet canonical mapping, warnings, errors, and accounting."""

    @pytest.fixture(autouse=True)
    def parsed_fixture(self, gstr1_adapter):
        self.result = gstr1_adapter.parse(SAMPLE_GSTR1_PATH)
        assert self.result.errors == []
        self.rows_by_sheet = {}
        for r in self.result.rows:
            self.rows_by_sheet.setdefault(r.sheet_name, []).append(r)

    def test_09_b2b_normalization(self):
        b2b_rows = self.rows_by_sheet['b2b']
        assert len(b2b_rows) == 3
        first = b2b_rows[0].normalized_data
        assert first['invoice_number'] == 'INV-2025-001'
        assert first['customer_gstin'] == '29AALCS5765L1ZP'
        assert first['taxable_value'] == Decimal('10000')
        assert first['tax_rate'] == Decimal('18')
        assert first['place_of_supply'] == '29'
        assert first['supply_type'] == 'B2B'
        assert first['invoice_type'] == 'Regular'

    def test_10_b2ba_normalization(self):
        b2ba_rows = self.rows_by_sheet['b2ba']
        assert len(b2ba_rows) == 1
        r = b2ba_rows[0].normalized_data
        assert r['amendment_flag'] is True
        assert r['original_invoice_number'] == 'INV-2024-999'
        assert r['invoice_number'] == 'INV-2024-999-REV'
        assert r['taxable_value'] == Decimal('12000')
        assert r['supply_type'] == 'B2BA'

    def test_11_b2cl_normalization(self):
        b2cl_rows = self.rows_by_sheet['b2cl']
        assert len(b2cl_rows) == 1
        r = b2cl_rows[0].normalized_data
        assert r['invoice_number'] == 'INV-2025-004'
        assert r['place_of_supply'] == '27'
        assert r['taxable_value'] == Decimal('254237.29')
        assert r['supply_type'] == 'B2CL'

    def test_12_b2cla_normalization(self):
        b2cla_rows = self.rows_by_sheet['b2cla']
        assert len(b2cla_rows) == 1
        r = b2cla_rows[0].normalized_data
        assert r['amendment_flag'] is True
        assert r['original_invoice_number'] == 'INV-2024-888'
        assert r['invoice_number'] == 'INV-2024-888-REV'
        assert r['place_of_supply'] == '27'
        assert r['taxable_value'] == Decimal('300000')
        assert r['supply_type'] == 'B2CLA'

    def test_13_b2cs_aggregate_handling(self):
        b2cs_rows = self.rows_by_sheet['b2cs']
        assert len(b2cs_rows) == 3
        for row in b2cs_rows:
            r = row.normalized_data
            assert r.get('is_aggregate') is True
            assert r.get('supply_type') == 'B2CS'
            assert r.get('invoice_number') in (None, '')
            assert r.get('taxable_value') > Decimal('0')

    def test_14_b2csa_aggregate_amendment_handling(self):
        b2csa_rows = self.rows_by_sheet['b2csa']
        assert len(b2csa_rows) == 1
        r = b2csa_rows[0].normalized_data
        assert r.get('is_aggregate') is True
        assert r.get('amendment_flag') is True
        assert r.get('supply_type') == 'B2CSA'
        assert r.get('place_of_supply') == '27'

    def test_15_cdnr_normalization(self):
        cdnr_rows = self.rows_by_sheet['cdnr']
        assert len(cdnr_rows) == 2
        cn = next(r.normalized_data for r in cdnr_rows if r.normalized_data['note_number'] == 'CN-2025-001')
        dn = next(r.normalized_data for r in cdnr_rows if r.normalized_data['note_number'] == 'DN-2025-001')
        assert cn['note_type'] == 'CREDIT'
        assert cn['original_invoice_number'] == 'INV-2025-001'
        assert cn['supply_type'] == 'CDNR'
        assert dn['note_type'] == 'DEBIT'
        assert dn['original_invoice_number'] == 'INV-2025-001'
        assert dn['supply_type'] == 'CDNR'

    def test_16_cdnra_normalization(self):
        cdnra_rows = self.rows_by_sheet['cdnra']
        assert len(cdnra_rows) == 1
        r = cdnra_rows[0].normalized_data
        assert r['amendment_flag'] is True
        assert r['original_invoice_number'] == 'CN-2024-901'
        assert r['note_number'] == 'CN-2024-901-REV'
        assert r['supply_type'] == 'CDNRA'

    def test_17_cdnur_normalization(self):
        cdnur_rows = self.rows_by_sheet['cdnur']
        assert len(cdnur_rows) == 1
        r = cdnur_rows[0].normalized_data
        assert r['note_number'] == 'CDNUR-001'
        assert r['original_invoice_number'] == 'INV-2025-004'
        assert r['note_type'] == 'CREDIT'
        assert r['supply_type'] == 'CDNUR'

    def test_18_cdnura_normalization(self):
        cdnura_rows = self.rows_by_sheet['cdnura']
        assert len(cdnura_rows) == 1
        r = cdnura_rows[0].normalized_data
        assert r['amendment_flag'] is True
        assert r['original_invoice_number'] == 'CDNUR-901'
        assert r['note_number'] == 'CDNUR-901-REV'
        assert r['supply_type'] == 'CDNURA'

    def test_19_exp_normalization(self):
        exp_rows = self.rows_by_sheet['exp']
        assert len(exp_rows) == 1
        r = exp_rows[0].normalized_data
        assert r['invoice_number'] == 'EXP-2025-001'
        assert r['taxable_value'] == Decimal('50000')
        assert r['supply_type'] == 'EXPORT'

    def test_20_expa_normalization(self):
        expa_rows = self.rows_by_sheet['expa']
        assert len(expa_rows) == 1
        r = expa_rows[0].normalized_data
        assert r['amendment_flag'] is True
        assert r['original_invoice_number'] == 'EXP-2024-901'
        assert r['invoice_number'] == 'EXP-2024-901-REV'
        assert r['taxable_value'] == Decimal('60000')
        assert r['supply_type'] == 'EXPA'

    def test_21_nil_aggregate_handling(self):
        nil_rows = self.rows_by_sheet['nil']
        assert len(nil_rows) == 4
        for row in nil_rows:
            r = row.normalized_data
            assert r['is_aggregate'] is True
            assert r['source_sheet'] == 'nil'
            assert 'source_metadata' in r
            meta = r['source_metadata']
            assert 'nil_sply_ty' in meta
            assert 'nil_rated_supplies' in meta
            assert 'exempt_supplies' in meta
            assert 'non_gst_supplies' in meta

    def test_22_hsn_aggregate_handling(self):
        hsn_rows = self.rows_by_sheet['hsn']
        assert len(hsn_rows) == 2
        for row in hsn_rows:
            r = row.normalized_data
            assert r['is_aggregate'] is True
            assert r['supply_type'] == 'HSN'
            assert r['hsn_sac'] in ('6109', '6203')
            assert r['quantity'] in (Decimal('50'), Decimal('20'))

    def test_23_hsn_b2c_handling(self):
        hsnb2c_rows = self.rows_by_sheet['hsnb2c']
        assert len(hsnb2c_rows) == 2
        for row in hsnb2c_rows:
            r = row.normalized_data
            assert r['is_aggregate'] is True
            assert r['supply_type'] == 'HSNB2C'
            assert r['hsn_sac'] in ('6109', '6203')
            assert r['quantity'] in (Decimal('100'), Decimal('30'))

    def test_24_document_summary_handling(self):
        docs_summary = self.result.metadata.get('docs_summary')
        assert docs_summary is not None
        assert len(docs_summary) == 3
        natures = [d.get('Nature of Document') for d in docs_summary]
        assert 'Invoices for outward supply' in natures
        assert 'Credit Note' in natures
        assert 'Debit Note' in natures

    def test_25_source_row_preservation(self):
        for row in self.result.rows:
            assert row.row_number >= 2
            assert row.sheet_name in GSTR1GovtAdapter.SHEET_NAMES
            assert row.normalized_data.get('source_sheet') == row.sheet_name

    def test_26_no_silent_row_loss(self):
        # 3 + 1 + 1 + 1 + 3 + 1 + 2 + 1 + 1 + 1 + 1 + 1 + 4 + 2 + 2 = 25 rows
        assert self.result.total_rows == 25
        assert len(self.result.rows) == 25

    def test_27_warning_error_accounting(self):
        assert self.result.error_rows == 0
        assert self.result.warning_rows > 0
        assert self.result.success_rows + self.result.warning_rows == self.result.total_rows


class TestPhase5DPersistenceAndLifecycle:
    """Requirement 28-31, 33: Persistence without fake invoice numbers, duplicates, multi-tenant isolation, return period."""

    def test_33_persistence_without_fake_invoice_numbers(self, seller_context):
        res = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
        )
        assert res.status in ('COMPLETED', 'PARTIAL')
        assert res.total_rows == 25

        txs = Transaction.query.filter_by(profile_id=seller_context["profile_id"]).all()
        assert len(txs) == 25

        # Verify aggregate rows have invoice_number = None and invoice_type = 'aggregate'
        agg_txs = [tx for tx in txs if tx.invoice_type == 'aggregate']
        # b2cs (3) + b2csa (1) + nil (4) + hsn (2) + hsnb2c (2) = 12 aggregate rows
        assert len(agg_txs) == 12
        for tx in agg_txs:
            assert tx.invoice_number is None or tx.invoice_number == ''
            meta = json.loads(tx.source_metadata) if tx.source_metadata else {}
            assert meta.get('is_aggregate') is True

        # Verify regular invoice rows have their genuine document numbers
        regular_txs = [tx for tx in txs if tx.invoice_type != 'aggregate']
        assert len(regular_txs) == 13
        for tx in regular_txs:
            assert tx.invoice_number is not None and len(tx.invoice_number) > 0

    def test_28_duplicate_handling(self, seller_context):
        # First import
        res1 = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
        )
        assert res1.status in ('COMPLETED', 'PARTIAL')

        # Re-import same file without duplicate override -> REJECTED_DUPLICATE_FILE
        res2 = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
            allow_duplicate_file=False,
        )
        assert res2.status == 'REJECTED_DUPLICATE_FILE'

    def test_29_multi_tenant_isolation(self, app, db, seller_context):
        # Import data for Seller 1
        process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
        )

        # Create Seller 2
        suffix = uuid.uuid4().hex[:8]
        user2 = User(email=f"other_{suffix}@example.com", username=f"other_{suffix}")
        user2.set_password("Pass123!")
        db.session.add(user2)
        db.session.commit()

        digits2 = str(1000 + (int(suffix, 16) % 8999))
        profile2 = GSTProfile(
            user_id=user2.id,
            gstin=f"27ABCDE{digits2}F1Z8",
            legal_name="Other Seller Ltd",
            state_code="27",
            state_name="Maharashtra",
            financial_year="2024-25",
            filing_frequency="Monthly",
        )
        db.session.add(profile2)
        db.session.commit()

        # Seller 2 should have 0 transactions
        seller2_txs = Transaction.query.filter_by(profile_id=profile2.id).all()
        assert len(seller2_txs) == 0

        # Seller 1 should have exactly 25 transactions
        seller1_txs = Transaction.query.filter_by(profile_id=seller_context["profile_id"]).all()
        assert len(seller1_txs) == 25

    def test_30_return_period_isolation(self, seller_context):
        # Import for 072024
        process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period="072024",
            financial_year="2024-25",
            platform_name="GSTR1_Govt",
        )

        ih_july = ImportHistory.query.filter_by(
            profile_id=seller_context["profile_id"], return_period="072024"
        ).first()
        assert ih_july is not None

        ih_aug = ImportHistory.query.filter_by(
            profile_id=seller_context["profile_id"], return_period="082024"
        ).first()
        assert ih_aug is None

    def test_31_re_import_behavior(self, seller_context):
        # First import
        res1 = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
        )
        assert res1.status in ('COMPLETED', 'PARTIAL')

        # Allowed re-import: duplicates are skipped idempotently
        res2 = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
            allow_duplicate_file=True,
        )
        assert res2.status in ('COMPLETED', 'PARTIAL')
        # Total transactions remains 25 because rows were detected as existing
        tx_count = Transaction.query.filter_by(profile_id=seller_context["profile_id"]).count()
        assert tx_count == 25


class TestPhase5DBidirectionalRoundTripAndReconciliation:
    """Requirement 34 & Step 8: Complete round-trip from Import -> DB -> Generator -> Excel & JSON."""

    def test_34_bidirectional_round_trip(self, seller_context):
        # Step 1: Import fixture
        import_res = process_import(
            file_path=SAMPLE_GSTR1_PATH,
            profile_id=seller_context["profile_id"],
            user_id=seller_context["user_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            platform_name="GSTR1_Govt",
        )
        assert import_res.status in ('COMPLETED', 'PARTIAL')
        assert import_res.total_rows == 25

        # Step 2: Generate GSTR-1
        clean_report = {
            'status': 'SUCCESS',
            'is_generation_blocked': False,
            'critical_failures': 0,
            'warnings': 0,
            'differences': []
        }
        gen_res = generate_gstr1(
            profile_id=seller_context["profile_id"],
            return_period=seller_context["return_period"],
            financial_year=seller_context["financial_year"],
            reconciliation_report=clean_report
        )

        assert os.path.exists(gen_res.excel_path), f"Excel output missing: {gen_res.excel_path}"
        assert os.path.exists(gen_res.json_path), f"JSON output missing: {gen_res.json_path}"

        # Step 3: Validate generated JSON schema and business rules
        assert gen_res.validation_result.is_valid is True
        assert gen_res.validation_result.errors == []

        # Step 4: Re-open generated Excel workbook and verify sheets & row counts
        gen_wb = openpyxl.load_workbook(gen_res.excel_path, data_only=True)
        expected_counts = {
            'b2b': 3,
            'b2ba': 1,
            'b2cl': 1,
            'b2cla': 1,
            'b2cs': 3,
            'b2csa': 1,
            'cdnr': 2,
            'cdnra': 1,
            'cdnur': 1,
            'cdnura': 1,
            'exp': 1,
            'expa': 1,
            'nil': 4,
            'hsn': 4,
            'docs': 3,
        }

        for sheet_name, expected_rows in expected_counts.items():
            assert sheet_name in gen_wb.sheetnames, f"Generated Excel missing sheet '{sheet_name}'"
            ws = gen_wb[sheet_name]
            data_rows = [row for row in ws.iter_rows(values_only=True) if any(row)]
            # First row is headers, remainder are data rows
            assert len(data_rows) - 1 == expected_rows, (
                f"Sheet '{sheet_name}' has {len(data_rows) - 1} rows, expected {expected_rows}"
            )
        gen_wb.close()
