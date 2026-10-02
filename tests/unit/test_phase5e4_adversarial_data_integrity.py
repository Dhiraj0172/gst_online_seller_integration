"""Adversarial Verification Track 2: Data Integrity & Schema Compliance.

Tests:
1. Section Counts & Tax Math Integrity:
   - total_b2b strictly counts B2B invoices and NOT total invoices or B2CS or line items.
   - stats dictionary correctly computes all 11 counts:
     total_b2b, total_b2cs, total_b2cl, total_cdnr, total_cdnur, total_exp, total_nil,
     total_hsn_b2b, total_hsn_b2c, total_ecom, total_invoices.
   - Tax totals: total_taxable_value, total_cgst, total_sgst, total_igst, total_cess, total_tax
     sum accurately without floating-point drift under micro-transactions and boundary numbers.
2. Option Toggles & Schema Validation:
   - include_hsn=False: HSN data is omitted/empty in JSON, HSN sheet is omitted from Excel,
     and GSTR1Validator still accepts the payload.
   - include_hsn=True: HSN summary matches item-level taxes in both combined (pre-May 2025)
     and separate_b2b_b2c (post-May 2025) reporting modes.
3. Persistence Integrity:
   - All 15 columns in GSTR1Generation populated and persisted to database without truncation or NPEs:
     total_b2b, total_b2cs, total_b2cl, total_cdnr, total_cdnur, total_nil,
     total_hsn_b2b, total_hsn_b2c, total_ecom,
     total_taxable_value, total_cgst, total_sgst, total_igst, total_cess, total_tax.
   - Empty periods, high numbers, JSON serialization robustness.
"""
import os
import json
import uuid
from datetime import date, datetime
from decimal import Decimal
import openpyxl
import pytest
from flask_login import login_user

from app.extensions import db as _db
from app.models import User, GSTProfile, ImportHistory, RawImport, Transaction, GSTR1Generation
from app.services.gstr1_generator import generate_gstr1, GenerationResult
from app.services.gstr1_json_validator import GSTR1Validator
from app.routes.generate import _save_generation_record


@pytest.fixture
def adversarial_env(app, db):
    """Fixture providing a clean user and profile with distinct GSTIN."""
    suffix = uuid.uuid4().hex[:6]
    rand_digits = f"{uuid.uuid4().int % 9000 + 1000:04d}"
    letter = chr(65 + (uuid.uuid4().int % 26))
    gstin = f"27ABCDE{rand_digits}{letter}1Z5"
    with app.app_context():
        user = User(username=f'adv_user_{suffix}', email=f'adv_{suffix}@example.com')
        user.set_password('AdvPass123!')
        _db.session.add(user)
        _db.session.commit()

        profile = GSTProfile(
            user_id=user.id,
            gstin=gstin,
            legal_name='Adversarial Testing Corp',
            trade_name='AdvCorp',
            state_code='27',
            state_name='Maharashtra',
            financial_year='2024-25',
            filing_frequency='Monthly'
        )
        _db.session.add(profile)
        _db.session.commit()

        return {
            'user': user,
            'user_id': user.id,
            'profile_id': profile.id,
            'gstin': gstin,
            'profile': profile
        }


# =========================================================================
# Track 2 - Category 1: Section Counts & Tax Math Integrity
# =========================================================================

def test_adv_01_total_b2b_strictly_counts_invoices_not_items_or_b2cs(app, db, adversarial_env):
    """Verify total_b2b strictly counts distinct B2B invoices and NOT line items or B2CS."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_sales.xlsx',
            original_file_name='adv_sales.xlsx',
            platform_name='AMAZON',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # Invoice 1: 3 line items for Customer 1 (B2B)
        for i in range(3):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number='INV-B2B-001',
                invoice_date=date(2025, 1, 15),
                supply_type='B2B',
                customer_gstin='27ABCDE1234F1Z5',
                place_of_supply='27',
                hsn_sac=f'847{i}',
                quantity=Decimal('1.00'),
                taxable_value=Decimal('1000.00'),
                tax_rate=Decimal('18.00'),
                cgst_amount=Decimal('90.00'),
                sgst_amount=Decimal('90.00'),
                total_tax=Decimal('180.00'),
                invoice_value=Decimal('1180.00'),
                is_deleted=False
            ))

        # Invoice 2: 2 line items for Customer 2 (B2B inter-state)
        for i in range(2):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number='INV-B2B-002',
                invoice_date=date(2025, 1, 16),
                supply_type='B2B',
                customer_gstin='29ABCDE1234F1Z8',
                place_of_supply='29',
                hsn_sac='8471',
                quantity=Decimal('2.00'),
                taxable_value=Decimal('2000.00'),
                tax_rate=Decimal('18.00'),
                igst_amount=Decimal('360.00'),
                total_tax=Decimal('360.00'),
                invoice_value=Decimal('2360.00'),
                is_deleted=False
            ))

        # 5 B2CS transactions (consumer sales - must NEVER be counted in total_b2b)
        for i in range(5):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number=f'INV-B2CS-{i}',
                invoice_date=date(2025, 1, 20),
                supply_type='B2CS',
                place_of_supply='27',
                taxable_value=Decimal('500.00'),
                tax_rate=Decimal('18.00'),
                cgst_amount=Decimal('45.00'),
                sgst_amount=Decimal('45.00'),
                total_tax=Decimal('90.00'),
                invoice_value=Decimal('590.00'),
                is_deleted=False
            ))

        _db.session.commit()

        res = generate_gstr1(str(profile_id), period, reconciliation_report={'status': 'SUCCESS'})

        # Total B2B must be EXACTLY 2 (2 distinct invoices), NOT 5 (line items), NOT 7, NOT 10
        assert res.stats['total_b2b'] == 2, f"Expected 2 B2B invoices, got {res.stats['total_b2b']}"
        assert res.stats['total_b2cs'] == 1, f"Expected 1 aggregated B2CS row (intra 27 @ 18%), got {res.stats['total_b2cs']}"
        assert res.stats['total_invoices'] == 3, f"Expected total_invoices=3 (2 b2b + 1 b2cs group), got {res.stats['total_invoices']}"


def test_adv_02_stats_computes_all_11_counts_correctly(app, db, adversarial_env):
    """Verify stats dictionary correctly computes all 11 section counts with every section represented."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_all_sections.xlsx',
            original_file_name='adv_all_sections.xlsx',
            platform_name='FLIPKART',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # 1. B2B (1 invoice) with e-commerce GSTIN
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-B2B-10',
            invoice_date=date(2025, 1, 5),
            supply_type='B2B',
            customer_gstin='27ABCDE5678F1Z5',
            place_of_supply='27',
            ecommerce_gstin='27ECOM1234E1Z1',
            hsn_sac='8471',
            taxable_value=Decimal('1000.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('90.00'),
            sgst_amount=Decimal('90.00'),
            total_tax=Decimal('180.00'),
            invoice_value=Decimal('1180.00'),
            is_deleted=False
        ))

        # 2. B2CS (2 different rates -> 2 aggregated rows)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-B2CS-18',
            invoice_date=date(2025, 1, 6),
            supply_type='B2CS',
            place_of_supply='27',
            ecommerce_gstin='27ECOM1234E1Z1',
            hsn_sac='8472',
            taxable_value=Decimal('500.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('45.00'),
            sgst_amount=Decimal('45.00'),
            total_tax=Decimal('90.00'),
            invoice_value=Decimal('590.00'),
            is_deleted=False
        ))
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-B2CS-12',
            invoice_date=date(2025, 1, 7),
            supply_type='B2CS',
            place_of_supply='27',
            hsn_sac='8473',
            taxable_value=Decimal('400.00'),
            tax_rate=Decimal('12.00'),
            cgst_amount=Decimal('24.00'),
            sgst_amount=Decimal('24.00'),
            total_tax=Decimal('48.00'),
            invoice_value=Decimal('448.00'),
            is_deleted=False
        ))

        # 3. CDNR (Registered Credit Note)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-B2B-10',
            note_number='CN-001',
            note_type='C',
            note_date=date(2025, 1, 8),
            supply_type='CDNR',
            customer_gstin='27ABCDE5678F1Z5',
            place_of_supply='27',
            taxable_value=Decimal('200.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('18.00'),
            sgst_amount=Decimal('18.00'),
            total_tax=Decimal('36.00'),
            invoice_value=Decimal('236.00'),
            is_deleted=False
        ))

        # 4. CDNUR (Unregistered Credit Note)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            note_number='CNUR-001',
            note_type='C',
            note_date=date(2025, 1, 9),
            supply_type='CDNUR',
            place_of_supply='29',
            taxable_value=Decimal('300.00'),
            tax_rate=Decimal('18.00'),
            igst_amount=Decimal('54.00'),
            total_tax=Decimal('54.00'),
            invoice_value=Decimal('354.00'),
            is_deleted=False
        ))

        # 5. EXP (Export)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='EXP-001',
            invoice_date=date(2025, 1, 10),
            supply_type='EXPORT',
            place_of_supply='96',
            hsn_sac='8471',
            taxable_value=Decimal('5000.00'),
            tax_rate=Decimal('18.00'),
            igst_amount=Decimal('900.00'),
            total_tax=Decimal('900.00'),
            invoice_value=Decimal('5900.00'),
            is_deleted=False
        ))

        # 6. NIL supply
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='NIL-001',
            invoice_date=date(2025, 1, 11),
            supply_type='NIL',
            place_of_supply='27',
            taxable_value=Decimal('150.00'),
            tax_rate=Decimal('0.00'),
            cgst_amount=Decimal('0.00'),
            sgst_amount=Decimal('0.00'),
            total_tax=Decimal('0.00'),
            invoice_value=Decimal('150.00'),
            is_deleted=False
        ))

        _db.session.commit()

        res = generate_gstr1(str(profile_id), period, include_hsn=True, reconciliation_report={'status': 'SUCCESS'})
        stats = res.stats

        # Check all required keys exist
        required_keys = [
            'total_b2b', 'total_b2cs', 'total_b2cl', 'total_cdnr', 'total_cdnur',
            'total_exp', 'total_nil', 'total_hsn_b2b', 'total_hsn_b2c', 'total_ecom',
            'total_invoices', 'total_taxable_value', 'total_cgst', 'total_sgst',
            'total_igst', 'total_cess', 'total_tax'
        ]
        for k in required_keys:
            assert k in stats, f"Missing key in stats: {k}"

        assert stats['total_b2b'] == 1
        assert stats['total_b2cs'] == 2  # 2 distinct rates (18% and 12%)
        assert stats['total_b2cl'] == 0
        assert stats['total_cdnr'] == 1
        assert stats['total_cdnur'] == 1
        assert stats['total_exp'] == 1
        assert stats['total_nil'] >= 1
        assert stats['total_ecom'] == 2  # 2 transactions had ecommerce_gstin
        expected_total_inv = (
            stats['total_b2b'] + stats['total_b2cs'] + stats['total_b2cl'] +
            stats['total_cdnr'] + stats['total_cdnur'] + stats['total_exp'] + stats['total_nil']
        )
        assert stats['total_invoices'] == expected_total_inv


def test_adv_03_tax_totals_sum_accurately_without_floating_point_drift(app, db, adversarial_env):
    """Stress-test tax totals with 100 fractional decimal transactions (0.07, 0.33, etc.) to detect drift."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_fractional_drift.xlsx',
            original_file_name='adv_fractional_drift.xlsx',
            platform_name='AMAZON',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # Create 100 micro B2CS transactions with amounts known to cause float rounding drift
        # e.g., taxable = 100.33, rate = 18%, cgst = 9.03, sgst = 9.03, cess = 1.01
        expected_taxable = Decimal('0.00')
        expected_cgst = Decimal('0.00')
        expected_sgst = Decimal('0.00')
        expected_cess = Decimal('0.00')

        for i in range(100):
            txval = Decimal('100.33')
            cgst = Decimal('9.03')
            sgst = Decimal('9.03')
            cess = Decimal('1.01')
            tot_tax = cgst + sgst + cess
            inv_val = txval + tot_tax

            expected_taxable += txval
            expected_cgst += cgst
            expected_sgst += sgst
            expected_cess += cess

            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number=f'INV-FRAC-{i}',
                invoice_date=date(2025, 1, 1 + (i % 25)),
                supply_type='B2CS',
                place_of_supply='27',
                hsn_sac='8471',
                taxable_value=txval,
                tax_rate=Decimal('18.00'),
                cgst_amount=cgst,
                sgst_amount=sgst,
                cess_amount=cess,
                total_tax=tot_tax,
                invoice_value=inv_val,
                is_deleted=False
            ))

        _db.session.commit()

        res = generate_gstr1(str(profile_id), period, reconciliation_report={'status': 'SUCCESS'})
        stats = res.stats

        # 100 * 100.33 = 10033.00
        assert expected_taxable == Decimal('10033.00')
        assert expected_cgst == Decimal('903.00')
        assert expected_sgst == Decimal('903.00')
        assert expected_cess == Decimal('101.00')
        expected_total_tax = expected_cgst + expected_sgst + expected_cess  # 1907.00

        # Verify stats matches exact cents
        assert round(stats['total_taxable_value'], 2) == float(expected_taxable)
        assert round(stats['total_cgst'], 2) == float(expected_cgst)
        assert round(stats['total_sgst'], 2) == float(expected_sgst)
        assert round(stats['total_cess'], 2) == float(expected_cess)
        assert round(stats['total_tax'], 2) == float(expected_total_tax)

        # Critical math property: total_tax == total_cgst + total_sgst + total_igst + total_cess
        computed_tax_sum = round(stats['total_cgst'] + stats['total_sgst'] + stats['total_igst'] + stats['total_cess'], 2)
        assert round(stats['total_tax'], 2) == computed_tax_sum


# =========================================================================
# Track 2 - Category 2: Option Toggles & Schema Validation
# =========================================================================

def test_adv_04_include_hsn_false_omits_hsn_from_excel_and_passes_validator(app, db, adversarial_env):
    """Verify include_hsn=False omits HSN data, omits HSN sheet from Excel, and GSTR1Validator accepts payload."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_hsn_toggle.xlsx',
            original_file_name='adv_hsn_toggle.xlsx',
            platform_name='AMAZON',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-B2B-HSN',
            invoice_date=date(2025, 1, 10),
            supply_type='B2B',
            customer_gstin='27GHIJK1234L1Z9',
            place_of_supply='27',
            hsn_sac='84713010',
            taxable_value=Decimal('50000.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('4500.00'),
            sgst_amount=Decimal('4500.00'),
            total_tax=Decimal('9000.00'),
            invoice_value=Decimal('59000.00'),
            is_deleted=False
        ))
        _db.session.commit()

        # Run with include_hsn=False
        res = generate_gstr1(str(profile_id), period, include_hsn=False, reconciliation_report={'status': 'SUCCESS'})

        # 1. Check stats
        assert res.stats['total_hsn_b2b'] == 0
        assert res.stats['total_hsn_b2c'] == 0

        # 2. Check JSON file on disk
        assert os.path.exists(res.json_path)
        with open(res.json_path, 'r', encoding='utf-8') as f:
            json_content = f.read()
            data = json.loads(json_content)

        # HSN data must be empty
        hsn_records = data.get('hsn', {}).get('data', [])
        assert len(hsn_records) == 0, f"Expected 0 HSN records in JSON when include_hsn=False, got {len(hsn_records)}"

        # 3. Check Excel file on disk
        assert os.path.exists(res.excel_path)
        wb = openpyxl.load_workbook(res.excel_path)
        # When include_hsn=False, the 'hsn' sheet must NOT be created
        assert 'hsn' not in wb.sheetnames, f"'hsn' sheet should be omitted from Excel workbook, but found in {wb.sheetnames}"
        assert 'hsnb2c' not in wb.sheetnames, f"'hsnb2c' sheet should be omitted from Excel workbook, but found in {wb.sheetnames}"

        # 4. JSON Validator must still accept the payload (is_valid == True)
        val_res = GSTR1Validator.validate_gstr1_json(json_content)
        assert val_res.is_valid is True, f"JSON Validator rejected payload with include_hsn=False: {val_res.errors}"


def test_adv_05_include_hsn_true_summary_matches_item_level_taxes(app, db, adversarial_env):
    """Verify include_hsn=True produces HSN summary exactly matching item-level taxes."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_hsn_match.xlsx',
            original_file_name='adv_hsn_match.xlsx',
            platform_name='MEESHO',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # Add 3 transactions under HSN 6109 (T-Shirts) at 12%
        for i in range(3):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number=f'INV-TSHIRT-{i}',
                invoice_date=date(2025, 1, 10 + i),
                supply_type='B2CS',
                place_of_supply='27',
                hsn_sac='6109',
                quantity=Decimal('5.00'),
                taxable_value=Decimal('1000.00'),
                tax_rate=Decimal('12.00'),
                cgst_amount=Decimal('60.00'),
                sgst_amount=Decimal('60.00'),
                total_tax=Decimal('120.00'),
                invoice_value=Decimal('1120.00'),
                is_deleted=False
            ))

        # Add 2 transactions under HSN 8471 (Computers) at 18% inter-state
        for i in range(2):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number=f'INV-COMP-{i}',
                invoice_date=date(2025, 1, 15 + i),
                supply_type='B2B',
                customer_gstin='29ABCDE5678F1Z2',
                place_of_supply='29',
                hsn_sac='8471',
                quantity=Decimal('1.00'),
                taxable_value=Decimal('25000.00'),
                tax_rate=Decimal('18.00'),
                igst_amount=Decimal('4500.00'),
                total_tax=Decimal('4500.00'),
                invoice_value=Decimal('29500.00'),
                is_deleted=False
            ))

        _db.session.commit()

        # Run with include_hsn=True
        res = generate_gstr1(str(profile_id), period, include_hsn=True, reconciliation_report={'status': 'SUCCESS'})

        with open(res.json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        hsn_list = data.get('hsn', {}).get('data', [])
        assert len(hsn_list) == 2, f"Expected 2 HSN summary lines, got {len(hsn_list)}"

        # Check HSN sums
        hsn_txval = sum(Decimal(str(h['txval'])) for h in hsn_list)
        hsn_cgst = sum(Decimal(str(h['camt'])) for h in hsn_list)
        hsn_sgst = sum(Decimal(str(h['samt'])) for h in hsn_list)
        hsn_igst = sum(Decimal(str(h['iamt'])) for h in hsn_list)

        expected_txval = Decimal('3000.00') + Decimal('50000.00')
        expected_cgst = Decimal('180.00')
        expected_sgst = Decimal('180.00')
        expected_igst = Decimal('9000.00')

        assert hsn_txval == expected_txval, f"HSN taxable {hsn_txval} != expected {expected_txval}"
        assert hsn_cgst == expected_cgst, f"HSN CGST {hsn_cgst} != expected {expected_cgst}"
        assert hsn_sgst == expected_sgst, f"HSN SGST {hsn_sgst} != expected {expected_sgst}"
        assert hsn_igst == expected_igst, f"HSN IGST {hsn_igst} != expected {expected_igst}"

        # Excel should have 'hsn' sheet populated
        wb = openpyxl.load_workbook(res.excel_path)
        assert 'hsn' in wb.sheetnames
        ws_hsn = wb['hsn']
        assert ws_hsn.max_row >= 3  # Header + 2 data rows


# =========================================================================
# Track 2 - Category 3: Persistence Integrity (All 15 Columns)
# =========================================================================

def test_adv_06_gstr1_generation_persistence_all_15_columns(app, db, adversarial_env):
    """Verify that all 15 columns in GSTR1Generation are correctly populated and persisted."""
    with app.test_request_context():
        login_user(adversarial_env['user'])
        profile = adversarial_env['profile']
        profile_id = adversarial_env['profile_id']
        period = '012025'

        # Create a mock GenerationResult with known stats
        stats = {
            'total_b2b': 5,
            'total_b2cs': 12,
            'total_b2cl': 2,
            'total_cdnr': 3,
            'total_cdnur': 1,
            'total_nil': 4,
            'total_hsn_b2b': 5,
            'total_hsn_b2c': 8,
            'total_ecom': 7,
            'total_invoices': 27,
            'total_taxable_value': 876543.21,
            'total_cgst': 78888.89,
            'total_sgst': 78888.89,
            'total_igst': 12345.67,
            'total_cess': 543.21,
            'total_tax': 170666.66
        }

        mock_gen_result = GenerationResult(
            generation_id='test-uuid-adv',
            excel_path='/tmp/fake.xlsx',
            json_path='/tmp/fake.json',
            validation_result=None,
            reconciliation_report={'status': 'CLEAN', 'differences': []},
            stats=stats
        )

        gen_rec = _save_generation_record(
            profile=profile,
            return_period=period,
            gen_result=mock_gen_result,
            recon_report={'status': 'CLEAN'}
        )

        assert gen_rec.id is not None

        # Re-query fresh from DB session to verify DB persistence
        saved = _db.session.get(GSTR1Generation, gen_rec.id)
        assert saved is not None

        # 9 Section Count Columns
        assert saved.total_b2b == 5
        assert saved.total_b2cs == 12
        assert saved.total_b2cl == 2
        assert saved.total_cdnr == 3
        assert saved.total_cdnur == 1
        assert saved.total_nil == 4
        assert saved.total_hsn_b2b == 5
        assert saved.total_hsn_b2c == 8
        assert saved.total_ecom == 7

        # 6 Tax Total Columns
        assert float(saved.total_taxable_value) == pytest.approx(876543.21, 0.001)
        assert float(saved.total_cgst) == pytest.approx(78888.89, 0.001)
        assert float(saved.total_sgst) == pytest.approx(78888.89, 0.001)
        assert float(saved.total_igst) == pytest.approx(12345.67, 0.001)
        assert float(saved.total_cess) == pytest.approx(543.21, 0.001)
        assert float(saved.total_tax) == pytest.approx(170666.66, 0.001)

        # Audit & Meta columns
        assert saved.generation_status == 'COMPLETED'
        assert saved.schema_version == '1.1'
        assert saved.rule_version == '2024-25'
        assert saved.financial_year == '2024-25'
        assert json.loads(saved.reconciliation_report) == {'status': 'CLEAN'}


def test_adv_07_persistence_handles_zero_transactions_without_null_pointer(app, db, adversarial_env):
    """Verify persistence on empty return period stores 0 and 0.00 without NULL or crash."""
    with app.test_request_context():
        login_user(adversarial_env['user'])
        profile = adversarial_env['profile']
        profile_id = adversarial_env['profile_id']
        period = '022025'

        res = generate_gstr1(str(profile_id), period, reconciliation_report={'status': 'SUCCESS'})
        gen_rec = _save_generation_record(
            profile=profile,
            return_period=period,
            gen_result=res,
            recon_report={'status': 'CLEAN'}
        )

        saved = _db.session.get(GSTR1Generation, gen_rec.id)
        assert saved is not None

        # Verify none of the 15 columns are None
        for col_name in [
            'total_b2b', 'total_b2cs', 'total_b2cl', 'total_cdnr', 'total_cdnur',
            'total_nil', 'total_hsn_b2b', 'total_hsn_b2c', 'total_ecom'
        ]:
            val = getattr(saved, col_name)
            assert val is not None, f"Column {col_name} is None"
            assert val == 0, f"Column {col_name} should be 0, got {val}"

        for col_name in [
            'total_taxable_value', 'total_cgst', 'total_sgst', 'total_igst', 'total_cess', 'total_tax'
        ]:
            val = getattr(saved, col_name)
            assert val is not None, f"Column {col_name} is None"
            assert float(val) == 0.0, f"Column {col_name} should be 0.0, got {val}"


def test_adv_08_persistence_large_values_boundary_test(app, db, adversarial_env):
    """Boundary test: 10-digit taxable values up to 9,999,999,999.99 do not overflow Numeric(15, 2)."""
    with app.test_request_context():
        login_user(adversarial_env['user'])
        profile = adversarial_env['profile']
        profile_id = adversarial_env['profile_id']
        period = '032025'

        large_taxable = Decimal('9999999999.99')
        large_tax = Decimal('1800000000.00')

        stats = {
            'total_b2b': 10000,
            'total_b2cs': 50000,
            'total_b2cl': 500,
            'total_cdnr': 200,
            'total_cdnur': 100,
            'total_nil': 50,
            'total_hsn_b2b': 2000,
            'total_hsn_b2c': 1500,
            'total_ecom': 25000,
            'total_invoices': 60850,
            'total_taxable_value': float(large_taxable),
            'total_cgst': float(large_tax / 2),
            'total_sgst': float(large_tax / 2),
            'total_igst': 0.0,
            'total_cess': 0.0,
            'total_tax': float(large_tax)
        }

        mock_gen_result = GenerationResult(
            generation_id='test-large-boundary',
            excel_path='/tmp/large.xlsx',
            json_path='/tmp/large.json',
            validation_result=None,
            reconciliation_report={'status': 'CLEAN'},
            stats=stats
        )

        gen_rec = _save_generation_record(
            profile=profile,
            return_period=period,
            gen_result=mock_gen_result,
            recon_report={'status': 'CLEAN'}
        )

        saved = _db.session.get(GSTR1Generation, gen_rec.id)
        assert saved.total_taxable_value == large_taxable
        assert saved.total_tax == large_tax
        assert saved.total_b2b == 10000


def test_adv_09_may2025_separate_hsn_reporting_mode(app, db, adversarial_env):
    """Verify period 052025 (May 2025+) separates HSN into B2B and B2C tables and sheets."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '052025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_may2025_hsn.xlsx',
            original_file_name='adv_may2025_hsn.xlsx',
            platform_name='AMAZON',
            financial_year='2025-26',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # B2B Transaction with HSN 8471
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-MAY-B2B',
            invoice_date=date(2025, 5, 10),
            supply_type='B2B',
            customer_gstin='27ABCDE1234F1Z5',
            place_of_supply='27',
            hsn_sac='8471',
            taxable_value=Decimal('10000.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('900.00'),
            sgst_amount=Decimal('900.00'),
            total_tax=Decimal('1800.00'),
            invoice_value=Decimal('11800.00'),
            is_deleted=False
        ))

        # B2CS Transaction with HSN 6109
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-MAY-B2C',
            invoice_date=date(2025, 5, 15),
            supply_type='B2CS',
            place_of_supply='27',
            hsn_sac='6109',
            taxable_value=Decimal('5000.00'),
            tax_rate=Decimal('12.00'),
            cgst_amount=Decimal('300.00'),
            sgst_amount=Decimal('300.00'),
            total_tax=Decimal('600.00'),
            invoice_value=Decimal('5600.00'),
            is_deleted=False
        ))
        _db.session.commit()

        res = generate_gstr1(str(profile_id), period, include_hsn=True, reconciliation_report={'status': 'SUCCESS'})

        assert res.stats['total_hsn_b2b'] == 1, f"Expected 1 B2B HSN, got {res.stats['total_hsn_b2b']}"
        assert res.stats['total_hsn_b2c'] == 1, f"Expected 1 B2C HSN, got {res.stats['total_hsn_b2c']}"

        with open(res.json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        assert 'data' in data['hsn']
        assert 'b2c_data' in data['hsn']
        assert len(data['hsn']['data']) == 1
        assert len(data['hsn']['b2c_data']) == 1

        # Check Excel sheets
        wb = openpyxl.load_workbook(res.excel_path)
        assert 'hsn' in wb.sheetnames
        assert 'hsnb2c' in wb.sheetnames

        # Validation passes
        val_res = GSTR1Validator.validate_gstr1_json(json.dumps(data))
        assert val_res.is_valid is True, f"May 2025 JSON failed validation: {val_res.errors}"


def test_adv_10_extreme_float_drift_b2cs_accumulation(app, db, adversarial_env):
    """Stress-test B2CS float accumulation with non-binary fractions (0.1, 0.2) to evaluate float drift."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_float_drift.xlsx',
            original_file_name='adv_float_drift.xlsx',
            platform_name='AMAZON',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # Add 30 identical transactions with taxable=0.10, tax_rate=18%, cgst=0.01, sgst=0.01
        for i in range(30):
            _db.session.add(Transaction(
                profile_id=profile_id,
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number=f'INV-DRIFT-{i}',
                invoice_date=date(2025, 1, 1),
                supply_type='B2CS',
                place_of_supply='27',
                taxable_value=Decimal('0.10'),
                tax_rate=Decimal('18.00'),
                cgst_amount=Decimal('0.01'),
                sgst_amount=Decimal('0.01'),
                total_tax=Decimal('0.02'),
                invoice_value=Decimal('0.12'),
                is_deleted=False
            ))
        _db.session.commit()

        res = generate_gstr1(str(profile_id), period, reconciliation_report={'status': 'SUCCESS'})
        stats = res.stats

        # 30 * 0.10 = 3.00, 30 * 0.01 = 0.30
        assert abs(stats['total_taxable_value'] - 3.00) < 1e-5
        assert abs(stats['total_cgst'] - 0.30) < 1e-5
        assert abs(stats['total_sgst'] - 0.30) < 1e-5
        assert abs(stats['total_tax'] - 0.60) < 1e-5


def test_adv_11_api_generate_populates_all_15_columns_and_returns_stats(client, app, db, adversarial_env):
    """Verify POST /api/generate creates a database record with all 15 columns populated and returns stats."""
    _login_user = adversarial_env['user']
    with app.test_client() as c:
        with c.session_transaction() as sess:
            sess['_user_id'] = str(_login_user.id)
            sess['active_profile_id'] = adversarial_env['profile_id']

        period = '012025'
        with app.app_context():
            imp = ImportHistory(
                profile_id=adversarial_env['profile_id'],
                user_id=adversarial_env['user_id'],
                return_period=period,
                file_name='api_test.xlsx',
                original_file_name='api_test.xlsx',
                platform_name='AMAZON',
                financial_year='2024-25',
                processing_status='COMPLETED'
            )
            _db.session.add(imp)
            _db.session.commit()

            raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
            _db.session.add(raw)
            _db.session.commit()

            _db.session.add(Transaction(
                profile_id=adversarial_env['profile_id'],
                import_history_id=imp.id,
                raw_import_id=raw.id,
                invoice_number='INV-API-01',
                invoice_date=date(2025, 1, 10),
                supply_type='B2B',
                customer_gstin='27ABCDE1234F1Z5',
                place_of_supply='27',
                hsn_sac='8471',
                taxable_value=Decimal('2000.00'),
                tax_rate=Decimal('18.00'),
                cgst_amount=Decimal('180.00'),
                sgst_amount=Decimal('180.00'),
                total_tax=Decimal('360.00'),
                invoice_value=Decimal('2360.00'),
                is_deleted=False
            ))
            _db.session.commit()

        resp = c.post('/api/generate', json={'return_period': period, 'include_hsn': True})
        assert resp.status_code == 201
        data = resp.get_json()['data']

        # Verify response contains all 15 column values
        expected_fields = [
            'total_b2b', 'total_b2cs', 'total_b2cl', 'total_cdnr', 'total_cdnur',
            'total_nil', 'total_hsn_b2b', 'total_hsn_b2c', 'total_ecom',
            'total_taxable_value', 'total_cgst', 'total_sgst', 'total_igst', 'total_cess', 'total_tax'
        ]
        for field in expected_fields:
            assert field in data, f"Missing field in API response: {field}"
            assert data[field] is not None, f"Field {field} is None in API response"

        # Check DB row
        with app.app_context():
            gen_db = _db.session.get(GSTR1Generation, data['id'])
            assert gen_db is not None
            assert gen_db.total_b2b == 1
            assert float(gen_db.total_taxable_value) == 2000.0
            assert float(gen_db.total_cgst) == 180.0
            assert float(gen_db.total_sgst) == 180.0
            assert float(gen_db.total_tax) == 360.0


def test_adv_12_amendment_transactions_section_counts(app, db, adversarial_env):
    """Verify amendment transactions (B2BA, CDNRA, CDNURA, EXPA) are counted in stats."""
    with app.app_context():
        profile_id = adversarial_env['profile_id']
        period = '012025'

        imp = ImportHistory(
            profile_id=profile_id,
            user_id=adversarial_env['user_id'],
            return_period=period,
            file_name='adv_amendments.xlsx',
            original_file_name='adv_amendments.xlsx',
            platform_name='AMAZON',
            financial_year='2024-25',
            processing_status='COMPLETED'
        )
        _db.session.add(imp)
        _db.session.commit()

        raw = RawImport(import_history_id=imp.id, sheet_name='Sheet1', row_number=1, status='SUCCESS')
        _db.session.add(raw)
        _db.session.commit()

        # Regular B2B
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-REG-1',
            invoice_date=date(2025, 1, 5),
            supply_type='B2B',
            customer_gstin='27ABCDE1234F1Z5',
            place_of_supply='27',
            taxable_value=Decimal('1000.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('90.00'),
            sgst_amount=Decimal('90.00'),
            total_tax=Decimal('180.00'),
            invoice_value=Decimal('1180.00'),
            is_deleted=False
        ))

        # B2B Amendment (B2BA)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-AMD-1',
            original_invoice_number='INV-OLD-1',
            invoice_date=date(2025, 1, 6),
            supply_type='B2BA',
            customer_gstin='27ABCDE1234F1Z5',
            place_of_supply='27',
            taxable_value=Decimal('1500.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('135.00'),
            sgst_amount=Decimal('135.00'),
            total_tax=Decimal('270.00'),
            invoice_value=Decimal('1770.00'),
            is_deleted=False
        ))

        # CDNR Amendment (CDNRA)
        _db.session.add(Transaction(
            profile_id=profile_id,
            import_history_id=imp.id,
            raw_import_id=raw.id,
            invoice_number='INV-REG-1',
            note_number='CN-AMD-1',
            original_invoice_number='CN-OLD-1',
            note_type='C',
            note_date=date(2025, 1, 7),
            supply_type='CDNRA',
            customer_gstin='27ABCDE1234F1Z5',
            place_of_supply='27',
            taxable_value=Decimal('200.00'),
            tax_rate=Decimal('18.00'),
            cgst_amount=Decimal('18.00'),
            sgst_amount=Decimal('18.00'),
            total_tax=Decimal('36.00'),
            invoice_value=Decimal('236.00'),
            is_deleted=False
        ))
        _db.session.commit()

        clean_report = {
            'status': 'SUCCESS',
            'is_generation_blocked': False,
            'critical_failures': 0,
            'warnings': 0,
            'differences': []
        }
        res = generate_gstr1(str(profile_id), period, reconciliation_report=clean_report)

        # total_b2b must count both B2B and B2BA (1 + 1 = 2)
        assert res.stats['total_b2b'] == 2, f"Expected total_b2b=2 (1 B2B + 1 B2BA), got {res.stats['total_b2b']}"
        # total_cdnr must count CDNRA (1)
        assert res.stats['total_cdnr'] == 1, f"Expected total_cdnr=1 (CDNRA), got {res.stats['total_cdnr']}"

